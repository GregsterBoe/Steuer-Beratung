"""Prüft die erzeugten Mappen auf Zirkelbezüge, so wie Excel sie sieht.

Aufruf: python -m pruefung.zirkelbezug
Läuft ohne LibreOffice. Beendet sich mit Fehlercode 1, wenn ein Ring gefunden wird.

Excel hängt eine Formel an jede Zelle eines Bereichs, den sie nennt (SUMIFS, INDEX,
MATCH über einen Namen), auch an Zweige von IF, die gerade nicht greifen. LibreOffice
rechnet solche Ringe oft trotzdem durch; Excel meldet den Zirkelbezug und rechnet die
beteiligten Zellen nicht mehr. Die Prüfung baut daher den Abhängigkeitsgraphen aus
allen Bezügen und sucht darin Ringe (stark zusammenhängende Komponenten).
"""

import subprocess
import sys
import tempfile
from collections import Counter, deque
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils import get_column_letter, range_boundaries

KOSTENSTELLEN = Path(__file__).resolve().parent.parent / "vorlagen" / "Kostenstellen_BWA_Vorlage.xlsx"
MAPPEN = {
    "mit Testdaten": [],
    "mit Kostenstellen": ["--kostenstellen", str(KOSTENSTELLEN)],
    "Schnellcheck": ["--schnellcheck", "--kostenstellen", str(KOSTENSTELLEN)],
}


def _bereiche(ref: str, blatt: str, groesse: dict) -> list:
    """(Blatt, Spalte1, Zeile1, Spalte2, Zeile2) je Teil eines Bezugs, an das Blatt gekappt."""
    out = []
    for teil in ref.split(","):
        teil = teil.strip()
        if "!" in teil:
            bl, r = teil.rsplit("!", 1)
            bl = bl.strip("'").replace("''", "'")
        else:
            bl, r = blatt, teil
        if bl not in groesse:
            continue
        zeilen, spalten = groesse[bl]
        try:
            c1, r1, c2, r2 = range_boundaries(r.replace("$", ""))
        except (ValueError, TypeError):
            continue
        out.append((bl, c1 or 1, r1 or 1, min(c2 or spalten, spalten), min(r2 or zeilen, zeilen)))
    return out


def graph(wb) -> dict:
    """Zelle -> Vorgänger; Bereiche als eigene Knoten, damit große Namen nur einmal zählen."""
    namen = {n.lower(): dn.attr_text for n, dn in wb.defined_names.items()}
    groesse = {ws.title: (ws.max_row, ws.max_column) for ws in wb.worksheets}
    g, bereiche = {}, {}
    for ws in wb.worksheets:
        for zeile in ws.iter_rows():
            for c in zeile:
                if not (isinstance(c.value, str) and c.value.startswith("=")):
                    continue
                vor = set()
                for t in Tokenizer(c.value).items:
                    if t.type != "OPERAND" or t.subtype != "RANGE":
                        continue
                    ref = namen.get(t.value.lower(), t.value)
                    for b in _bereiche(ref, ws.title, groesse):
                        if b[1:3] == b[3:5]:
                            vor.add(b[:3])
                        else:
                            knoten = ("Bereich",) + b
                            vor.add(knoten)
                            bereiche.setdefault(knoten, b)
                g[(ws.title, c.column, c.row)] = vor
    for knoten, (bl, c1, r1, c2, r2) in bereiche.items():
        g[knoten] = [z for z in ((bl, c, r) for c in range(c1, c2 + 1) for r in range(r1, r2 + 1))
                     if z in g]
    return g


def ringe(g: dict) -> list:
    """Stark zusammenhängende Komponenten mit mehr als einem Knoten (Tarjan, iterativ)."""
    index, low, auf_stapel, stapel, gefunden, i = {}, {}, set(), [], [], 0
    for start in g:
        if start in index:
            continue
        index[start] = low[start] = i
        i += 1
        stapel.append(start)
        auf_stapel.add(start)
        arbeit = [(start, iter(g[start]))]
        while arbeit:
            v, it = arbeit[-1]
            for w in it:
                if w not in g:
                    continue
                if w not in index:
                    index[w] = low[w] = i
                    i += 1
                    stapel.append(w)
                    auf_stapel.add(w)
                    arbeit.append((w, iter(g[w])))
                    break
                if w in auf_stapel:
                    low[v] = min(low[v], index[w])
            else:
                arbeit.pop()
                if arbeit:
                    low[arbeit[-1][0]] = min(low[arbeit[-1][0]], low[v])
                if low[v] == index[v]:
                    komp = []
                    while True:
                        w = stapel.pop()
                        auf_stapel.discard(w)
                        komp.append(w)
                        if w == v:
                            break
                    if len(komp) > 1 or v in g[v]:
                        gefunden.append(komp)
    return gefunden


def _name(knoten) -> str:
    if knoten[0] == "Bereich":
        _, bl, c1, r1, c2, r2 = knoten
        return f"  über {bl}!{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}"
    return f"{knoten[0]}!{get_column_letter(knoten[1])}{knoten[2]}"


def kurzer_ring(g: dict, komp: list) -> list:
    """Ein kürzester Ring durch die erste Zelle der Komponente, zum Nachlesen in der Mappe."""
    menge = set(komp)
    start = next(k for k in komp if k[0] != "Bereich")
    vorher, offen = {start: None}, deque([start])
    while offen:
        v = offen.popleft()
        for w in g[v]:
            if w == start:
                pfad = []
                while v is not None:
                    pfad.append(v)
                    v = vorher[v]
                return pfad[::-1]
            if w in menge and w not in vorher:
                vorher[w] = v
                offen.append(w)
    return [start]


def main() -> int:
    fehler = 0
    with tempfile.TemporaryDirectory() as tmp:
        for titel, argumente in MAPPEN.items():
            datei = Path(tmp) / "mappe.xlsx"
            subprocess.run([sys.executable, "-m", "prognosemodell", *argumente,
                            "--ausgabe", str(datei)], check=True, capture_output=True)
            wb = load_workbook(datei)
            g = graph(wb)
            gefunden = ringe(g)
            print(f"{'OK  ' if not gefunden else 'FEHL'} {titel}: {len(gefunden)} Zirkelbezug/-bezüge")
            for komp in gefunden:
                fehler += 1
                spalten = Counter(f"{k[0]}!{get_column_letter(k[1])}" for k in komp
                                  if k[0] != "Bereich")
                print("     Spalten:", ", ".join(f"{s} ({n})" for s, n in sorted(spalten.items())))
                for k in kurzer_ring(g, komp):
                    print("     " + _name(k))
    print(f"\n{fehler} Zirkelbezug/-bezüge" if fehler else "\nKeine Zirkelbezüge.")
    return 1 if fehler else 0


if __name__ == "__main__":
    sys.exit(main())
