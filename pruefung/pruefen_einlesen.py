"""Prüfskript für die Einleseschicht: synthetische Kostenstellenblätter im Musterlayout.

Aufruf: python -m pruefung.pruefen_einlesen
Läuft ohne LibreOffice. Beendet sich mit Fehlercode 1 bei Abweichung.
"""

import dataclasses
import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook

from prognosemodell.einlesen import (EinleseFehler, anlagen_zusammenfuehren, lese_inventar,
                                     lese_kostenstellen, ordne_anlagen_zu, stand_aus_dateiname,
                                     zusammenfuehren)
from prognosemodell.modelle import Anlage, Objekt
from prognosemodell.testdaten import testobjekt
from prognosemodell.vorlagen import (INVENTAR_BEISPIELE, erstelle_inventar_vorlage,
                                     erstelle_vorlage)

TOLERANZ = 0.01

# BWA-Nr. -> Wert in der Jahresspalte 2026 (frei erfunden, nur Layout wie im Muster)
WERTE = {
    1020: 120_000.0, 1051: 120_000.0, 1090: 1_500.0,
    1100: 2_000.0, 1120: 10_000.0, 1140: 3_000.0, 1150: 4_000.0,
    1240: 16_000.0, 1250: -500.0, 1260: 1_000.0,
    1280: 36_000.0, 1310: 9_000.0,
}


def kostenstellenblatt(ws, kst: str, objekt: str, werte: dict, formel_statt_wert=None):
    """Blatt im Layout des Musters: B2/C2 Kopf, Zeile 4 Spaltenköpfe, ab Zeile 6 BWA-Zeilen."""
    ws["B2"], ws["C2"] = kst, objekt
    ws["B4"], ws["C4"] = "Nr.", "Bezeichnung kurz"
    ws["F4"], ws["S4"] = 2025, 2026  # G–R wären die Monate, T ff. Folgejahre
    for zeile, nr in enumerate(sorted(werte), start=6):
        ws.cell(row=zeile, column=2, value=nr)
        ws.cell(row=zeile, column=19,
                value="=SUM(G{0}:R{0})".format(zeile) if nr == formel_statt_wert
                else werte[nr])


def schreibe(pfad: Path, blaetter: list, annahmen: bool = False) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    if annahmen:  # Fremdblätter: B2 belegt; einmal ohne "Nr.", einmal mit, aber ohne Jahr
        ws = wb.create_sheet("Annahmen")
        ws["B2"], ws["B4"], ws["C4"] = "Mietsteigerung", "Jahr", 2025
        ws = wb.create_sheet("Übersicht")
        ws["B2"], ws["B4"], ws["C4"] = "Summe", "Nr.", "Bezeichnung"
    for titel, kst, objekt, werte, formel in blaetter:
        kostenstellenblatt(wb.create_sheet(titel), kst, objekt, werte, formel)
    wb.save(pfad)
    return pfad


def main() -> int:
    fehler = 0

    def pruefe(fall, ist, soll):
        nonlocal fehler
        gut = abs(ist - soll) <= TOLERANZ if isinstance(soll, float) else ist == soll
        print(f"{'OK    ' if gut else 'FEHLER'}  {fall}: Soll {soll!r}, Ist {ist!r}")
        fehler += not gut

    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        ohne_1090 = {k: v for k, v in WERTE.items() if k != 1090}
        pfad = schreibe(tmp / "kst.xlsx", [
            ("KSt 1", "KSt 1", "KC 24+26", WERTE, None),
            ("KSt 2", "KSt 2", "Objekt B", ohne_1090, None),
        ], annahmen=True)
        (lw1, lw2), uebersprungen = lese_kostenstellen(pfad, 2026)
        pruefe("Blätter mit anderem Format werden mit Grund übersprungen", uebersprungen, [
            ("Annahmen", "kein 'Nr.' in Zeile 4, Spalte B"),
            ("Übersicht", "keine Spalte 2026 in Zeile 4")])
        pruefe("ID aus B2", lw1.objekt_id, "KSt 1")
        pruefe("Name aus C2", lw1.name, "KC 24+26")
        pruefe("Miete = BWA 1020", lw1.miete, 120_000.0)
        pruefe("weitere Einnahmen = BWA 1090", lw1.weitere_einnahmen, 1_500.0)
        pruefe("Erhaltung = BWA 1250 (auch negativ)", lw1.erhaltung, -500.0)
        pruefe("weitere Ausgaben = 1100–1220 + 1260, ohne AfA/Erhaltung/Zins",
               lw1.weitere_ausgaben, 20_000.0)
        pruefe("AfA lt. BWA 1240", lw1.abschreibung, 16_000.0)
        pruefe("fehlende BWA-Zeile zählt 0", lw2.weitere_einnahmen, 0.0)

        try:
            lese_kostenstellen(pfad, 2030)
            pruefe("Basisjahr in keinem Blatt meldet Fehler", "kein Fehler", "EinleseFehler")
        except EinleseFehler:
            pruefe("Basisjahr in keinem Blatt meldet Fehler", "EinleseFehler", "EinleseFehler")

        nur_annahmen = schreibe(tmp / "annahmen.xlsx", [], annahmen=True)
        try:
            lese_kostenstellen(nur_annahmen, 2026)
            pruefe("ohne Kostenstellenblatt meldet Fehler", "kein Fehler", "EinleseFehler")
        except EinleseFehler:
            pruefe("ohne Kostenstellenblatt meldet Fehler", "EinleseFehler", "EinleseFehler")

        doppelt = schreibe(tmp / "doppelt.xlsx", [
            ("A", "KSt 1", "x", WERTE, None), ("B", "KSt 1", "y", WERTE, None)])
        try:
            lese_kostenstellen(doppelt, 2026)
            pruefe("doppelte Kostenstelle meldet Fehler", "kein Fehler", "EinleseFehler")
        except EinleseFehler:
            pruefe("doppelte Kostenstelle meldet Fehler", "EinleseFehler", "EinleseFehler")

        ungerechnet = schreibe(tmp / "formel.xlsx", [("A", "KSt 1", "x", WERTE, 1020)])
        try:
            lese_kostenstellen(ungerechnet, 2026)
            pruefe("Formel ohne Wert meldet Fehler", "kein Fehler", "EinleseFehler")
        except EinleseFehler:
            pruefe("Formel ohne Wert meldet Fehler", "EinleseFehler", "EinleseFehler")

        # Zielformat der Planungsreferenz: Kopf "Jahr 2026", Monatsspalten als Datum,
        # Planspalten "Plan 2027", 2028 …; Summenblatt "Alle Objekte" wird übersprungen
        vorlage = tmp / "vorlage.xlsx"
        erstelle_vorlage(basisjahr=2026).save(vorlage)
        (v1, v2), uebersprungen = lese_kostenstellen(vorlage, 2026)
        pruefe("Vorlage: Summenblatt übersprungen", uebersprungen,
               [("Alle Objekte", "Summenblatt aller Kostenstellen")])
        pruefe("Vorlage: Kostenstellen", (v1.objekt_id, v2.objekt_id), ("KSt 1", "KSt 2"))
        pruefe("Vorlage: Name aus C2", v1.name, "Musterstraße 1")
        pruefe("Vorlage: Miete aus Spalte 'Jahr 2026', nicht aus Vorjahr oder Monat",
               v1.miete, 60_000.0)
        pruefe("Vorlage: Erhaltung", v1.erhaltung, 8_000.0)
        pruefe("Vorlage: weitere Ausgaben 1140 + 1150 + 1260", v1.weitere_ausgaben, 3_600.0)
        pruefe("Vorlage: weitere Einnahmen KSt 2", v2.weitere_einnahmen, 1_500.0)
        pruefe("Vorlage: Planspalten leer, kein AfA-Plan", v1.afa_plan, {})

        # schon geplante AfA in den Planspalten (BWA 1240), Lücke 2029 bleibt dem Modell
        geplant = tmp / "geplant.xlsx"
        erstelle_vorlage(basisjahr=2026, plan={"KSt 1": {
            1240: {2027: 15_000, 2028: 14_000.5, 2030: 0}, 1020: {2027: 99}}}).save(geplant)
        g1 = lese_kostenstellen(geplant, 2026)[0][0]
        pruefe("AfA-Plan: Jahre mit Wert aus BWA 1240, auch 0, Lücke bleibt leer",
               g1.afa_plan, {2027: 15_000.0, 2028: 14_000.5, 2030: 0.0})
        pruefe("AfA-Plan: Basisjahr bleibt AfA lt. Buchhaltung", g1.abschreibung, 16_000.0)

        # Neukauf-Kostenstellen: alle Blätter hinter „KSt 9999“; die Marke selbst bleibt
        # eine Kostenstelle
        nk_pfad = tmp / "neukauf.xlsx"
        erstelle_vorlage(basisjahr=2026, neukauf=[
            ("KSt 31", "Neubau Nord", {1020: 50_000, 1250: 2_000, 1150: 1_000}),
            ("KSt 32", "Neubau Süd", {})], plan={"KSt 31": {1020: {2028: 52_000}}}).save(nk_pfad)
        alle, _ = lese_kostenstellen(nk_pfad, 2026)
        pruefe("Neukauf: Reihenfolge und Kennzeichen",
               [(lw.objekt_id, lw.neukauf) for lw in alle],
               [("KSt 1", False), ("KSt 2", False), ("KSt 9999", False), ("KSt 31", True),
                ("KSt 32", True)])
        nk = alle[3]
        pruefe("Neukauf: Miete je Jahr ab Basisjahr", nk.jahre["miete"],
               {2026: 50_000.0, 2028: 52_000.0})
        pruefe("Neukauf: weitere Ausgaben 1100–1220, 1260", nk.jahre["ausgaben"],
               {2026: 1_000.0})
        pruefe("Neukauf: Erhaltung", nk.jahre["erhaltung"], {2026: 2_000.0})
        pruefe("Neukauf: leeres Blatt ohne Jahreswerte", alle[4].jahre["miete"], {})
        pruefe("Neukauf: kein Bestandsobjekt",
               [o.objekt_id for o in zusammenfuehren([], alle)], ["KSt 1", "KSt 2", "KSt 9999"])
        pruefe("Vorlage: AfA lt. BWA KSt 2", v2.abschreibung, 30_000.0)
        try:
            (p1, _), _ = lese_kostenstellen(vorlage, 2027)
            pruefe("Vorlage: Planspalte 'Plan 2027' ist leer, Werte 0", p1.miete, 0.0)
        except EinleseFehler as e:
            pruefe("Vorlage: Planspalte 'Plan 2027' lesbar", str(e), "kein Fehler")
        try:
            lese_kostenstellen(vorlage, 2031)
            pruefe("Vorlage: Planjahr 2031 als Zahl erkannt", "kein Fehler", "kein Fehler")
        except EinleseFehler as e:
            pruefe("Vorlage: Planjahr 2031 als Zahl erkannt", str(e), "kein Fehler")

        stamm = testobjekt()  # OBJ-001
        stamm_kst1 = dataclasses.replace(stamm, objekt_id="KSt 1", name=None)
        objekte = zusammenfuehren([stamm_kst1, stamm], [lw1, lw2])
        pruefe("Zusammenführen: Anzahl Objekte", len(objekte), 3)
        pruefe("Zusammenführen: Miete überschrieben", objekte[0].miete, 120_000.0)
        pruefe("Zusammenführen: Stammdaten bleiben", objekte[0].ak_gebaeude, 800_000)
        pruefe("Zusammenführen: Objekt ohne Blatt unverändert", objekte[1].miete, 60_000)
        pruefe("Zusammenführen: neue Kostenstelle ohne AK", objekte[2].ak_gebaeude, None)

        # Aufschlüsselung der Abschreibungen unter der BWA (Stand Spalte "Jahr 2025")
        gruppen = {g.bezeichnung: g for g in v2.anlagen}
        pruefe("Aufschlüsselung: Gruppen ohne Buchwert und AfA entfallen", sorted(gruppen),
               ["Anbau im Bau", "Außenanlagen 7", "Wohnbau 7"])
        pruefe("Aufschlüsselung: Buchwert Stand", gruppen["Wohnbau 7"].bw_stand, 876_000.0)
        pruefe("Aufschlüsselung: AfA über Beschriftungsanfang (Wohnbau zu Wohnbau 7)",
               gruppen["Wohnbau 7"].afa, 24_000.0)
        pruefe("Aufschlüsselung: Zwischensumme ohne Beschriftung übersprungen",
               gruppen["Außenanlagen 7"].afa, 6_000.0)
        pruefe("Aufschlüsselung: ObjektID und eindeutige Nr.",
               (gruppen["Wohnbau 7"].objekt_id, gruppen["Wohnbau 7"].nr),
               ("KSt 2", "KSt 2 Wohnbau 7"))
        pruefe("Aufschlüsselung fehlt: keine Gruppen", lw1.anlagen, [])
        (_, ohne_stand), _ = lese_kostenstellen(vorlage, 2026, stand=2019)
        pruefe("Aufschlüsselung: Stand ohne Spalte, keine Gruppen", ohne_stand.anlagen, [])

        # Anlagenverzeichnis im Format des DATEV-Exports
        inventar_pfad = tmp / "Inventar_2025.xlsx"
        erstelle_inventar_vorlage().save(inventar_pfad)
        pruefe("Stand aus dem Dateinamen", stand_aus_dateiname(inventar_pfad), 2025)
        pruefe("Stand fehlt im Dateinamen", stand_aus_dateiname(tmp / "Anlagen.xlsx"), None)
        anlagen, abgang = lese_inventar(inventar_pfad)
        pruefe("Inventar: Abgang entfällt", (len(anlagen), abgang),
               (len(INVENTAR_BEISPIELE) - 1, 1))
        nach_nr = {a.nr: a for a in anlagen}
        geb = nach_nr["300002"]
        pruefe("Inventar: Gebäude", (geb.art, geb.methode, geb.satz, geb.ahk, geb.bw_stand),
               ("Gebäude", "linear", 0.02, 1_200_000.0, 876_000.0))
        pruefe("Inventar: Datum aus Text", str(geb.datum), "2012-07-01")
        pruefe("Inventar: AfA im Stand-Jahr = N-AfA Ende − Beginn", geb.afa_stand, 24_000.0)
        pruefe("Inventar: Art und Methode je Konto bzw. AfA-Art",
               [(nach_nr[n].art, nach_nr[n].methode) for n in
                ("100002", "310002", "750002", "690003", "910001")],
               [("G+B", "keine"), ("Gebäude", "linear"), ("im Bau", "keine"),
                ("BGA", "degressiv"), ("Finanzanlage", "keine")])
        pruefe("Inventar: AfA p. a. bleibt leer (Formel in der Mappe)", geb.afa, None)
        zugeordnet = {a.nr: a.objekt_id for a in ordne_anlagen_zu(anlagen, ["KSt 1", "KSt 2"])}
        pruefe("KOST1 zu ObjektID", [zugeordnet[n] for n in ("300001", "300002", "300009",
                                                             "100009")],
               ["KSt 1", "KSt 2", "KSt 9", None])
        andere = {a.nr: a.objekt_id for a in ordne_anlagen_zu(anlagen, ["Haus 2", "Haus 12"])}
        pruefe("KOST1 zu ObjektID mit gleicher Endnummer", andere["300002"], "Haus 2")

        # ohne KOST1: Zuordnung über die Inventarbezeichnung, gekennzeichnet
        objekte = [Objekt("KSt 1", name="Musterstraße 1"), Objekt("KSt 2", name="Beispielweg 7"),
                   Objekt("KSt 3", name="KC 24+26"), Objekt("KSt 4", name="KC 30")]
        ohne = [Anlage(nr=str(i), bw_stand=1.0, bezeichnung=b) for i, b in enumerate([
            "Grund und Boden Musterstr. 1", "Außenanlage Beispeilweg", "Wohngebäude KC 24",
            "Garage KC", "Grund und Boden ohne Kostenstelle", "Parkplatz Kst. 2", None])]
        bez = [(a.objekt_id, a.zuordnung) for a in ordne_anlagen_zu(ohne, objekte)]
        pruefe("Bezeichnung: abgekürzte Straße", bez[0],
               ("KSt 1", "Bezeichnung: musterstr, 1"))
        pruefe("Bezeichnung: Tippfehler", bez[1], ("KSt 2", "Bezeichnung: beispeilweg"))
        pruefe("Bezeichnung: Kürzel und Hausnummer", bez[2], ("KSt 3", "Bezeichnung: kc, 24"))
        pruefe("Bezeichnung: mehrdeutig bleibt ohne Objekt", bez[3],
               (None, "mehrdeutig: KSt 3, KSt 4"))
        pruefe("Bezeichnung: nur Füllwörter, keine Zuordnung", bez[4], (None, None))
        pruefe("Bezeichnung: KSt in der Bezeichnung", bez[5], ("KSt 2", "Bezeichnung: KSt 2"))
        pruefe("Bezeichnung fehlt", bez[6], (None, None))
        pruefe("KOST1 gekennzeichnet", ordne_anlagen_zu(anlagen, objekte)[0].zuordnung, "KOST1")
        pruefe("KOST1 geht vor der Bezeichnung",
               ordne_anlagen_zu([Anlage(nr="x", bw_stand=1.0, kost1="1",
                                        bezeichnung="Beispielweg 7")], objekte)[0].objekt_id,
               "KSt 1")
        mit_inventar = ordne_anlagen_zu(anlagen, ["KSt 1", "KSt 2"])
        zusammen, abgleich = anlagen_zusammenfuehren(mit_inventar, [v1, v2])
        pruefe("Anlagenverzeichnis geht vor der Aufschlüsselung", len(zusammen),
               len(mit_inventar))
        pruefe("Abgleich Buchwert je Kostenstelle", abgleich,
               [("KSt 1", 496_000.0, 496_000.0, 16_000.0),
                ("KSt 2", 941_000.0, 941_000.0, 30_000.0)])
        nur_kst1 = [a for a in mit_inventar if a.objekt_id != "KSt 2"]
        zusammen, _ = anlagen_zusammenfuehren(nur_kst1, [v1, v2])
        pruefe("Kostenstelle ohne Inventar: Gruppen aus der Aufschlüsselung",
               len(zusammen), len(nur_kst1) + 3)

        kein_inventar = tmp / "kein.xlsx"
        Workbook().save(kein_inventar)
        try:
            lese_inventar(kein_inventar)
            pruefe("Inventar ohne Kopf meldet Fehler", "kein Fehler", "EinleseFehler")
        except EinleseFehler:
            pruefe("Inventar ohne Kopf meldet Fehler", "EinleseFehler", "EinleseFehler")
        doppelt_inv = tmp / "doppelt_inv.xlsx"
        erstelle_inventar_vorlage(INVENTAR_BEISPIELE + INVENTAR_BEISPIELE[:1]).save(doppelt_inv)
        try:
            lese_inventar(doppelt_inv)
            pruefe("Inventar-Nr. doppelt meldet Fehler", "kein Fehler", "EinleseFehler")
        except EinleseFehler:
            pruefe("Inventar-Nr. doppelt meldet Fehler", "EinleseFehler", "EinleseFehler")

    print(f"\n{fehler} Abweichung(en)" if fehler else "\nAlle Prüfungen bestanden.")
    return 1 if fehler else 0


if __name__ == "__main__":
    sys.exit(main())
