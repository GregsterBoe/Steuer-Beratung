"""Kompakte Bedienungsanleitung für Steuerberater als .docx.

Aufruf: python docs/bedienungsanleitung.py docs/Bedienungsanleitung_Steuerberatung.docx
"""
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

doc = Document()
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.2)
    s.top_margin = s.bottom_margin = Cm(2)
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(10.5)
st.paragraph_format.space_after = Pt(4)
for name, size in (("Title", 22), ("Heading 1", 14), ("Heading 2", 11.5)):
    doc.styles[name].font.name = "Calibri"
    doc.styles[name].font.size = Pt(size)
    doc.styles[name].font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)


def absatz(text):
    p = doc.add_paragraph()
    _fett(p, text)
    return p


def _fett(p, text):
    """**fett** im Text auswerten."""
    teile = text.split("**")
    for i, t in enumerate(teile):
        if t:
            p.add_run(t).bold = i % 2 == 1


def punkte(liste, stil="List Bullet"):
    for t in liste:
        p = doc.add_paragraph(style=stil)
        p.paragraph_format.space_after = Pt(2)
        _fett(p, t)


def schattieren(zelle, farbe):
    tc = zelle._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), farbe)
    tc.append(shd)


def tabelle(kopf, zeilen, breiten, farben=None):
    t = doc.add_table(rows=1, cols=len(kopf))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, k in enumerate(kopf):
        z = t.rows[0].cells[i]
        z.text = ""
        z.paragraphs[0].add_run(k).bold = True
        schattieren(z, "D9E2F3")
    for n, zeile in enumerate(zeilen):
        zellen = t.add_row().cells
        for i, w in enumerate(zeile):
            zellen[i].text = ""
            _fett(zellen[i].paragraphs[0], w)
        if farben and farben[n]:
            schattieren(zellen[0], farben[n])
    t.autofit = False
    for i, b in enumerate(breiten):
        t.columns[i].width = Cm(b)
    for zeile in t.rows:
        for i, b in enumerate(breiten):
            zeile.cells[i].width = Cm(b)
            for p in zeile.cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.keep_with_next = True   # Tabelle nicht trennen
                for r in p.runs:
                    r.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


# ---------------------------------------------------------------- Titel
doc.add_paragraph("Prognosemodell V+V", style="Title")
absatz("**Kurzanleitung für die Steuerberatung:** 20-Jahres-Prognose für Immobilien im "
       "Betriebsvermögen einer GmbH mit AfA, Verkauf, § 6b-Rücklage und Reinvestition. "
       "Die Mappe vergleicht Halten, § 6b-Kette und sofortiges Versteuern.")

# ---------------------------------------------------------------- 1
doc.add_heading("1. Ablauf in fünf Schritten", level=1)
punkte([
    "**Blatt Start öffnen:** Handlungsempfehlung, Belastbarkeit der Daten und Farblegende ansehen.",
    "**Kostenstellen prüfen:** Werte des Basisjahrs in den Kostenstellenblättern kontrollieren "
    "und dort korrigieren.",
    "**Annahmen prüfen:** im Blatt Parameter zuerst die Steuerwelt (GmbH oder GmbH & Co. KG; "
    "für die KG die Sätze im Abschnitt GmbH & Co. KG) und die Kernparameter (★, fett: Steuersatz, "
    "Steigerungen, Alternativrendite), dann die blauen Zellen im Blatt Objekte. Wo bessere "
    "Werte bekannt sind, eintippen.",
    "**Planung erfassen:** geplante Verkäufe im Blatt Verkäufe, Käufe im Blatt Neuobjekte.",
    "**Prüfen und lesen:** „Plausibilität prüfen“ drücken, Fehler beheben, dann Start, "
    "Übersicht und Vergleich lesen.",
], "List Number")

# ---------------------------------------------------------------- 2
doc.add_heading("2. Die wichtigsten Blätter", level=1)
tabelle(["Blatt", "Wofür"], [
    ["**Start**", "Empfehlung (Option mit dem höchsten Endvermögen nach latenter Steuer), "
                  "Vorsprung gegenüber Halten, tiefster Liquiditätsstand, Anleitung mit Links"],
    ["**Parameter**", "Alle Sätze und Annahmen zentral; Schaltflächen (Abschnitt 7)"],
    ["**Objekte**", "Ein Bestandsobjekt je Zeile: Miete, Kosten, AK, Kaufjahr, Verkehrswert, "
                    "Status; Kostenstellen ohne Immobilie mit Art „Unternehmen“"],
    ["**Verkäufe**", "ObjektID, Verkauf aktiv (nein = ausgeschaltet, Rücklage entfällt), Jahr, "
                     "Preis, Kosten, Anteil G+B lt. Kaufvertrag, § 6b ja/nein, "
                     "„reinvestieren = ja“ legt ein Neuobjekt aus Annahmen an"],
    ["**Neuobjekte**", "Kaufjahr, Kaufpreis, Anteil G+B (AK G+B und AK Gebäude daraus "
                       "vorbelegt, in Euro überschreibbar; mit AK G+B in Euro darf der Anteil leer "
                       "bleiben und wird berechnet), AfA-Satz und -Methode, Sonder-AfA § 7b "
                       "(QNG, 5 % über 4 Jahre vor der regulären AfA), bis zu drei "
                       "Quell-Rücklagen, „Kostenstelle Neukauf“, Finanzierung (Abschnitt 4)"],
    ["**Darlehen**", "Tilgungsplan je Kredit: Zins, Tilgung und Restschuld je Jahr"],
    ["**Anlagen**", "Anlagenverzeichnis: AK, Buchwert und AfA je Anlage und Jahr"],
    ["**Übersicht / Vergleich**", "Ergebnis je Jahr und Szenarien nebeneinander (Abschnitt 5)"],
    ["**Ergebnisbrücke**", "Wirkung jedes Verkaufs und Kaufs auf das Ergebnis vor Steuern, "
                           "von Halten zum Plan (Abschnitt 5)"],
    ["**Cashflow**", "Kapitalflussrechnung je Jahr für ein wählbares Szenario: operativ, "
                     "Investition, Finanzierung, Kapitaldienst, Diagramm (Abschnitt 5)"],
    ["**Prüfung**", "Alle Plausibilitätsprüfungen als Fehler, Warnung oder Hinweis"],
    ["**Alle Objekte, KSt-Blätter**", "Ist und Plan im BWA-Format je Kostenstelle und "
                                      "als Summe; Verkauf und Kauf als eigenes Blatt"],
    ["**Alle Objekte halten**", "Summe im BWA-Format als reine Halten-Ansicht: ohne Verkäufe "
                                "und Käufe (Abschnitt 5)"],
], [4.2, 12.4])

# ---------------------------------------------------------------- 3
doc.add_heading("3. Farben der Eingabezellen", level=1)
tabelle(["Farbe", "Bedeutung", "Was tun?"], [
    ["rot", "Pflichtangabe fehlt oder Annahme gelöscht", "Wert eintragen"],
    ["orange", "Annahme bei verkauftem Objekt (bestimmt Gewinn und Rücklage) oder Wert "
               "weicht vom Kostenstellenblatt ab", "Wert belegen bzw. abgleichen"],
    ["blau", "Annahme aus dem Parameterblatt", "bei besserer Kenntnis überschreiben"],
    ["grün", "aus Buchhaltung, Kostenstellenblatt oder Anlagen", "nichts; zählt nicht als Annahme"],
    ["gelb", "händische Eingabe", "—"],
    ["grau", "Formel", "nicht überschreiben"],
], [2.2, 8.4, 6.0], ["F4B6B6", "F8CBAD", "BDD7EE", "C6E0B4", "FFF2CC", "D9D9D9"])
absatz("Ein eingetippter Wert ersetzt die Annahme. Wird ein Feld geleert, stellt die Schaltfläche "
       "„Annahmen wiederherstellen“ die Annahme wieder her.")
absatz("**Die Mappe erklärt sich selbst:** Spaltenköpfe in Objekte, Verkäufe, Neuobjekte und "
       "Anlagen sind rot mit * (Pflichtfeld), blau (optional) oder grau (berechnet). Ein rotes "
       "Dreieck am Kopf heißt: Maus darüber zeigt die Erläuterung, bei berechneten Spalten die "
       "Herleitung in Worten; das gilt auch für Prognose, Rücklagen, Darlehen, Liquidität und "
       "Auswertung. Im Blatt Parameter zeigt die Maus über der Bezeichnung, was der Parameter "
       "beeinflusst; Abschnitte ohne Kernparameter sind eingeklappt (+ am linken Rand).")

# ---------------------------------------------------------------- 4
doc.add_heading("4. Kostenstellenblätter sind führend", level=1)
punkte([
    "Die Blätter je Kostenstelle und „Alle Objekte“ sind die Datenbasis. Werte des Basisjahrs "
    "(gelbe Spalte) werden **dort** geändert; sie fließen automatisch in Objekte, Prognose und "
    "„Alle Objekte“.",
    "Im Blatt Objekte verweisen Miete, Erhaltung, weitere Einnahmen und Ausgaben, die AfA "
    "lt. Buchhaltung und der Zinsaufwand (1310) auf das Kostenstellenblatt (grün).",
    "Wird ein solcher Wert im Blatt Objekte überschrieben, rechnet nur die Prognose damit. Die "
    "Zelle wird orange, das Prüfungsblatt warnt. Die Schaltfläche **„Objekte -> Kostenstellen“** "
    "überträgt den Wert ins Kostenstellenblatt und stellt die Verknüpfung wieder her.",
    "Die Monatsspalten des Basisjahrs sind eingeklappt; das „+“ über der Jahresspalte öffnet sie.",
    "**Steigerung je Objekt:** Die Spalten Mietsteigerung und Erhaltungssteigerung im Blatt "
    "Objekte sind mit dem Parameter vorbelegt (blau) und lassen sich je Objekt überschreiben; "
    "leer gilt wieder der Parameter. „Miete eingefroren ab Jahr“: ab diesem Jahr bleibt die "
    "Miete auf dem Stand des Vorjahrs (weitere Einnahmen steigen weiter).",
    "**Mietsteigerung der Neuobjekte:** eigener Parameter „Mietsteigerung Neuobjekte p. a.“, "
    "getrennt vom Bestand. Im Blatt Neuobjekte kann jedes Neuobjekt einen eigenen Satz haben "
    "(leer = Parameter). „Mietsteigerung erst nach Jahren“, z. B. 10 bei Neubau oder "
    "Staffelmiete über dem Mietspiegel: so viele Jahre nach dem Kauf bleibt die Miete auf dem "
    "Anfangsstand und steigt erst danach. Das gilt auch für die Fortschreibung im Blatt "
    "Neukauf-KSt.",
])

doc.add_heading("Neukauf-Kostenstellen (KSt 31–35)", level=2)
absatz("Alle Kostenstellen hinter „KSt 9999“ gelten als geplante Käufe. Sie erhalten ein eigenes "
       "Blatt (Basisjahr und Planjahre gelb). Wählt man die Kostenstelle im Blatt Neuobjekte unter "
       "„Kostenstelle Neukauf“ aus, übernimmt die Prognose ab dem Jahr nach dem Kauf Miete, "
       "Einnahmen, Erhaltung und Ausgaben aus diesem Blatt; Jahre ohne Wert werden mit der "
       "Steigerung fortgeschrieben. Das Blatt ist dann zugleich das Blatt des Neuobjekts: AfA und "
       "Kreditzinsen kommen grau aus dem Modell, das Neuobjekt zählt in „Alle Objekte“. Kaufpreis, "
       "Anteil G+B und AfA-Satz bleiben Eingaben im Blatt Neuobjekte.")

doc.add_heading("Rücklagen und Finanzierung eines Kaufs", level=2)
punkte([
    "**Quell-Rücklagen:** bis zu drei Rücklagen je Neuobjekt aus der Auswahlliste. Übertragen "
    "werden erst alle Gebäudegewinne auf das Gebäude, dann die G+B-Gewinne auf G+B, der Rest auf "
    "das Gebäude. Ob die Übertragung im Einzelfall gewollt ist, entscheidet der Berater.",
    "**Die Rücklage wirkt nur steuerlich** (Übertragung, AfA-Basis). Die Finanzierung legt man "
    "getrennt fest.",
    "**Finanzierungsbedarf** = Kaufpreis + Nebenkosten. Mit „Verkaufserlös zur Finanzierung = "
    "ja“ abzüglich des Nettoerlöses der Quell-Verkäufe; leer oder „nein“ = ganze AK.",
    "**Finanzierung Rest:** leer bzw. „Eigenmittel“ zahlt den Bedarf aus der Liquidität; "
    "„Kredit“ finanziert ihn. Kreditbetrag leer = ganzer Bedarf, sonst der eingetragene Betrag, "
    "der Rest kommt aus Eigenmitteln. Alternativ gibt die Kreditquote den Kredit in % von "
    "Kaufpreis + Nebenkosten vor. Betrag und Quote stehen fest: der Verkaufserlös deckt dann "
    "nur den Eigenkapitalteil, der Rest bleibt für spätere Käufe bzw. in der Liquidität.",
    "**Kreditangaben:** Zinssatz, Tilgungsart (Annuität mit anfänglicher Tilgung, linear oder "
    "endfällig), Tilgung p. a., optional die Laufzeit, nach der die Restschuld getilgt wird, "
    "und tilgungsfreie Jahre (anfangs nur Zins).",
    "**Zweitkredit:** Ein Teil des Kredits zu eigenen Konditionen, z. B. ein günstiger "
    "Förderkredit. „Zweitkredit Anteil“ (in % des Kredits) oder „Zweitkredit Betrag“ (geht "
    "vor) und der Zinssatz; Tilgungsart, Tilgung, Laufzeit und tilgungsfreie Jahre leer = wie "
    "der Hauptkredit. Der Kredit insgesamt bleibt gleich, im Blatt Darlehen stehen die "
    "Zweitkredite unter den Hauptkrediten.",
    "**Restschuld des Zweitkredits nach der Laufzeit:** leer bzw. „ablösen“ tilgt sie nach "
    "der Laufzeit auf einmal aus der Liquidität. „auf Hauptkredit“ finanziert sie über den "
    "Hauptkredit weiter: Zum Ende des letzten Laufzeitjahrs (nach der regulären Rate) geht sie "
    "auf den Hauptkredit über, dessen Rate ab dem Folgejahr um übernommener Betrag × (Zinssatz "
    "+ Tilgung) steigt, bei linearer Tilgung um Betrag × Tilgung. Braucht eine Laufzeit; das "
    "Blatt Darlehen zeigt Übernahmejahr, Betrag und neue Rate.",
    "Zinsen mindern das Ergebnis und die Steuer, die Tilgung nur die Liquidität, die Restschuld "
    "das Gesamtvermögen. Ist der Kredit getilgt, entfallen Zins und Tilgung (Blatt Darlehen).",
])

doc.add_heading("Darlehen der Bestandsobjekte", level=2)
punkte([
    "**Grob, ohne weitere Daten:** Der Zinsaufwand des Basisjahrs (BWA 1310) läuft bis zum "
    "Verkaufsjahr weiter und sinkt je Jahr um den Satz vom Parameterblatt (Standard 3 %). "
    "Tilgung und Ablösung beim Verkauf fehlen; die Prüfung weist darauf hin.",
    "**Genau:** Im Blatt Objekte die **Restschuld zum 31.12.** des Basisjahrs eintragen "
    "(Saldenliste oder Bankauszug), möglichst auch Zinssatz und Jahresrate. Dann rechnet die "
    "Prognose Zins und Tilgung je Jahr und löst die Restschuld im Verkaufsjahr aus dem Erlös "
    "ab. Fehlen Zinssatz oder Rate, stehen Annahmen (blau) in den Zellen.",
    "Eine Vorfälligkeitsentschädigung gehört in die Verkaufskosten.",
])

# ---------------------------------------------------------------- 5
doc.add_heading("5. Ergebnis lesen", level=1)
punkte([
    "**Start:** empfohlene Option mit dem höchsten Endvermögen nach latenter Steuer, Vorsprung "
    "gegenüber Halten und tiefster Liquiditätsstand.",
    "**Übersicht:** Wert der Immobilien und Gesamtvermögen je Jahr, als Tabelle und Diagramm, "
    "mit und ohne die geplanten Verkäufe.",
    "**Vergleich:** Kennzahlen der Szenarien am Ende des Prognosezeitraums nebeneinander.",
    "**Ergebnisbrücke:** je Jahr vom Ergebnis vor Steuern bei Halten über die Pakete und den "
    "Zinseffekt der Liquidität zum Ergebnis vor Steuern im Plan. Ein Paket ist ein Verkauf "
    "(Gewinn und Rücklage minus das entfallende laufende Ergebnis) mit den Käufen aus seiner "
    "Rücklage oder seinem Erlös (Miete − Kosten − AfA − Kreditzinsen), mit Zwischensumme; "
    "Käufe ohne Quelle stehen getrennt, darunter alle Käufe einzeln mit ihrem Paket. Die Kontrollzeile muss 0 sein. Kostenstellen der Art „Unternehmen“ "
    "(Blatt Objekte, keine Immobilie) stehen getrennt im Ergebnis bei Halten.",
    "**Cashflow:** Kapitalflussrechnung je Jahr; das Szenario wählen Sie in Zelle B3. Der "
    "operative Cashflow rechnet indirekt vom Jahresüberschuss: plus AfA, minus "
    "Veräußerungsgewinne, plus Einstellung in die § 6b-Rücklage, minus Auflösung und "
    "Gewinnzuschlag. Darunter Investition (Verkaufserlöse, Käufe) und Finanzierung "
    "(Kreditauszahlung, planmäßige Tilgung, Ablösung bei Verkauf), die Liquidität am Jahresende "
    "und die Kapitaldienstdeckung (Einnahmen − Ausgaben je Euro Zins und Tilgung, unter 1 rot). "
    "Das Diagramm zeigt die drei Bereiche als Säulen und die Liquidität als Linie.",
    "**Varianten:** mit „Variante festhalten“ gesicherte Stände, etwa verschiedene Verkaufsjahre.",
    "**Alle Objekte halten:** die Summen-BWA, als würden alle Objekte gehalten und nichts "
    "gekauft (Baseline).",
    "**Vergleich mit Halten:** in „Alle Objekte“ D3 auf „ja“ stellen. Rechts neben den "
    "Planjahren steht dann je Position und Jahr Plan − Halten, grün mehr, rot weniger.",
])
tabelle(["Szenario", "Inhalt"], [
    ["A Plan", "§ 6b-Kette wie erfasst"],
    ["B", "Gewinn sofort versteuert, kein Neukauf aus der Rücklage, Geld angelegt"],
    ["C", "sofort versteuert, Neuobjekte trotzdem gekauft"],
    ["Baseline", "alles halten"],
], [3.0, 13.6])

# ---------------------------------------------------------------- 6
doc.add_heading("6. Prüfung", level=1)
punkte([
    "**Fehler:** Status ungleich OK in Objekten, Verkäufen oder Neuobjekten, Fristverstoß. "
    "Vor jeder Auswertung beheben.",
    "**Warnung:** Vorbesitzzeit für § 6b zu kurz, Rücklage nicht voll übertragen, Drei-Objekt-"
    "Grenze, Abweichung vom Kostenstellenblatt, unbekannte Neukauf-Kostenstelle, Verkauf "
    "ohne Restschuld trotz Zinsaufwand.",
    "**Hinweis:** Frist nach Prognoseende, negative Liquidität, fehlender Verkehrswert, Anlagen "
    "ohne Objekt.",
])

# ---------------------------------------------------------------- 7
doc.add_heading("7. Schaltflächen (Blatt Parameter)", level=1)
tabelle(["Schaltfläche", "Wirkung"], [
    ["Plausibilität prüfen", "zeigt alle Fehler, Warnungen und Hinweise"],
    ["Neu berechnen", "rechnet die Mappe vollständig neu"],
    ["Variante festhalten", "speichert die Kennzahlen des Vergleichs im Blatt Varianten"],
    ["Objekt anlegen / duplizieren / entfernen", "pflegt die Zeilen im Blatt Objekte"],
    ["Annahmen wiederherstellen", "setzt geleerte Felder wieder auf die Annahme"],
    ["Objekte -> Kostenstellen", "überträgt überschriebene Werte ins Kostenstellenblatt"],
    ["Leere Prognoseblöcke aus-/einblenden", "blendet ungenutzte Zeilen der Prognose aus"],
], [6.2, 10.4])

# ---------------------------------------------------------------- 8
doc.add_heading("8. Praxistipps", level=1)
punkte([
    "Bei verkauften Objekten AK, Kaufjahr und Anteil G+B belegen; orange Zellen bestimmen den "
    "Gewinn direkt.",
    "Graue Formelzellen nicht überschreiben; Änderungen immer in gelben oder blauen Zellen.",
    "Bei Objekten mit Darlehen die Restschuld eintragen, vor allem vor einem Verkauf: sonst "
    "fehlt die Ablösung, Liquidität und Kapitalanlage sind zu hoch.",
    "**Alle Steuersätze, Fristen und Annahmen vor dem Echteinsatz durch den zuständigen "
    "Berufsträger prüfen.**",
])

doc.add_heading("9. Neue Version der Mappe: Eingaben übernehmen", level=1)
absatz("Eine neue Version wird mit denselben Dateien erzeugt und bekommt die alte Mappe "
       "mit: python -m prognosemodell --kostenstellen Kostenstellen.xlsx --inventar "
       "Inventar.xlsx --uebernehmen Alte_Mappe.xlsx --ausgabe Neue_Mappe.xlsx. Die alte "
       "Mappe bleibt unverändert, der Bericht nennt je Blatt die übernommenen Werte.")
punkte([
    "**Übernommen werden eingetippte Werte** aus Parameter, Objekte, Anlagen (Zuordnung zum "
    "Objekt, AfA von Hand), AfA-Plan, Neukauf-KSt, Verkäufe, Neuobjekte, BWA-Zuordnung, "
    "Schalter „Vergleich mit Halten“ und Varianten. Blaue Annahmen (Formeln) rechnet die neue Mappe selbst.",
    "**Spalten nach Überschrift:** neue Spalten einer neuen Version behalten ihre "
    "Vorbelegung; Spalten, die es nicht mehr gibt, meldet der Bericht.",
    "**Kostenstellen-Datei geht vor:** Werte, die beim Einlesen entstanden sind, kommen aus "
    "der neuen Kostenstellen-Datei; nur abweichend eingetippte Werte (orange) bleiben. "
    "Änderungen in den Kostenstellenblättern der Mappe werden nicht übernommen, dafür die "
    "Kostenstellen-Datei pflegen.",
    "Verkäufe und Neuobjekte behalten ihre Zeilennummer (Rücklage und Reinvestition hängen "
    "daran); Objekte, Anlagen und Neukauf-Kostenstellen werden über ihre ID zugeordnet.",
])

absatz("**Übersichtskostenstelle:** Gehören mehrere Kostenstellen zu einem Gebäude, fasst "
       "--uebersicht \"Musterstraße 5=KSt 10,KSt 11\" sie in einem eigenen BWA-Blatt "
       "zusammen (mehrfach möglich; ObjektID oder Blattname). Das Blatt summiert nur die "
       "Kostenstellenblätter, rechnet nicht mit und erscheint im Mandanten-Export unter BWA "
       "je Objekt.")

doc.add_heading("10. Mandanten-Export", level=1)
absatz("Für den Mandanten gibt es eine schlanke Fassung der Mappe mit bis zu neun Blättern, nur "
       "mit Werten: Mappe in Excel speichern und schließen, dann python -m "
       "prognosemodell.export Prognose.xlsx. Daneben entsteht Prognose_Mandant.xlsx; mit "
       "--ausgabe ein anderer Name. Eine frisch erzeugte, nie geöffnete Mappe rechnet "
       "LibreOffice vorher durch.")
tabelle(["Blatt", "Inhalt"], [
    ["Überblick", "Empfehlung, Endvermögen je Option, Datenlage, Verkehrswert und "
                  "Gesamtvermögen je Jahr mit Diagrammen, Gesamtvermögen aller Optionen "
                  "je Jahr in einem Diagramm"],
    ["Jahresrechnung", "je Jahr vom Ergebnis über Steuer und Geldfluss zum Vermögen, Plan "
                       "und Halten, Differenz; jede Spalte rechnet von oben nach unten"],
    ["Ergebnisbrücke", "vom Ergebnis bei Halten je Verkauf und Kauf zum Ergebnis im Plan "
                       "(entfällt ohne Verkauf und Kauf)"],
    ["Cashflow", "Kapitalflussrechnung je Jahr für das im Blatt Cashflow gewählte Szenario: "
                 "operativ, Investition, Finanzierung, Liquidität, Kennzahlen, Diagramm"],
    ["BWA gesamt", "Summen-BWA aller Objekte, Ist-Jahre und Planjahre, Herleitung der "
                   "Sonderposten; bei „Vergleich mit Halten = ja“ auch die Differenz zu Halten"],
    ["BWA halten", "dieselbe Summen-BWA als Status quo: alle Objekte behalten, ohne Verkäufe "
                   "und Neuobjekte (entfällt ohne Verkauf und Kauf)"],
    ["BWA je Objekt", "alle Kostenstellenblätter untereinander, mit Inhaltsverzeichnis; "
                      "dazu die Übersichtskostenstellen"],
    ["Verkauf und Kauf", "Verkäufe, Neuobjekte, Ergebnis und Herleitung je Vorgang "
                         "(entfällt ohne Verkauf und Kauf)"],
    ["Annahmen", "zentrale Annahmen, Daten je Objekt mit Herkunft (blau Annahme, grün "
                 "Buchhaltung), auffällige Prüfungen"],
], [3.6, 13.0])
punkte([
    "Zeilen ohne Wert, Monatsspalten, Hilfszeilen und ausgeschaltete Vorgänge entfallen; "
    "Ist-Vorjahre erscheinen nur, wenn sie eingelesen sind.",
    "Jedes Blatt ist für den Druck eingerichtet: Querformat, eine Seite breit, Blattname und "
    "Seitenzahl in der Fußzeile.",
])

doc.save(sys.argv[1])
