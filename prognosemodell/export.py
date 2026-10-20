"""Mandanten-Export: die gerechnete Mappe kompakt, mit nachvollziehbaren Formeln (Projektplan
Abschnitte 30 und 36).

Aufruf: python -m prognosemodell.export MAPPE [--ausgabe PFAD]

Die Mappe muss durchgerechnet sein, also in Excel gespeichert. Ist sie es nicht (frisch
erzeugt und nie geöffnet), rechnet LibreOffice sie durch, sofern es installiert ist.

Statt aller Eingabe-, Rechen- und Kontrollblätter enthält der Export bis zu neun Blätter:
- Überblick: Empfehlung, Endvermögen je Option, Datenlage, Verkehrswert und Vermögen je Jahr
- Jahresrechnung: je Jahr Ergebnis und Steuer, Geldfluss und Vermögen, Plan und Halten
- Ergebnisbrücke: vom Ergebnis bei Halten je Verkauf und Kauf zum Ergebnis im Plan
  (entfällt ohne Verkauf und Kauf)
- Cashflow: Kapitalflussrechnung je Jahr für das in der Mappe gewählte Szenario, mit Diagramm
- BWA gesamt: Summen-BWA aller Objekte; mit „Vergleich mit Halten = ja“ auch die Differenz
  zu Halten
- BWA halten: Summen-BWA als Status quo, alle Objekte behalten (entfällt, wenn sie der BWA
  gesamt gleicht)
- BWA je Objekt: alle Kostenstellenblätter untereinander, dazu die Übersichtskostenstellen
- Verkauf und Kauf: Verkäufe, Neuobjekte und die Sicht je Vorgang (entfällt ohne beides)
- Annahmen: zentrale Annahmen, Daten je Objekt mit Herkunft, auffällige Prüfungen
Zeilen ohne Wert entfallen, ebenso Monatsspalten, Hilfszeilen und ausgeschaltete Vorgänge.
Summen und Herleitungen aus Zeilen desselben Blatts sind Formeln, Werte aus anderen Blättern
stehen als Zahl.
"""

import argparse
import datetime
import re
import shutil
import subprocess
import sys
import tempfile
from copy import copy
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formula.tokenizer import Token, Tokenizer
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

from .bwa import (HALTEBLATT, SONDERBEREICH, SUMMENBLATT, SUMMENZEILEN, UEBERSICHT_KST,
                  UEBERSICHT_ZELLE)
from .formeln import deckt_anlagen, hspalte, ispalte
from .mappe import (FILL_ANNAHME, FILL_EINGELESEN, FILL_KOPF, FONT_KOPF, OPTIONEN, _diagramm,
                    diagramm_cashflow, diagramm_szenarien, zentrale_annahmen)
from .modelle import (DARLEHEN_ANNAHMEN, FMT_EURO, FMT_JAHR, FMT_PROZENT, FMT_ZAHL,
                      KST_VERKNUEPFT, OBJEKT_EINGELESEN, OBJEKT_FELDER, PARAMETER, PRUEFUNGEN,
                      QUELLEN, STATUS_AUS, SZ_A, SZ_B, SZ_BASELINE, SZ_C)
from .vorlagen import ERSTE_ZEILE, SPALTE_JAHR, SPALTE_PLAN, SPALTE_VORJAHRE

UEBERBLICK = "Überblick"
JAHRESRECHNUNG = "Jahresrechnung"
BRUECKE = "Ergebnisbrücke"
CASHFLOW = "Cashflow"
BWA_GESAMT = "BWA gesamt"
BWA_HALTEN = "BWA halten"
BWA_OBJEKTE = "BWA je Objekt"
VERKAUF_KAUF = "Verkauf und Kauf"
ANNAHMEN = "Annahmen"
INHALT = [
    (JAHRESRECHNUNG, "je Jahr Ergebnis und Steuer, Geldfluss und Vermögen; Plan und Halten"),
    (BRUECKE, "Wirkung jedes Verkaufs und Kaufs auf das Ergebnis vor Steuern, von Halten "
              "zum Plan"),
    (CASHFLOW, "Kapitalflussrechnung je Jahr: operativ, Investition, Finanzierung, Liquidität, "
               "Kapitaldienst, mit Diagramm"),
    (BWA_GESAMT, "Summen-BWA über alle Objekte: Ist-Jahre aus der Buchhaltung, dann die "
                 "Planjahre"),
    (BWA_HALTEN, "Summen-BWA als Status quo: alle Objekte behalten, ohne Verkäufe und "
                 "Neuobjekte"),
    (BWA_OBJEKTE, "BWA je Kostenstelle untereinander, mit Inhaltsverzeichnis"),
    (VERKAUF_KAUF, "geplante Verkäufe und Neuobjekte, Ergebnis und Herleitung je Vorgang"),
    (ANNAHMEN, "zentrale Annahmen, Daten je Objekt mit Herkunft, Plausibilitätsprüfung"),
]

FMT_BETRAG = '#,##0;[Red]-#,##0;"–"'
FMT_BWA = '#,##0;[Red]-#,##0;""'
FMT_DIFF = '[Color10]+#,##0;[Red]-#,##0;"–"'   # grün mehr, rot weniger als bei Halten
FONT_TITEL = Font(bold=True, size=14)
FONT_ABSCHNITT = Font(bold=True, size=12)
FONT_HINWEIS = Font(italic=True, color="595959")
FONT_FETT = Font(bold=True)
FONT_LINK = Font(color="0563C1", underline="single")
FILL_ABSCHNITT = PatternFill("solid", fgColor="D9E1F2")
RAHMEN_SUMME = Border(top=Side(style="thin"))
SCHWELLE = 0.005   # kleiner gilt als 0
BREITE_TEXT = 46
BREITE_ZAHL = 12


# --- Werte der Quellmappe ---

def _bereich(wb, name: str) -> list:
    """Werte eines benannten Bereichs, zeilen- oder spaltenweise als Liste."""
    blatt, ref = next(iter(wb.defined_names[name].destinations))
    teile = ref.replace("$", "").split(":")
    zellen = wb[blatt][teile[0]:teile[-1]]
    return [c.value for zeile in zellen for c in zeile]


def _wert(wb, name: str):
    return _bereich(wb, name)[0]


def _zahl(wert) -> float:
    return wert if isinstance(wert, (int, float)) and not isinstance(wert, bool) else 0


def _ungleich_null(werte) -> bool:
    return any(abs(_zahl(w)) > SCHWELLE for w in werte)


def _ist_zahl(wert) -> bool:
    return isinstance(wert, (int, float)) and not isinstance(wert, bool)


# --- Formeln im Export ---
#
# Summen und Herleitungen, deren Teile im selben Blatt stehen, sind Formeln, damit der Mandant
# die Zusammensetzung sieht; Werte aus anderen Blättern stehen als Zahl. Eine Formel steht nur,
# wenn sie den Wert der Mappe ergibt, sonst bleibt der Wert.

BEZUG = re.compile(r"^(\$?)([A-Z]{1,3})(\$?)(\d+)$")
GENAU = 0.01   # Formel gegen den Wert der Mappe


class Formeln:
    """Formeln der Mappe je Blatt als {(Zeile, Spalte): Formel}; ohne Mappe leer."""

    def __init__(self, wb):
        self.wb, self._blaetter = wb, {}

    def __call__(self, blatt: str) -> dict:
        if self.wb is None or blatt not in self.wb.sheetnames:
            return {}
        if blatt not in self._blaetter:
            self._blaetter[blatt] = {
                (z, s): v
                for z, werte in enumerate(self.wb[blatt].iter_rows(min_row=1, min_col=1,
                                                                   values_only=True), start=1)
                for s, v in enumerate(werte, start=1)
                if isinstance(v, str) and v.startswith("=")}
        return self._blaetter[blatt]


def _formel(ws, zeile: int, spalte: int, formel: str):
    """Formel in die Zelle; nur so gesetzte Zellen bleiben im Export Formeln."""
    c = ws.cell(row=zeile, column=spalte, value=formel)
    ws.parent.formelzellen.add((ws.title, c.coordinate))
    return c


def _literal(wert) -> str:
    if isinstance(wert, str):
        return '"' + wert.replace('"', '""') + '"'
    return f"({wert!r})" if wert < 0 else repr(wert)


def uebersetze(formel, zeilen: dict, spalten: dict, wert, linear: bool = False,
               texte: tuple = ()):
    """Formel einer Zelle der Mappe für den Export, oder None.

    Bezüge auf Zellen desselben Blatts zeigen auf deren Zeile (zeilen) und Spalte (spalten) im
    Export; ein Bezug auf eine Zeile, die im Export fehlt, wird ihr Wert (wert(Zeile, Spalte)),
    in einer Summe entfällt er, wenn er 0 ist; ein Bereich schrumpft auf die Zeilen im Export,
    wenn die übrigen 0 oder leer sind. Bezüge auf die Spalten texte (Beschriftung) werden der
    Text. None bei Namen, Bezügen auf andere Blätter oder auf leere Zellen außerhalb des Exports
    und ohne einen Bezug im Export. linear: nur Grundrechenarten und SUM (Differenzen zweier
    BWA)."""
    if not isinstance(formel, str) or not formel.startswith("="):
        return None
    tok = Tokenizer(formel)
    im_export, nullen = False, set()
    for t in tok.items:
        if linear and t.type == Token.FUNC and t.subtype == Token.OPEN and t.value != "SUM(":
            return None
        if t.type != Token.OPERAND or t.subtype != Token.RANGE:
            continue
        teile = [BEZUG.match(x) for x in t.value.split(":")]
        if len(teile) > 2 or not all(teile):
            return None
        # Zelle wie ein Bereich aus einer Zelle
        (s1, sp1, z1, r1), (s2, sp2, z2, r2) = [m.groups() for m in teile] * (3 - len(teile))
        c1, c2 = column_index_from_string(sp1), column_index_from_string(sp2)
        r1, r2 = int(r1), int(r2)
        if len(teile) == 1 and c1 in texte:
            t.value = _literal(str(wert(r1, c1) or ""))
            continue
        if any(c not in spalten for c in range(c1, c2 + 1)):
            return None
        if len(teile) == 1:
            if r1 in zeilen:
                t.value = f"{s1}{get_column_letter(spalten[c1])}{z1}{zeilen[r1]}"
                im_export = True
                continue
            w = wert(r1, c1)
            if not _ist_zahl(w) and not isinstance(w, str):
                return None
            if _ist_zahl(w) and abs(w) <= SCHWELLE:
                w = 0
                nullen.add(id(t))
            t.value = _literal(w)
            continue
        drin = [r for r in range(r1, r2 + 1) if r in zeilen]
        if not drin or any(abs(_zahl(wert(r, c))) > SCHWELLE for r in range(r1, r2 + 1)
                           if r not in zeilen for c in range(c1, c2 + 1)):
            return None
        t.value = f"{s1}{get_column_letter(spalten[c1])}{z1}{zeilen[drin[0]]}"
        if (c1, drin[0]) != (c2, drin[-1]):
            t.value += f":{s2}{get_column_letter(spalten[c2])}{z2}{zeilen[drin[-1]]}"
        im_export = True
    if not im_export:
        return None
    tok.items = _ohne_nullen(tok.items, nullen)
    return _vereinfacht(tok.render())


def _ohne_nullen(items: list, nullen: set) -> list:
    """Summanden 0 weglassen: a+0-0 → a, 0+a → a; nicht neben Punktrechnung oder Vergleich."""
    def infix(t, zeichen="+-"):
        return t is not None and t.type == Token.OP_IN and t.value in zeichen

    def strich_oder_ende(t):   # rechts des Summanden endet die Summe oder sie geht weiter
        return t is None or not (t.type == Token.OP_IN and t.value not in "+-")

    items = list(items)
    weiter = True
    while weiter:
        weiter = False
        for k, t in enumerate(items):
            if id(t) not in nullen:
                continue
            vor = items[k - 1] if k else None
            nach = items[k + 1] if k + 1 < len(items) else None
            if infix(vor) and strich_oder_ende(nach) and \
                    not (k >= 2 and items[k - 2].type == Token.OP_IN):
                del items[k - 1:k + 1]
            elif (vor is None or vor.subtype == Token.OPEN or vor.type == Token.SEP) \
                    and infix(nach, "+"):
                del items[k:k + 2]
            else:
                continue
            weiter = True
            break
    return items


def _vereinfacht(formel: str) -> str:
    """Bedingungen, die im Export immer erfüllt sind, entfallen: =IF(1=1,X,"") → =X (Vorgang
    rechnet mit), =IF("Text"="","",X) → =X (Zeile mit Beschriftung)."""
    for muster in (r'=IF\(1=1,(.*),""\)', r'=IF\("(?:[^"]|"")+"="","",(.*)\)'):
        treffer = re.fullmatch(muster, formel)
        if not treffer:
            continue
        innen, tiefe = treffer.group(1), 0
        for zeichen in innen:
            tiefe += {"(": 1, ")": -1}.get(zeichen, 0)
            if tiefe < 0 or (tiefe == 0 and zeichen == ","):
                return formel
        return "=" + innen if tiefe == 0 else formel
    return formel


def _zelle(ws, zeile: int, spalte: int, wert, formel: str = None):
    """Formel, sonst der Wert."""
    if formel:
        return _formel(ws, zeile, spalte, formel)
    return ws.cell(row=zeile, column=spalte, value=wert)


def _passt(werte, soll) -> bool:
    return len(werte) == len(soll) and all(abs(_zahl(a) - _zahl(b)) <= GENAU
                                           for a, b in zip(werte, soll))


# --- Gestaltung ---

def _titel(ws, titel: str, hinweis: str) -> int:
    ws["A1"] = titel
    ws["A1"].font = FONT_TITEL
    ws["A2"] = hinweis
    ws["A2"].font = FONT_HINWEIS
    return 4


def _abschnitt(ws, zeile: int, text: str, bis: int = 1) -> None:
    for spalte in range(1, bis + 1):
        ws.cell(row=zeile, column=spalte).fill = FILL_ABSCHNITT
    ws.cell(row=zeile, column=1, value=text).font = FONT_ABSCHNITT


def _kopfzeile(ws, zeile: int, texte: list, start: int = 1) -> None:
    for spalte, text in enumerate(texte, start=start):
        c = ws.cell(row=zeile, column=spalte, value=text)
        c.font, c.fill = FONT_KOPF, FILL_KOPF
        c.alignment = Alignment(horizontal="center" if spalte > start else "left",
                                vertical="center", wrap_text=True)


def _link(zelle, ziel: str) -> None:
    zelle.hyperlink = Hyperlink(ref=zelle.coordinate, location=ziel, display=str(zelle.value))
    zelle.font = FONT_LINK


def _druck(ws, titelzeilen: str = None) -> None:
    """Querformat, eine Seite breit, Blattname und Seitenzahl in der Fußzeile."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.left.text = "&A"
    ws.oddFooter.right.text = "Seite &P von &N"
    if titelzeilen:
        ws.print_title_rows = titelzeilen


# --- Blatt Überblick ---

def _blatt_ueberblick(ex, wb, quelle: str, blaetter: list) -> None:
    ws = ex.active
    ws.title = UEBERBLICK
    basis = _wert(wb, "par_Basisjahr")
    jahre = _bereich(wb, "liq_Jahr")
    zeile = _titel(ws, "Prognose Immobilien – Überblick",
                   f"Basisjahr {basis} (Ist), Prognose {jahre[0]}–{jahre[-1]}. Erstellt am "
                   f"{datetime.date.today():%d.%m.%Y}" + (f" aus {quelle}" if quelle else "")
                   + ". Alle Werte beruhen auf den Annahmen im Blatt Annahmen; die Steuer ist "
                   "mit dem Grenzsteuersatz gerechnet.")

    _abschnitt(ws, zeile, "Ergebnis", 4)
    zeile += 1
    for name, font in (("start_Empfehlung", Font(bold=True, size=12, color="305496")),
                       ("start_Belastbarkeit", None)):
        c = ws.cell(row=zeile, column=1, value=_wert(wb, name))
        if font:
            c.font = font
        zeile += 1
    zeile += 1
    _kopfzeile(ws, zeile, [f"Option (Endvermögen Ende {jahre[-1]})",
                           "Endvermögen nach latenter Steuer", "Differenz zu Halten", "Rang"])
    zeile += 1
    werte = [_zahl(_wert(wb, name)) for _, name in OPTIONEN]
    erste, letzte = zeile, zeile + len(OPTIONEN) - 1   # erste Option: Halten
    for (text, _), wert in zip(OPTIONEN, werte):
        rang = sum(w > wert + SCHWELLE for w in werte) + 1
        ws.cell(row=zeile, column=1, value=text).font = FONT_FETT if rang == 1 else Font()
        for spalte, inhalt, formel, fmt in (
                (2, wert, None, FMT_BETRAG),
                (3, wert - werte[0], f"=B{zeile}-B${erste}", FMT_DIFF),
                (4, rang, f"=RANK(B{zeile},B${erste}:B${letzte})", "0")):
            _zelle(ws, zeile, spalte, inhalt, formel).number_format = fmt
        zeile += 1
    zeile += 1
    kum = _bereich(wb, "liq_Kum")
    tiefst = min(range(len(kum)), key=lambda i: _zahl(kum[i]))
    for text, wert, fmt in [
        ("Wert der § 6b-Kette (Plan gegen sofort versteuern und kaufen)",
         _wert(wb, "vg_DiffC"), FMT_DIFF),
        ("Steuer gesamt im Plan über alle Jahre", sum(map(_zahl, _bereich(wb, "liq_Steuer"))),
         FMT_BETRAG),
        (f"tiefster Liquiditätsstand im Plan (im Jahr {jahre[tiefst]}; negativ = "
         "Finanzierungsbedarf)", kum[tiefst], FMT_BETRAG),
    ]:
        ws.cell(row=zeile, column=1, value=text)
        ws.cell(row=zeile, column=2, value=wert).number_format = fmt
        zeile += 1

    zeile += 1
    _abschnitt(ws, zeile, "Datenlage", 4)
    zeile += 1
    for text, name in (("Objekte im Modell", "ueb_Objekte"),
                       ("davon mit Annahmen (Blatt Annahmen, blau)", "ueb_MitAnnahmen"),
                       ("geplante Verkäufe", "ueb_Verkaeufe"),
                       ("Neuobjekte", "ueb_Neuobjekte"),
                       ("Plausibilitätsprüfung", "pr_Gesamt")):
        ws.cell(row=zeile, column=1, value=text)
        ws.cell(row=zeile, column=2, value=_wert(wb, name)).alignment = Alignment(
            horizontal="right")
        zeile += 1

    zeile += 1
    _abschnitt(ws, zeile, "Inhalt", 4)
    zeile += 1
    for blatt, text in INHALT:
        if blatt in blaetter:
            _link(ws.cell(row=zeile, column=1, value=blatt), f"'{blatt}'!A1")
            ws.cell(row=zeile, column=2, value=text)
            zeile += 1

    zeile += 1
    _abschnitt(ws, zeile, "Entwicklung je Jahr (Ende des Jahres)", 7)
    zeile += 1
    ws.cell(row=zeile, column=2, value="Immobilien (Verkehrswert)").font = FONT_FETT
    ws.cell(row=zeile, column=5, value="Gesamtvermögen nach latenter Steuer").font = FONT_FETT
    kopf = zeile + 1
    _kopfzeile(ws, kopf, ["Jahr", "Halten", "Plan", "Differenz Plan − Halten",
                          "Halten", "Plan", "Differenz Plan − Halten"])
    spalten = [_bereich(wb, n) for n in ("ueb_Jahr", "ueb_Baseline", "ueb_Plan",
                                         "ueb_Differenz", "ueb_VermBaseline", "ueb_VermPlan",
                                         "ueb_VermDifferenz")]
    erste = kopf + 1
    for i, werte_jahr in enumerate(zip(*spalten)):
        z = erste + i
        # Differenz Plan − Halten aus den beiden Spalten links daneben
        differenz = {4: (f"=C{z}-B{z}", 2), 7: (f"=F{z}-E{z}", 5)}
        for spalte, wert in enumerate(werte_jahr, start=1):
            formel = None
            if spalte in differenz:
                formel, links = differenz[spalte]
                if not _passt([_zahl(werte_jahr[links]) - _zahl(werte_jahr[links - 1])],
                              [wert]):
                    formel = None
            c = _zelle(ws, z, spalte, wert, formel)
            c.number_format = "0" if spalte == 1 else (
                FMT_DIFF if spalte in (4, 7) else FMT_BETRAG)
            if spalte == 1:
                c.alignment = Alignment(horizontal="left")
    letzte = erste + len(spalten[0]) - 1
    # Diagramme rechts neben Ergebnis und Datenlage, die nur die Spalten A bis D nutzen
    for i, (titel, spalte) in enumerate((("Gesamtvermögen nach latenter Steuer", 5),
                                         ("Verkehrswert am Jahresende", 2))):
        chart = _diagramm(ws, titel, spalte, kopf, erste, letzte)
        chart.height, chart.width = 7, 17
        ws.add_chart(chart, f"F{4 + i * 14}")

    # alle Optionen nebeneinander, nicht nur Halten gegen Plan; Reihenfolge wie OPTIONEN
    zeile = letzte + 2
    _abschnitt(ws, zeile, "Gesamtvermögen nach latenter Steuer je Option (Ende des Jahres)", 7)
    kopf = zeile + 1
    szenarien = [SZ_BASELINE, SZ_A, SZ_C, SZ_B]
    _kopfzeile(ws, kopf, ["Jahr", "Halten", "Plan (§ 6b-Rücklage)", "sofort versteuern, kaufen",
                          "sofort versteuern, anlegen"])
    ws.row_dimensions[kopf].height = 30
    reihen = [_bereich(wb, "vgj_Jahr")] + [_bereich(wb, f"vgj_{sz.key}") for sz in szenarien]
    erste = kopf + 1
    for i, werte_jahr in enumerate(zip(*reihen)):
        for spalte, wert in enumerate(werte_jahr, start=1):
            c = ws.cell(row=erste + i, column=spalte, value=wert)
            c.number_format = "0" if spalte == 1 else FMT_BETRAG
            if spalte == 1:
                c.alignment = Alignment(horizontal="left")
    letzte = erste + len(reihen[0]) - 1
    chart = diagramm_szenarien(ws, kopf, erste, letzte, szenarien,
                               "Gesamtvermögen nach latenter Steuer je Option")
    chart.height, chart.width = 9, 17
    ws.add_chart(chart, f"G{erste + 1}")

    ws.column_dimensions["A"].width = 62
    for spalte in "BCDEFG":
        ws.column_dimensions[spalte].width = 17
    _druck(ws)


# --- Blatt Jahresrechnung ---

@dataclass(frozen=True)
class Posten:
    text: str
    quelle: str = None      # "liq" (Blatt Liquidität) oder "aus" (Blatt Auswertung)
    feld: str = None        # Name ohne Präfix, z. B. "Einnahmen"
    vorzeichen: int = 1     # −1: in der Quelle positiv, hier abgezogen
    art: str = "fluss"      # fluss: Summe über die Jahre, summe: dazu fett, bestand: Stand


JAHRESRECHNUNG_POSTEN = [
    Posten("Ergebnis und Steuer"),
    Posten("Mieten und weitere Einnahmen", "liq", "Einnahmen"),
    Posten("./. Erhaltung und weitere Ausgaben", "liq", "Ausgaben", -1),
    Posten("./. Abschreibungen (AfA)", "liq", "AfA", -1),
    Posten("= laufendes Ergebnis", "liq", "Ergebnis", art="summe"),
    Posten("+ steuerwirksam aus Verkauf und § 6b-Rücklage", "liq", "Verkauf"),
    Posten("+ Zinsertrag Alternativanlage", "liq", "Zins"),
    Posten("./. Zinsen Darlehen", "liq", "KreditZins", -1),
    Posten("= Ergebnis vor Verlustvortrag", "liq", "ZvE", art="summe"),
    Posten("./. Verlustvortrag genutzt", "liq", "VortragGenutzt", -1),
    Posten("= Bemessungsgrundlage", "liq", "Bemessung", art="summe"),
    Posten("Steuer", "liq", "Steuer", art="summe"),
    Posten("Verlustvortrag am Jahresende", "liq", "Vortrag", art="bestand"),
    Posten("Geldfluss"),
    Posten("laufendes Ergebnis", "liq", "Ergebnis"),
    Posten("+ Abschreibungen (nicht zahlungswirksam)", "liq", "AfA"),
    Posten("+ Zinsertrag Alternativanlage", "liq", "Zins"),
    Posten("./. Zinsen Darlehen", "liq", "KreditZins", -1),
    Posten("./. Steuer", "liq", "Steuer", -1),
    Posten("+ Verkaufserlöse netto", "liq", "Verkaufserloes"),
    Posten("./. Kauf Neuobjekte inkl. Nebenkosten", "liq", "Kauf", -1),
    Posten("+ Kreditauszahlung", "liq", "Kredit"),
    Posten("./. Tilgung Darlehen (mit Ablösung bei Verkauf)", "liq", "Tilgung", -1),
    Posten("= freier Mittelzufluss", "liq", "Zufluss", art="summe"),
    Posten("Liquidität am Jahresende (Alternativanlage)", "liq", "Kum", art="bestand"),
    Posten("Vermögen am Jahresende"),
    Posten("Verkehrswert Immobilien", "aus", "Verkehrswert", art="bestand"),
    Posten("+ Liquidität", "aus", "Liquiditaet", art="bestand"),
    Posten("./. Restschuld Darlehen", "aus", "Restschuld", -1, art="bestand"),
    Posten("= Gesamtvermögen vor latenter Steuer", "aus", "Vermoegen", art="bestand"),
    Posten("./. latente Steuer auf stille Reserven und Rücklage", "aus", "LatenteSteuer", -1,
           art="bestand"),
    Posten("= Gesamtvermögen nach latenter Steuer", "aus", "VermoegenNetto", art="bestand"),
    Posten("nachrichtlich: Buchwert Immobilien", "aus", "Buchwert", art="bestand"),
    Posten("nachrichtlich: stille Reserven", "aus", "StilleReserven", art="bestand"),
    Posten("nachrichtlich: § 6b-Rücklage", "aus", "Ruecklage", art="bestand"),
]
DIFFERENZ_POSTEN = [
    Posten("laufendes Ergebnis", "liq", "Ergebnis"),
    Posten("Steuer", "liq", "Steuer"),
    Posten("freier Mittelzufluss", "liq", "Zufluss"),
    Posten("Liquidität am Jahresende", "liq", "Kum", art="bestand"),
    Posten("Verkehrswert Immobilien", "aus", "Verkehrswert", art="bestand"),
    Posten("Restschuld Darlehen", "aus", "Restschuld", art="bestand"),
    Posten("Gesamtvermögen nach latenter Steuer", "aus", "VermoegenNetto", art="bestand"),
]
# Präfixe der Namen je Sicht: Plan = Szenario A, Halten = Baseline
SICHTEN = {"Plan": {"liq": "liq", "aus": "aus"}, "Halten": {"liq": "lqb", "aus": "asb"}}


def _reihe(wb, sicht: str, p: Posten) -> list:
    return [p.vorzeichen * _zahl(w) for w in _bereich(wb, f"{SICHTEN[sicht][p.quelle]}_{p.feld}")]


# Posten außerhalb der Rechenkette ihres Abschnitts (nachrichtlich bzw. Stand)
NICHT_IN_SUMME = ("Steuer", "Verlustvortrag am Jahresende",
                  "Liquidität am Jahresende (Alternativanlage)")
# derselbe Wert unter einem anderen Namen weiter oben
GLEICHER_WERT = {("aus", "Liquiditaet"): ("liq", "Kum")}


class Herleitung:
    """Formeln einer Sicht der Jahresrechnung, {s} steht für die Spalte. Eine „=“-Zeile summiert
    die letzte „=“-Zeile des Abschnitts und die Zeilen dazwischen; ein Posten, der weiter oben
    schon steht, verweist auf ihn."""

    def __init__(self):
        self.zeilen = {}   # (Quelle, Feld) → (Zeile, Vorzeichen, Werte)
        self.teile = []    # (Zeile, Werte) seit der letzten „=“-Zeile

    def __call__(self, p: Posten, werte: list, zeile: int):
        if p.quelle is None:
            self.teile = []
            return None
        formel = None
        if p.text.startswith("="):
            summe = [sum(t[i] for _, t in self.teile) for i in range(len(werte))]
            if self.teile and _passt(summe, werte):
                formel = "=" + "+".join(f"{{s}}{z}" for z, _ in self.teile)
            self.teile = [(zeile, werte)]
        else:
            formel = self.verweis(p, werte)
            if p.text not in NICHT_IN_SUMME and not p.text.startswith("nachrichtlich"):
                self.teile.append((zeile, werte))
        self.zeilen.setdefault((p.quelle, p.feld), (zeile, p.vorzeichen, werte))
        return formel

    def verweis(self, p: Posten, werte: list, minus: "Herleitung" = None):
        """Formel auf denselben Posten weiter oben; mit minus: dieser minus der in minus."""
        schluessel = (p.quelle, p.feld)
        bezug = self.zeilen.get(schluessel) or self.zeilen.get(GLEICHER_WERT.get(schluessel))
        abzug = minus.zeilen.get(schluessel) if minus else None
        if not bezug or (minus and not abzug):
            return None
        zeile, vorzeichen, werte_bezug = bezug
        faktor = p.vorzeichen * vorzeichen   # ±1
        if minus:
            soll = [faktor * (a - b) for a, b in zip(werte_bezug, abzug[2])]
            formel = (f"={{s}}{zeile}-{{s}}{abzug[0]}" if faktor > 0
                      else f"={{s}}{abzug[0]}-{{s}}{zeile}")
        else:
            soll = [faktor * w for w in werte_bezug]
            formel = f"={'' if faktor > 0 else '-'}{{s}}{zeile}"
        return formel if _passt(werte, soll) else None


def _jahreszeilen(ws, zeile: int, posten: list, reihen, n: int, fmt: str = FMT_BETRAG,
                  herleitung=None) -> int:
    """Je Posten eine Zeile: Text, Jahreswerte, Summe (nur Flussgrößen). herleitung(p, werte,
    zeile) liefert eine Formel mit {s} für die Spalte oder None (dann die Werte)."""
    letzte = get_column_letter(n + 1)
    for p in posten:
        if p.quelle is None:
            ws.cell(row=zeile, column=1, value=p.text).font = FONT_FETT
            if herleitung:
                herleitung(p, [], zeile)
            zeile += 1
            continue
        werte = reihen(p)
        formel = herleitung(p, werte, zeile) if herleitung else None
        ws.cell(row=zeile, column=1, value=p.text)
        for i, wert in enumerate(werte, start=2):
            _zelle(ws, zeile, i, wert, formel and formel.format(s=get_column_letter(i))
                   ).number_format = fmt
        if p.art != "bestand":
            _formel(ws, zeile, n + 2, f"=SUM(B{zeile}:{letzte}{zeile})").number_format = fmt
        if p.text.startswith("="):
            for spalte in range(1, n + 3):
                c = ws.cell(row=zeile, column=spalte)
                c.font, c.border = FONT_FETT, RAHMEN_SUMME
        zeile += 1
    return zeile


def _blatt_jahresrechnung(ex, wb) -> None:
    ws = ex.create_sheet(JAHRESRECHNUNG)
    jahre = _bereich(wb, "liq_Jahr")
    n = len(jahre)
    zeile = _titel(ws, "Jahresrechnung: vom Ergebnis zum Vermögen",
                   "Plan = mit den geplanten Verkäufen und Neuobjekten (§ 6b wie erfasst), "
                   "Halten = alle Objekte behalten. Beträge in €, „./.“ wird abgezogen; jede "
                   "Spalte rechnet von oben nach unten. Summe = über alle Jahre.")
    _kopfzeile(ws, zeile, ["Position"] + jahre + ["Summe"])
    ws.freeze_panes = ws.cell(row=zeile + 1, column=2)
    kopf = zeile
    zeile += 1
    plan_gleich_halten = all(
        all(abs(a - b) <= SCHWELLE for a, b in zip(_reihe(wb, "Plan", p), _reihe(wb, "Halten", p)))
        for p in JAHRESRECHNUNG_POSTEN if p.quelle)
    herleitungen = {}
    for sicht, titel in (("Plan", "Plan: mit Verkäufen und Neuobjekten"),
                         ("Halten", "Halten: alle Objekte behalten")):
        if sicht == "Halten" and plan_gleich_halten:
            break
        zeile += 1
        _abschnitt(ws, zeile, titel, n + 2)
        herleitungen[sicht] = Herleitung()
        zeile = _jahreszeilen(ws, zeile + 1, JAHRESRECHNUNG_POSTEN,
                              lambda p, s=sicht: _reihe(wb, s, p), n,
                              herleitung=herleitungen[sicht])
    zeile += 1
    if plan_gleich_halten:
        ws.cell(row=zeile, column=1, value="Ohne Verkäufe und Neuobjekte gleicht der Plan dem "
                                           "Halten.").font = FONT_HINWEIS
    else:
        _abschnitt(ws, zeile, "Differenz Plan − Halten (+ Plan besser)", n + 2)
        zeile = _jahreszeilen(
            ws, zeile + 1, DIFFERENZ_POSTEN,
            lambda p: [a - b for a, b in zip(_reihe(wb, "Plan", p), _reihe(wb, "Halten", p))],
            n, FMT_DIFF,
            herleitung=lambda p, werte, _: herleitungen["Plan"].verweis(
                p, werte, minus=herleitungen["Halten"]))
    ws.column_dimensions["A"].width = BREITE_TEXT
    for spalte in range(2, n + 3):
        ws.column_dimensions[get_column_letter(spalte)].width = BREITE_ZAHL
    _druck(ws, f"{kopf}:{kopf}")


# --- Blatt Ergebnisbrücke ---

def _uebertragen(ws, eintraege: list, quelle, formeln: dict, spalte: int, n: int) -> None:
    """Zeilen eines Blatts mit Jahresspalten in den Export: eintraege (Zeile der Quelle, Zeile
    im Export, Werte, Summe ja/nein); Formeln der Mappe, soweit sie im Blatt bleiben, sonst die
    Werte. Summe = über die Jahre."""
    zeilen = {q: z for q, z, _, _ in eintraege}
    spalten = {spalte + i: 2 + i for i in range(n)}

    def wert(z, s):
        return quelle.cell(row=z, column=s).value

    letzte = get_column_letter(n + 1)
    for q, z, werte, summe in eintraege:
        for i, w in enumerate(werte):
            formel = uebersetze(formeln.get((q, spalte + i)), zeilen, spalten, wert,
                                texte=(1,)) if _ist_zahl(wert(q, spalte + i)) else None
            _zelle(ws, z, 2 + i, w, formel)
        if summe:
            _formel(ws, z, n + 2, f"=SUM(B{z}:{letzte}{z})")


def _blatt_bruecke(ex, wb, formeln: Formeln) -> bool:
    """Blatt Ergebnisbrücke ohne leere Pakete und Vorgänge, ohne Paketspalte und
    Kontrollzeile; ohne
    Verkauf und Kauf kein Blatt (False)."""
    if not any(_ungleich_null(_bereich(wb, f"brk_{name}"))
               for name in ("Veraeusserung", "Laufend", "Neuobjekte")):
        return False
    quelle = wb[BRUECKE]
    blatt, ref = next(iter(wb.defined_names["brk_Halten"].destinations))
    erste = wb[blatt][ref.replace("$", "").split(":")[0]]
    kopf, spalte = erste.row - 1, erste.column
    jahre = [c.value for c in quelle[kopf][spalte - 1:]
             if isinstance(c.value, (int, float))]
    n = len(jahre)
    kontrolle = next(iter(wb.defined_names["brk_Kontrolle"].destinations))[1]
    kontrollzeile = int(kontrolle.split("$")[2].split(":")[0])
    ws = ex.create_sheet(BRUECKE)
    zeile = _titel(ws, "Ergebnisbrücke: Wirkung der Entscheidungen",
                   "Vom Ergebnis vor Steuern bei Halten je Verkauf und Kauf zum Ergebnis vor "
                   "Steuern im Plan. Verkauf: Gewinn und § 6b-Rücklage minus das laufende "
                   "Ergebnis, das mit dem Objekt entfällt. Kauf: Miete − Kosten − AfA − "
                   "Kreditzinsen. Summe = über alle Jahre.")
    _kopfzeile(ws, zeile, ["Position"] + jahre + ["Summe"])
    ws.freeze_panes = ws.cell(row=zeile + 1, column=2)
    kopfzeile = zeile
    zeile += 1
    eintraege = []
    for q in range(kopf + 1, quelle.max_row + 1):
        text = quelle.cell(row=q, column=1).value
        if text in (None, "") or q == kontrollzeile:
            continue
        werte = [_zahl(quelle.cell(row=q, column=spalte + i).value) for i in range(n)]
        einzug = quelle.cell(row=q, column=1).alignment.indent
        if quelle.cell(row=q, column=2).value is None:   # Abschnitt
            zeile += 1
            _abschnitt(ws, zeile, text, n + 2)
            zeile += 1
            continue
        # leere Pakete und Vorgänge entfallen, feste Summenzeilen bleiben
        if not _ungleich_null(werte) and not text.startswith(
                ("Ergebnis", "Summe", "Zinseffekt", "Wirkung")):
            continue
        fett = bool(quelle.cell(row=q, column=1).font.bold)
        c = ws.cell(row=zeile, column=1, value=text)
        c.alignment = Alignment(indent=einzug)
        for i in range(2, n + 3):
            ws.cell(row=zeile, column=i).number_format = FMT_BETRAG
        if fett:
            for i in range(1, n + 3):
                c = ws.cell(row=zeile, column=i)
                c.font, c.border = FONT_FETT, RAHMEN_SUMME
        eintraege.append((q, zeile, werte, True))
        zeile += 1
    _uebertragen(ws, eintraege, quelle, formeln(BRUECKE), spalte, n)
    ws.column_dimensions["A"].width = BREITE_TEXT
    for i in range(2, n + 3):
        ws.column_dimensions[get_column_letter(i)].width = BREITE_ZAHL
    _druck(ws, f"{kopfzeile}:{kopfzeile}")
    return True


# --- Blatt Cashflow ---

def _blatt_cashflow(ex, wb, formeln: Formeln) -> None:
    """Blatt Cashflow für das dort gewählte Szenario: Jahre, dann die Summe; ohne
    Kontrollzeile, Zeilen ohne Wert entfallen bis auf die fetten Summen; Diagramm darunter."""
    quelle = wb[CASHFLOW]
    blatt, ref = next(iter(wb.defined_names["cf_Jahresueberschuss"].destinations))
    erste = wb[blatt][ref.replace("$", "").split(":")[0]]
    spalte = erste.column
    kopf = next(z for z in range(erste.row, 0, -1)
                if quelle.cell(row=z, column=1).value == "Position")
    n = column_index_from_string(ref.replace("$", "").split(":")[-1].rstrip("0123456789")) \
        - spalte + 1
    jahre = [quelle.cell(row=kopf, column=spalte + i).value for i in range(n)]

    def zeile_von(name):
        return int(next(iter(wb.defined_names[name].destinations))[1].split("$")[2].split(":")[0])

    kontrollzeile = zeile_von("cf_Kontrolle")
    diagramm = {key: zeile_von(f"cf_{name}") for key, name in (
        ("operativ", "Operativ"), ("investition", "Investition"),
        ("finanzierung", "Finanzierung"), ("ende", "Ende"))}
    ws = ex.create_sheet(CASHFLOW)
    zeile = _titel(ws, f"Cashflow: Kapitalflussrechnung je Jahr ({_wert(wb, 'cf_Szenario')})",
                   "Operativer Cashflow indirekt: Jahresüberschuss plus Aufwand ohne Auszahlung "
                   "(AfA, Einstellung in die § 6b-Rücklage) minus Ertrag ohne Einzahlung "
                   "(Veräußerungsgewinne, Auflösung, Gewinnzuschlag); Zinsen gehören dazu. "
                   "Zuflüsse positiv, Abflüsse negativ. Kapitaldienstdeckung = (Einnahmen − "
                   "Ausgaben) / (Zinsen + planmäßige Tilgung). Summe = über alle Jahre.")
    _kopfzeile(ws, zeile, ["Position"] + jahre + ["Summe"])
    ws.freeze_panes = ws.cell(row=zeile + 1, column=2)
    kopfzeile = zeile
    zeilen = {}
    eintraege = []
    zeile += 1
    for q in range(kopf + 1, quelle.max_row + 1):
        text = quelle.cell(row=q, column=1).value
        if text in (None, "") or q == kontrollzeile:
            continue
        roh = [quelle.cell(row=q, column=spalte + i).value for i in range(n)]
        if quelle.cell(row=q, column=spalte).value is None and not any(roh):   # Abschnitt
            zeile += 1
            _abschnitt(ws, zeile, text, n + 2)
            zeile += 1
            continue
        fett = bool(quelle.cell(row=q, column=1).font.bold)
        quote = quelle.cell(row=q, column=spalte).number_format == "0.00"
        werte = [w if isinstance(w, (int, float)) else None for w in roh] if quote \
            else [_zahl(w) for w in roh]
        if not fett and not quote and not _ungleich_null(werte) and q not in diagramm.values():
            continue
        if quote and not any(w is not None for w in werte):
            continue
        c = ws.cell(row=zeile, column=1, value=text)
        c.alignment = Alignment(indent=quelle.cell(row=q, column=1).alignment.indent)
        fmt = "0.00" if quote else FMT_BETRAG
        for i in range(2, n + 3):
            ws.cell(row=zeile, column=i).number_format = fmt
        # Summe nur bei Flussgrößen
        eintraege.append((q, zeile, werte, quelle.cell(row=q, column=2).value is not None))
        if fett:
            for i in range(1, n + 3):
                c = ws.cell(row=zeile, column=i)
                c.font, c.border = FONT_FETT, RAHMEN_SUMME
        for key, q_key in diagramm.items():
            if q == q_key:
                zeilen[key] = zeile
        zeile += 1
    _uebertragen(ws, eintraege, quelle, formeln(CASHFLOW), spalte, n)
    ws.column_dimensions["A"].width = 58
    for i in range(2, n + 3):
        ws.column_dimensions[get_column_letter(i)].width = BREITE_ZAHL
    ws.add_chart(diagramm_cashflow(ws, kopfzeile, zeilen, 2, n + 1), f"A{zeile + 1}")
    _druck(ws, f"{kopfzeile}:{kopfzeile}")


# --- BWA-Blätter ---

@dataclass
class BwaZeile:
    art: str          # bwa (Zeile mit BWA-Nr.), titel (Abschnitt), daten (Herleitung u. Ä.)
    nr: int
    text: str
    ziel: object      # Spalte D: BWA-Nr. des Postens bzw. Anlagennummer
    werte: list
    zeile: int = None  # Zeile im BWA-Blatt


def ist_bwa_blatt(ws) -> bool:
    return ws["B4"].value == "Nr." and ws.cell(row=ERSTE_ZEILE - 1, column=3).value \
        == "Jahr (für Formeln)"


def _bwa_spalten(wb) -> list:
    """(Spalte im BWA-Blatt, Kopf im Export): Ist-Vorjahre nur, wenn eingelesen."""
    basis = _wert(wb, "par_Basisjahr")
    summe = _bwa_lesen(wb[SUMMENBLATT], [(SPALTE_VORJAHRE, ""), (SPALTE_VORJAHRE + 1, "")])
    vorjahre = [(SPALTE_VORJAHRE + i, f"Ist {basis - 2 + i}") for i in range(2)
                if _ungleich_null(z.werte[i] for z in summe if z.art == "bwa")]
    plan = [(SPALTE_PLAN + i, jahr) for i, jahr in enumerate(_bereich(wb, "liq_Jahr"))]
    return vorjahre + [(SPALTE_JAHR, f"Ist {basis}")] + plan


def _bwa_lesen(ws, spalten: list) -> list:
    """Zeilen eines BWA-Blatts ab der ersten BWA-Zeile: Werte der gewählten Spalten."""
    zeilen = []
    for nummer, zeile in enumerate(ws.iter_rows(min_row=ERSTE_ZEILE, values_only=True),
                                   start=ERSTE_ZEILE):
        zeile = zeile + (None,) * (max(s for s, _ in spalten) - len(zeile))
        nr, text, ziel = zeile[1], zeile[2], zeile[3]
        if not isinstance(text, str) or not text.strip():
            continue
        werte = [zeile[s - 1] for s, _ in spalten]
        if isinstance(nr, int) and not text.endswith(":"):
            art = "bwa"
        elif isinstance(nr, int) or not any(isinstance(w, (int, float)) for w in werte):
            art = "titel"
        else:
            art = "daten"
        zeilen.append(BwaZeile(art, nr if isinstance(nr, int) else None, text, ziel, werte,
                               nummer))
    return zeilen


def _bwa_filtern(zeilen: list) -> list:
    """Zeilen ohne Wert weglassen; Summenzeilen bleiben, Abschnitte nur mit Inhalt."""
    ergebnis, abschnitt = [], []

    def abschliessen():
        if any(z.art != "titel" for z in abschnitt):
            ergebnis.extend(abschnitt)
        abschnitt.clear()

    for z in zeilen:
        if z.art == "bwa":
            abschliessen()
            if z.nr in SUMMENZEILEN or _ungleich_null(z.werte):
                ergebnis.append(z)
        elif z.art == "titel":
            abschliessen()
            abschnitt.append(z)
        elif _ungleich_null(z.werte):
            abschnitt.append(z)
    abschliessen()
    return ergebnis


def _bwa_block(ws, zeile: int, zeilen: list, spalten: list, fmt: str = FMT_BWA,
               formeln: dict = None, wert=None, linear: bool = False) -> int:
    """Kopf und Zeilen einer BWA ab zeile; liefert die nächste freie Zeile. formeln: die des
    BWA-Blatts, wert(Zeile, Spalte) dessen Werte (Summenzeilen als Formel, siehe uebersetze)."""
    _kopfzeile(ws, zeile, ["Bezeichnung", "Nr. / →"] + [k for _, k in spalten])
    zeile += 1
    im_export = {z.zeile: zeile + i for i, z in enumerate(zeilen)}
    nach_spalte = {s: i for i, (s, _) in enumerate(spalten, start=3)}
    for z in zeilen:
        ws.cell(row=zeile, column=1, value=z.text)
        ziel = z.nr if z.art == "bwa" else z.ziel
        if ziel not in (None, "", "→ Nr."):
            ws.cell(row=zeile, column=2, value=ziel).alignment = Alignment(horizontal="center")
        if z.art == "titel":
            ws.cell(row=zeile, column=1).font = FONT_FETT
        else:
            for (s, _), w in zip(spalten, z.werte):
                if _ist_zahl(w):
                    formel = uebersetze(formeln.get((z.zeile, s)), im_export, nach_spalte, wert,
                                        linear) if formeln else None
                    _zelle(ws, zeile, nach_spalte[s], w, formel).number_format = fmt
        if z.nr in SUMMENZEILEN:
            for spalte in range(1, len(spalten) + 3):
                c = ws.cell(row=zeile, column=spalte)
                c.font, c.border = FONT_FETT, RAHMEN_SUMME
        zeile += 1
    return zeile


def _bwa_breiten(ws, spalten: list) -> None:
    ws.column_dimensions["A"].width = BREITE_TEXT
    ws.column_dimensions["B"].width = 9
    for i in range(len(spalten)):
        ws.column_dimensions[get_column_letter(i + 3)].width = BREITE_ZAHL


def _bwa_werte(ws):
    """Werte eines BWA-Blatts je (Zeile, Spalte), für Bezüge auf ausgelassene Zeilen."""
    return lambda z, s: ws.cell(row=z, column=s).value


def _blatt_bwa_gesamt(ex, wb, spalten: list, formeln: Formeln) -> None:
    ws = ex.create_sheet(BWA_GESAMT)
    vergleich = str(_wert(wb, "bwa_VglHalten")).strip().lower() == "ja"
    zeile = _titel(ws, "BWA gesamt: alle Objekte",
                   "Ist-Spalten aus der Buchhaltung, Planspalten aus der Prognose (mit Verkäufen "
                   "und Neuobjekten), Steuer und Zinsertrag der Alternativanlage nur hier. "
                   "Unter der BWA die Herleitung der Sonderposten mit der BWA-Nr., in die sie "
                   "eingehen. Zeilen ohne Wert sind ausgelassen.")
    plan = _bwa_lesen(wb[SUMMENBLATT], spalten)
    _abschnitt(ws, zeile, "Plan: mit Verkäufen und Neuobjekten", len(spalten) + 2)
    zeile = _bwa_block(ws, zeile + 1, _bwa_filtern(plan), spalten, formeln=formeln(SUMMENBLATT),
                       wert=_bwa_werte(wb[SUMMENBLATT])) + 1
    # Halten selbst steht im Blatt BWA halten
    if vergleich:
        _abschnitt(ws, zeile, f"Differenz Plan − Halten (+ Plan höher; Halten im Blatt "
                              f"{BWA_HALTEN})", len(spalten) + 2)
        halten = {z.nr: z.werte for z in _bwa_lesen(wb[HALTEBLATT], spalten) if z.art == "bwa"}
        plan_ab = _plan_ab(spalten)
        diff = [BwaZeile("bwa", z.nr, z.text, None,
                         [None] * plan_ab + [_zahl(a) - _zahl(b) for a, b in
                                             zip(z.werte[plan_ab:], halten[z.nr][plan_ab:])],
                         z.zeile)
                for z in plan if z.art == "bwa" and z.nr in halten]
        # Summenzeilen der Differenz wie im Plan: Differenzen summieren sich wie die Beträge
        werte = {(z.zeile, s): w for z in diff for (s, _), w in zip(spalten, z.werte)}
        _bwa_block(ws, zeile + 1, _bwa_filtern(diff), spalten, FMT_DIFF,
                   formeln=formeln(SUMMENBLATT), wert=lambda z, s: werte.get((z, s)),
                   linear=True)
    _bwa_breiten(ws, spalten)
    ws.freeze_panes = "C5"
    _druck(ws)


def _plan_ab(spalten: list) -> int:
    """Index der ersten Planspalte in spalten."""
    return next(i for i, (s, _) in enumerate(spalten) if s >= SPALTE_PLAN)


def _blatt_bwa_halten(ex, wb, spalten: list, formeln: Formeln) -> bool:
    """Summen-BWA bei Halten (Status quo). False und kein Blatt, wenn sie in den Planjahren
    der BWA gesamt gleicht, also ohne Verkauf und Kauf."""
    halten = _bwa_lesen(wb[HALTEBLATT], spalten)
    plan = {z.nr: z.werte for z in _bwa_lesen(wb[SUMMENBLATT], spalten) if z.art == "bwa"}
    ab = _plan_ab(spalten)
    if all(abs(_zahl(a) - _zahl(b)) <= SCHWELLE
           for z in halten if z.art == "bwa"
           for a, b in zip(z.werte[ab:], plan.get(z.nr, [None] * len(spalten))[ab:])):
        return False
    ws = ex.create_sheet(BWA_HALTEN)
    zeile = _titel(ws, "BWA halten: Status quo, alle Objekte behalten",
                   "Wie die BWA gesamt, aber ohne Verkäufe und Neuobjekte: alle Bestandsobjekte "
                   "bleiben über das ganze Raster, ihre Darlehen laufen weiter. Ist-Spalten wie "
                   "in der BWA gesamt, Planspalten aus der Baseline (Steuer mit Verlustvortrag). "
                   "Zeilen ohne Wert sind ausgelassen.")
    _abschnitt(ws, zeile, "Halten: alle Objekte behalten", len(spalten) + 2)
    _bwa_block(ws, zeile + 1, _bwa_filtern(halten), spalten, formeln=formeln(HALTEBLATT),
               wert=_bwa_werte(wb[HALTEBLATT]))
    _bwa_breiten(ws, spalten)
    ws.freeze_panes = "C5"
    _druck(ws)
    return True


def _blatt_bwa_objekte(ex, wb, spalten: list, formeln: Formeln) -> None:
    ws = ex.create_sheet(BWA_OBJEKTE)
    blaetter = [b for b in wb.worksheets
                if b.title not in (SUMMENBLATT, HALTEBLATT) and ist_bwa_blatt(b)]
    zeile = _titel(ws, "BWA je Objekt",
                   "Je Kostenstelle die BWA: Ist-Spalten aus der Buchhaltung, Planspalten aus "
                   "der Prognose; nach einem Verkauf leer, Neuobjekte ab dem Jahr nach dem Kauf. "
                   "Steuer und Zinsertrag der Alternativanlage stehen nur im Blatt BWA gesamt. "
                   "Eine Übersicht fasst die Kostenstellen eines Gebäudes zusammen (Summe der "
                   "genannten Kostenstellen, zusätzlich zu deren BWA).")
    _abschnitt(ws, zeile, "Inhalt", 2)
    inhalt = zeile + 1
    zeile = inhalt + len(blaetter) + 1
    for i, quelle in enumerate(blaetter):
        kst, name = quelle["B2"].value, quelle["C2"].value
        titel = f"{kst} – {name}" if name else str(kst)
        if kst == UEBERSICHT_KST and quelle[UEBERSICHT_ZELLE].value:
            titel += f" ({quelle[UEBERSICHT_ZELLE].value})"
        _link(ws.cell(row=inhalt + i, column=1, value=titel), f"'{BWA_OBJEKTE}'!A{zeile}")
        _abschnitt(ws, zeile, titel, len(spalten) + 2)
        zeile = _bwa_block(ws, zeile + 1, _bwa_filtern(_bwa_lesen(quelle, spalten)),
                           spalten, formeln=formeln(quelle.title),
                           wert=_bwa_werte(quelle)) + 1
    if not blaetter:
        ws.cell(row=inhalt, column=1, value="keine Kostenstellen").font = FONT_HINWEIS
    _bwa_breiten(ws, spalten)
    ws.freeze_panes = "C4"
    _druck(ws)


# --- Blatt Verkauf und Kauf ---

@dataclass(frozen=True)
class Angabe:
    """Zeile der Tabellen Verkäufe und Neuobjekte, je Vorgang eine Spalte."""
    text: str
    name: str = None        # benannter Bereich oder berechnete Reihe (#…), Wert je Vorgang
    fmt: str = None
    key: str = None         # Schlüssel für Formeln anderer Zeilen
    formel: str = None      # aus Zeilen derselben Spalte, z. B. "{preis}+{kosten}"
    vorzeichen: int = 1     # −1: abgezogen („./.“)
    leer: object = None     # Anzeige einer leeren Eingabe, z. B. "ja" oder 0
    ebene: int = 0          # Gliederung: Zeilen der Ebene 1 und 2 lassen sich einklappen
    kopf: bool = False      # Überschrift einer Gruppe, ohne Werte
    nur_mit: str = None     # nur, wenn die Zeile dieses Schlüssels einen Betrag ≠ 0 hat


FMT_QM = '#,##0.00 "€"'
VERKAUF_ANGABEN = [
    Angabe("ObjektID", "vk_ID"),
    Angabe("Bezeichnung", "vk_Bez"),
    Angabe("weitere Kostenstellen im Verkauf", "vk_Weitere"),
    Angabe("Verkauf aktiv", "vk_Aktiv", leer="ja"),
    Angabe("Status", "vk_Status"),
    Angabe("Verkaufsjahr (zum Jahresende)", "vk_Jahr", FMT_JAHR),
    Angabe("Vorbesitzzeit Jahre", "vk_Vorbesitz", FMT_ZAHL),
    Angabe("Verkaufspreis", "vk_Preis", FMT_BETRAG, "preis"),
    Angabe("./. Verkaufskosten", "vk_Kosten", FMT_BETRAG, "kosten", vorzeichen=-1, leer=0),
    Angabe("= Nettoerlös", "vk_Nettoerloes", FMT_BETRAG, "netto", "{preis}+{kosten}"),
    Angabe("./. Buchwert Gebäude Ende Verkaufsjahr", "vk_BuchwertGeb", FMT_BETRAG, "bw_geb",
           vorzeichen=-1),
    Angabe("./. Buchwert Grund und Boden (AK)", "vk_AKGuB", FMT_BETRAG, "bw_gub",
           vorzeichen=-1),
    Angabe("= Veräußerungsgewinn", "vk_Gewinn", FMT_BETRAG, "gewinn",
           "{netto}+{bw_geb}+{bw_gub}"),
    Angabe("Aufteilung auf Gebäude und Grund und Boden", kopf=True),
    Angabe("Anteil G+B am Preis", "vk_AnteilGuB", FMT_PROZENT, "anteil", ebene=1),
    Angabe("Erlösanteil G+B (Nettoerlös × Anteil)", "vk_ErloesGuB", FMT_BETRAG, "erl_gub",
           "{netto}*{anteil}", ebene=1),
    Angabe("Erlösanteil Gebäude (Nettoerlös − G+B)", "vk_ErloesGeb", FMT_BETRAG, "erl_geb",
           "{netto}-{erl_gub}", ebene=1),
    Angabe("Gewinn Gebäude", "vk_GewinnGeb", FMT_BETRAG, formel="{erl_geb}+{bw_geb}", ebene=1),
    Angabe("Gewinn G+B", "vk_GewinnGuB", FMT_BETRAG, formel="{erl_gub}+{bw_gub}", ebene=1),
    Angabe("§ 6b nutzen", "vk_6b", leer="nein"),
    Angabe("§ 6b Neubau begonnen", "vk_6bNeubau", leer="nein"),
    Angabe("§ 6b-Rücklage gebildet", "rl_Betrag", FMT_BETRAG),
    Angabe("Fristjahr der Rücklage", "rl_Fristjahr", FMT_JAHR),
    Angabe("reinvestieren (Neuobjekt nach den Annahmen)", "vk_Reinvest", leer="nein"),
    Angabe("Miete und Darlehen des Objekts", kopf=True),
    Angabe("Miete im Verkaufsjahr", "#miete", FMT_BETRAG, ebene=1),
    Angabe("Zinssatz Darlehen", "#zinssatz", FMT_PROZENT, ebene=1),
    Angabe("Ablösung Darlehen bei Verkauf (Restschuld)", "#abloesung", FMT_BETRAG, ebene=1),
]
# Tilgungsplan im Blatt Darlehen: #dl_… Hauptkredit, #dl2_… Zweitkredit
KREDIT_ANGABEN = [("Zinssatz", "Zinssatz", FMT_PROZENT), ("Tilgungsart", "Art", None),
                  ("Tilgung p. a.", "Tilgungssatz", FMT_PROZENT),
                  ("Laufzeit Jahre", "Laufzeit", FMT_ZAHL),
                  ("tilgungsfreie Jahre", "Tilgungsfrei", FMT_ZAHL),
                  ("Rate p. a. (Annuität) bzw. Tilgung p. a. (linear)", "Rate", FMT_BETRAG)]
NEU_ANGABEN = [
    Angabe("NeuID", "ne_ID"),
    Angabe("Name", "ne_Name"),
    Angabe("Kauf aktiv", "ne_Aktiv", leer="ja"),
    Angabe("Status", "ne_Status"),
    Angabe("Kaufjahr (zum Jahresende)", "ne_Kaufjahr", FMT_JAHR),
    Angabe("§ 6b-Rücklage aus (RücklageID)", "#quellen"),
    Angabe("Kaufpreis", "ne_Kaufpreis", FMT_BETRAG, "kp"),
    Angabe("+ Kaufnebenkosten", "ne_Nebenkosten", FMT_BETRAG, "nk", leer=0),
    Angabe("= Kaufpreis + Nebenkosten", "ne_AKGesamt", FMT_BETRAG, "ak", "{kp}+{nk}"),
    Angabe("davon Grund und Boden", "ne_AKGuBNeu", FMT_BETRAG, "ak_gub"),
    Angabe("davon Gebäude", "ne_AKGebNeu", FMT_BETRAG, "ak_geb", "{ak}-{ak_gub}"),
    Angabe("Anteil G+B", "ne_AnteilGuBVerw", FMT_PROZENT, formel="{ak_gub}/{ak}"),
    Angabe("Übertrag § 6b-Rücklage", "ne_UeGesamt", FMT_BETRAG, "ue", "{ue1}+{ue2}+{ue3}"),
    Angabe("ü1 Gebäudegewinn auf Gebäude", "ne_Ue1", FMT_BETRAG, "ue1", ebene=1),
    Angabe("ü2 G+B-Gewinn auf G+B", "ne_Ue2", FMT_BETRAG, "ue2", ebene=1),
    Angabe("ü3 G+B-Gewinn auf Gebäude", "ne_Ue3", FMT_BETRAG, "ue3", ebene=1),
    Angabe("AfA-Basis Gebäude (Gebäude − ü1 − ü3)", "ne_AfABasis", FMT_BETRAG,
           formel="{ak_geb}-{ue1}-{ue3}"),
    Angabe("steuerliche AK G+B (G+B − ü2)", "ne_AKGuB", FMT_BETRAG, formel="{ak_gub}-{ue2}"),
    Angabe("AfA-Satz", "ne_AfASatz", FMT_PROZENT),
    Angabe("Gebäude und Miete", kopf=True),
    Angabe("AfA-Methode", "ne_AfAMethode", leer="linear", ebene=1),
    Angabe("Sonder-AfA § 7b (QNG)", "ne_QNG", leer="nein", ebene=1),
    Angabe("Fläche m²", "ne_Flaeche", "#,##0", "fl", ebene=1),
    Angabe("Mietpreis je m² mtl.", "ne_MieteQm", FMT_QM, "mqm", ebene=1),
    Angabe("Jahresmiete aus m² (Fläche × Mietpreis × 12)", "ne_MieteQmJahr", FMT_BETRAG,
           formel="{fl}*{mqm}*12", ebene=1),
    Angabe("Mietrendite auf Kaufpreis", "ne_Mietrendite", FMT_PROZENT, ebene=1),
    Angabe("Erhaltung auf Kaufpreis", "ne_ErhQuote", FMT_PROZENT, ebene=1),
    Angabe("Mietsteigerung p. a.", "ne_MietsteigV", FMT_PROZENT, ebene=1),
    Angabe("Mietsteigerung erst nach Jahren", "ne_MietWarte", FMT_ZAHL, ebene=1),
    Angabe("Kostenstelle Neukauf", "ne_KSt", ebene=1),
    Angabe("Miete im ersten vollen Jahr", "#miete", FMT_BETRAG, ebene=1),
    Angabe("Finanzierung", kopf=True),
    Angabe("Kaufpreis + Nebenkosten", "ne_AKGesamt", FMT_BETRAG, "ak_fin", "{ak}"),
    Angabe("./. Einsatz Verkaufserlös (Quell-Verkäufe)", "ne_Erloes", FMT_BETRAG, "erloes",
           vorzeichen=-1),
    Angabe("= Finanzierungsbedarf", "ne_Bedarf", FMT_BETRAG, "bedarf", "{ak_fin}+{erloes}"),
    Angabe("Verkaufserlös zur Finanzierung", "ne_ErloesFin", leer="nein"),
    Angabe("Finanzierung Rest", "ne_FinArt", leer="Eigenmittel"),
    Angabe("davon Eigenmittel", "ne_Eigen", FMT_BETRAG, formel="{bedarf}-{kredit}"),
    Angabe("davon Kredit", "ne_Kredit", FMT_BETRAG, "kredit"),
    Angabe("Kreditquote (Kredit / Kaufpreis + Nebenkosten)", fmt=FMT_PROZENT,
           formel="{kredit}/{ak}", ebene=1, nur_mit="kredit"),
    Angabe("Hauptkredit (Kredit − Zweitkredit)", "#dl_Betrag", FMT_BETRAG,
           formel="{kredit}-{zweit}", ebene=1, nur_mit="kredit"),
    *[Angabe(text, f"#dl_{name}", fmt, ebene=1, nur_mit="kredit")
      for text, name, fmt in KREDIT_ANGABEN],
    Angabe("Zweitkredit", "ne_Kredit2", FMT_BETRAG, "zweit", ebene=1, nur_mit="zweit"),
    Angabe("Anteil am Kredit", fmt=FMT_PROZENT, formel="{zweit}/{kredit}", ebene=2,
           nur_mit="zweit"),
    *[Angabe(text, f"#dl2_{name}", fmt, ebene=2, nur_mit="zweit")
      for text, name, fmt in KREDIT_ANGABEN],
    Angabe("Restschuld nach der Laufzeit (leer = ablösen)", "ne_Kredit2Uebernahme", ebene=2,
           nur_mit="zweit"),
]
SB_AUSLASSEN = ("rechnet mit (1 = ja)",)
GRAU = "E7E6E6"   # FILL_BERECHNET der Mappe


def _prognose(wb) -> list:
    """Zeilen des Blatts Prognose als (ID, Jahr, Neuobjekt, Miete, Ablösung)."""
    return list(zip(*(_bereich(wb, n) for n in ("prg_ID", "prg_Jahr", "prg_Neu", "prg_Miete",
                                                "prg_Abloesung"))))


def _reihen_verkauf(wb) -> dict:
    """Werte je Verkaufszeile: benannte Bereiche, dazu Miete, Zinssatz und Ablösung des
    Objekts bzw. Pakets im Verkaufsjahr."""
    reihen = {a.name: _bereich(wb, a.name) for a in VERKAUF_ANGABEN
              if a.name and not a.name.startswith("#")}
    prognose = _prognose(wb)
    objekte = dict(zip(_bereich(wb, "obj_ID"), zip(_bereich(wb, "obj_Zinssatz"),
                                                    _bereich(wb, "obj_Restschuld"))))
    miete, zinssatz, abloesung = [], [], []
    for umfang, jahr in zip(_bereich(wb, "vk_Umfang"), reihen["vk_Jahr"]):
        ids = [i for i in str(umfang or "").split(";") if i]
        zeilen = [z for z in prognose if z[0] in ids and z[1] == jahr and not _zahl(z[2])]
        miete.append(sum(_zahl(z[3]) for z in zeilen) if zeilen else None)
        abloesung.append(sum(_zahl(z[4]) for z in zeilen) if zeilen else None)
        # Zinssatz der Objekte mit Darlehen, nur wenn alle gleich
        saetze = {objekte[i][0] for i in ids if i in objekte and _zahl(objekte[i][1]) > 0}
        zinssatz.append(saetze.pop() if len(saetze) == 1 else None)
    return {**reihen, "#miete": miete, "#zinssatz": zinssatz, "#abloesung": abloesung}


def _reihen_neu(wb) -> dict:
    """Werte je Zeile des Blatts Neuobjekte, dazu RücklageIDs, Miete im ersten vollen Jahr und
    der Tilgungsplan von Haupt- und Zweitkredit (Blatt Darlehen)."""
    reihen = {a.name: _bereich(wb, a.name) for a in NEU_ANGABEN
              if a.name and not a.name.startswith("#")}
    n = len(reihen["ne_ID"])
    quellen = zip(*(_bereich(wb, name) for _, name, _ in QUELLEN))
    reihen["#quellen"] = ["; ".join(str(q) for q in qs if q not in (None, "")) or None
                          for qs in quellen]
    prognose = _prognose(wb)
    reihen["#miete"] = [
        next((z[3] for z in prognose if z[0] == id_ and _zahl(z[2]) and _ist_zahl(jahr)
              and z[1] == jahr + 1), None) if id_ not in (None, "") else None
        for id_, jahr in zip(reihen["ne_ID"], reihen["ne_Kaufjahr"])]
    for _, name, _ in KREDIT_ANGABEN + [(None, "Betrag", None)]:
        plan = _bereich(wb, f"dl_{name}")
        reihen[f"#dl_{name}"], reihen[f"#dl2_{name}"] = plan[:n], plan[n:2 * n]
    return reihen


def _auswerten(formel: str, werte: dict):
    """Wert einer Formel aus Angaben derselben Spalte (leer = 0); None, wenn nicht rechenbar."""
    try:
        return eval(formel.format(**{k: f"({_zahl(v)!r})" for k, v in werte.items()}),
                    {"__builtins__": {}})
    except ZeroDivisionError:
        return None


def _vorgaenge(ws, zeile: int, angaben: list, reihen: dict) -> tuple:
    """Tabelle je Vorgang eine Spalte, je Angabe eine Zeile; Herleitungen als Formel, Gruppen
    zum Einklappen, ausgeschaltete Vorgänge grau durchgestrichen. Liefert (Zeile, Anzahl)."""
    nummern = [i for i, id_ in enumerate(reihen[angaben[0].name]) if id_ not in (None, "")]
    if not nummern:
        return zeile, 0
    werte = {}
    for a in angaben:
        if a.name:
            reihe = reihen[a.name]
            werte[a] = [a.vorzeichen * reihe[i] if _ist_zahl(reihe[i])
                        else a.leer if reihe[i] in (None, "") else reihe[i] for i in nummern]
    nach_key = {a.key: werte.get(a) for a in angaben if a.key}
    for a in angaben:   # Zeilen nur aus Formeln
        if not a.name and a.formel:
            werte[a] = [_auswerten(a.formel, {k: v[j] for k, v in nach_key.items() if v})
                        for j in range(len(nummern))]

    def zeigen(a) -> bool:
        if a.nur_mit and not _ungleich_null(nach_key[a.nur_mit]):
            return False
        return any(w not in (None, "") for w in werte[a])

    sichtbar = []
    for k, a in enumerate(angaben):
        if a.kopf:   # Überschrift nur mit sichtbaren Zeilen bis zur nächsten
            gruppe = []
            for b in angaben[k + 1:]:
                if b.kopf:
                    break
                gruppe.append(b)
            if any(zeigen(b) for b in gruppe):
                sichtbar.append(a)
        elif zeigen(a):
            sichtbar.append(a)
    zeile_von = {a.key: zeile + k for k, a in enumerate(sichtbar) if a.key}
    aktiv = next(a for a in angaben if a.name and a.name.endswith("_Aktiv"))
    aus = [str(w).strip().lower() == "nein" for w in werte[aktiv]]
    ws.sheet_properties.outlinePr.summaryBelow = False
    for k, a in enumerate(sichtbar):
        c = ws.cell(row=zeile, column=1, value=a.text)
        c.alignment = Alignment(indent=a.ebene)
        if a.ebene:
            ws.row_dimensions[zeile].outline_level = a.ebene
        if a.kopf:
            c.font = FONT_FETT
        for j, w in enumerate(werte.get(a, []) if not a.kopf else []):
            spalte = get_column_letter(j + 2)
            formel = None
            if a.formel and all(key in zeile_von for key in re.findall(r"{(\w+)}", a.formel)):
                soll = _auswerten(a.formel, {key: v[j] for key, v in nach_key.items() if v})
                if soll is not None and (not a.name or (_ist_zahl(w) and abs(soll - w) <= GENAU)):
                    formel = "=" + a.formel.format(**{key: f"{spalte}{z}"
                                                      for key, z in zeile_von.items()})
            c = _zelle(ws, zeile, j + 2, w, formel)
            if a.fmt:
                c.number_format = a.fmt
            if not formel and not _ist_zahl(w):
                c.alignment = Alignment(horizontal="right")
        if k == 0:
            _kopfzeile(ws, zeile, [a.text] + [werte[a][j] for j in range(len(nummern))])
        for spalte in range(1, len(nummern) + 2) if k else ():
            summe = a.text.startswith("=")
            c = ws.cell(row=zeile, column=spalte)
            if spalte > 1 and aus[spalte - 2]:
                c.font = Font(bold=summe, color="808080", strike=True)
            elif summe:
                c.font = FONT_FETT
            if summe:
                c.border = RAHMEN_SUMME
        zeile += 1
    return zeile, len(nummern)


def _sonderbereich(ws, zeile: int, quelle, formeln: dict) -> int:
    """Blöcke je Vorgang und Detailsicht Kauf aus dem Blatt Verkauf und Kauf, nur mit Inhalt;
    Herleitungen innerhalb eines Blocks als Formel."""
    zeilen = list(quelle.iter_rows(min_row=1, max_col=5))
    bloecke, aktuell = [], None
    for z in zeilen:
        text = z[0].value
        if isinstance(text, str) and (text.startswith("Vorgang ")
                                      or text.startswith("Detailsicht Kauf")):
            aktuell = []
            bloecke.append(aktuell)
        elif isinstance(text, str) and text.startswith("Weitere Vorgänge"):
            aktuell = None
        if aktuell is not None:
            aktuell.append(z)

    def wert(z, s):
        return quelle.cell(row=z, column=s).value

    for block in bloecke:
        titel = block[0][0].value
        if titel.startswith("Vorgang "):
            mit = next((z[1].value for z in block if z[0].value == SB_AUSLASSEN[0]), 0)
            if not _zahl(mit):
                continue
            spalten = range(5)
        else:   # Detailsicht Kauf: nur Spalten der erfassten Neuobjekte
            ids = next((z for z in block if z[0].value == "NeuID (Kostenstelle)"), None)
            spalten = [0] + [i for i in range(1, 5) if ids and ids[i].value not in (None, "")]
            if len(spalten) == 1:
                continue
        block = [z for z in block if z[0].value not in SB_AUSLASSEN]
        im_export = {z[0].row: zeile + k for k, z in enumerate(block)}
        nach_spalte = {i + 1: ziel for ziel, i in enumerate(spalten, start=1)}
        for z in block:
            for ziel, i in enumerate(spalten, start=1):
                q = z[i]
                formel = uebersetze(formeln.get((q.row, q.column)), im_export, nach_spalte,
                                    wert) if _ist_zahl(q.value) else None
                c = _zelle(ws, zeile, ziel, q.value, formel)
                if q.has_style:
                    c.font, c.alignment = copy(q.font), copy(q.alignment)
                    if not str(q.fill.fgColor.rgb).endswith(GRAU):   # grau = Formel
                        c.fill = copy(q.fill)
                    c.number_format = (FMT_BETRAG if q.number_format == FMT_EURO
                                       else q.number_format)
            zeile += 1
        zeile += 1
    return zeile


def _blatt_verkauf_kauf(ex, wb, formeln: Formeln) -> bool:
    """Verkäufe, Neuobjekte und Sonderbereich; ohne beides kein Blatt (False)."""
    ws = ex.create_sheet(VERKAUF_KAUF)
    zeile = _titel(ws, "Verkauf und Kauf",
                   "Geplante Verkäufe und Neuobjekte, je Vorgang eine Spalte; Herleitungen "
                   "rechnen mit den Zeilen darüber, Gruppen (+/− am Rand) lassen sich ein- und "
                   "ausblenden. Darunter je Verkauf das Ergebnis (Erlös, Reinvestition, "
                   "Kapitalanlage, Vergleich mit Halten) und die Herleitung von Gewinn und "
                   "§ 6b-Rücklage, dann je Neuobjekt Kauf und Finanzierung. Ausgeschaltete "
                   "Vorgänge sind durchgestrichen und rechnen nicht mit.")
    reihen_vk, reihen_neu = _reihen_verkauf(wb), _reihen_neu(wb)
    breite = 1 + max(sum(v not in (None, "") for v in reihen_vk["vk_ID"]),
                     sum(v not in (None, "") for v in reihen_neu["ne_ID"]), 4)
    _abschnitt(ws, zeile, "Verkäufe", breite)
    zeile, verkaeufe = _vorgaenge(ws, zeile + 1, VERKAUF_ANGABEN, reihen_vk)
    zeile += 1
    _abschnitt(ws, zeile, "Neuobjekte", breite)
    zeile, neu = _vorgaenge(ws, zeile + 1, NEU_ANGABEN, reihen_neu)
    if not verkaeufe and not neu:
        ex.remove(ws)
        return False
    _sonderbereich(ws, zeile + 1, wb[SONDERBEREICH], formeln(SONDERBEREICH))
    ws.column_dimensions["A"].width = 66
    for spalte in range(2, breite + 1):
        ws.column_dimensions[get_column_letter(spalte)].width = 16
    _druck(ws)
    return True


# --- Blatt Annahmen ---

OBJEKT_SPALTEN = ("objekt_id", "name", "miete", "erhaltung", "weitere_einnahmen",
                  "weitere_ausgaben", "afa_bwa", "zinsen", "verkehrswert", "vk_quote_gebaeude",
                  "kaufjahr", "baujahr", "ak_gebaeude", "ak_gub", "restbuchwert", "afa_satz",
                  "afa_jahr", "restschuld", "zinssatz", "rate")
ANNAHME, BUCHHALTUNG = "Annahme", "Buchhaltung"
HERKUNFT_FILL = {ANNAHME: FILL_ANNAHME, BUCHHALTUNG: FILL_EINGELESEN}


def _spalte(buchstabe: str) -> int:
    return column_index_from_string(buchstabe) - 1


def herkunft(formel_zeile: tuple, wert_zeile: tuple) -> dict:
    """Herkunft je Objektfeld wie die Farblogik des Blatts Objekte: Annahme (blau, Formel der
    Auffülllogik), Buchhaltung (grün, Kostenstellenblatt, eingelesen oder Blatt Anlagen) oder
    keine Angabe (eingetippt)."""
    def zahl(k):
        return _zahl(wert_zeile[_spalte(hspalte(k))])

    ergebnis = {}
    restschuld = 0
    for i, f in enumerate(OBJEKT_FELDER):
        inhalt = formel_zeile[i] if i < len(formel_zeile) else None
        eingelesen = (wert_zeile[_spalte(ispalte(f.key))]
                      if f.key in OBJEKT_EINGELESEN else None)
        if f.key == "restschuld":
            restschuld = _zahl(wert_zeile[i])
        if isinstance(inhalt, str) and inhalt.startswith("="):
            if f.key in KST_VERKNUEPFT and eingelesen not in (None, ""):
                ergebnis[f.key] = BUCHHALTUNG
            elif deckt_anlagen(f.key, zahl):
                ergebnis[f.key] = BUCHHALTUNG
            elif f.key in DARLEHEN_ANNAHMEN:
                if restschuld > 0:
                    ergebnis[f.key] = ANNAHME
            elif f.annahme:
                ergebnis[f.key] = ANNAHME
        elif eingelesen not in (None, "") and inhalt == eingelesen:
            ergebnis[f.key] = BUCHHALTUNG
    return ergebnis


def _objektzeilen(ws, anzahl: int) -> list:
    return [tuple(z) for z in ws.iter_rows(min_row=2, max_row=anzahl + 1, values_only=True)]


def _blatt_annahmen(ex, wb, formeln_wb) -> None:
    ws = ex.create_sheet(ANNAHMEN)
    zeile = _titel(ws, "Annahmen und Datenlage",
                   "Die zentralen Annahmen gelten für alle Objekte, soweit keine Daten aus der "
                   "Buchhaltung oder eigene Werte vorliegen. Je Objekt zeigt die Farbe, woher "
                   "ein Wert stammt.")
    _abschnitt(ws, zeile, "Zentrale Annahmen", 2)
    zeile += 1
    for p in [q for q in PARAMETER if q.name in ("par_Basisjahr", "par_Prognosejahre")] \
            + zentrale_annahmen():
        ws.cell(row=zeile, column=1, value=p.bezeichnung)
        c = ws.cell(row=zeile, column=2, value=_wert(wb, p.name))
        c.number_format = FMT_BETRAG if p.format == FMT_EURO else p.format
        c.alignment = Alignment(horizontal="right")
        zeile += 1

    zeile += 1
    felder = [f for f in OBJEKT_FELDER if f.key in OBJEKT_SPALTEN]
    _abschnitt(ws, zeile, "Daten je Objekt (Basisjahr)", len(felder) + 1)
    zeile += 1
    for fill, text in ((FILL_ANNAHME, "Annahme: fehlende Angabe, nach den zentralen Annahmen "
                                      "ergänzt"),
                       (FILL_EINGELESEN, "aus der Buchhaltung (BWA) oder dem "
                                         "Anlagenverzeichnis"),
                       (None, "ohne Farbe: eigens erfasster Wert")):
        c = ws.cell(row=zeile, column=1, value=text)
        if fill:
            c.fill = fill
        zeile += 1
    _kopfzeile(ws, zeile, [f.ueberschrift.rstrip(" *") for f in felder] + ["Status"])
    ws.row_dimensions[zeile].height = 45
    zeile += 1
    anzahl = len(_bereich(wb, "obj_ID"))
    werte_zeilen = _objektzeilen(wb["Objekte"], anzahl)
    formel_zeilen = (_objektzeilen(formeln_wb["Objekte"], anzahl) if formeln_wb is not None
                     else [()] * anzahl)
    status = _bereich(wb, "obj_Status")
    for wert_zeile, formel_zeile, st in zip(werte_zeilen, formel_zeilen, status):
        if wert_zeile[0] in (None, ""):
            continue
        quelle = herkunft(formel_zeile, wert_zeile) if formel_zeile else {}
        for spalte, f in enumerate(felder, start=1):
            wert = wert_zeile[OBJEKT_FELDER.index(f)]
            c = ws.cell(row=zeile, column=spalte, value=wert)
            c.number_format = FMT_BETRAG if f.format == FMT_EURO else f.format
            if f.key in quelle:
                c.fill = HERKUNFT_FILL[quelle[f.key]]
        ws.cell(row=zeile, column=len(felder) + 1, value=st)
        zeile += 1

    zeile += 1
    _abschnitt(ws, zeile, f"Plausibilitätsprüfung: {_wert(wb, 'pr_Gesamt')}", 3)
    zeile += 1
    auffaellig = [(text, art, n) for text, art, n in zip(_bereich(wb, "pr_Bezeichnung"),
                                                         _bereich(wb, "pr_Art"),
                                                         _bereich(wb, "pr_Anzahl"))
                  if _zahl(n) > 0]
    if auffaellig:
        _kopfzeile(ws, zeile, ["Prüfung", "Art", "betroffen"])
        zeile += 1
        for text, art, n in auffaellig:
            ws.cell(row=zeile, column=1, value=text).alignment = Alignment(wrap_text=True)
            ws.cell(row=zeile, column=2, value=art)
            ws.cell(row=zeile, column=3, value=n)
            zeile += 1
    else:
        ws.cell(row=zeile, column=1, value="keine Auffälligkeiten").font = FONT_HINWEIS
    ws.column_dimensions["A"].width = BREITE_TEXT
    for spalte in range(2, len(felder) + 2):
        ws.column_dimensions[get_column_letter(spalte)].width = 14
    _druck(ws)


# --- Mappe ---

def erstelle_export(wb, formeln_wb=None, quelle: str = "") -> Workbook:
    """Export aus der gerechneten Mappe wb (Werte); formeln_wb (dieselbe Mappe mit Formeln)
    liefert die Herkunft der Objektdaten und die Formeln der Summenzeilen in Cashflow,
    Ergebnisbrücke, BWA und Sonderbereich; ohne sie bleiben die Zellen ohne Farbe und dort
    Werte."""
    ex = Workbook()
    ex.formelzellen = set()   # (Blatt, Zelle) der Formeln, siehe _formel
    formeln = Formeln(formeln_wb)
    spalten = _bwa_spalten(wb)
    blaetter = [UEBERBLICK, JAHRESRECHNUNG]
    _blatt_jahresrechnung(ex, wb)
    if _blatt_bruecke(ex, wb, formeln):
        blaetter.append(BRUECKE)
    _blatt_cashflow(ex, wb, formeln)
    blaetter.append(CASHFLOW)
    blaetter.append(BWA_GESAMT)
    _blatt_bwa_gesamt(ex, wb, spalten, formeln)
    if _blatt_bwa_halten(ex, wb, spalten, formeln):
        blaetter.append(BWA_HALTEN)
    blaetter.append(BWA_OBJEKTE)
    _blatt_bwa_objekte(ex, wb, spalten, formeln)
    if _blatt_verkauf_kauf(ex, wb, formeln):
        blaetter.append(VERKAUF_KAUF)
    _blatt_annahmen(ex, wb, formeln_wb)
    blaetter.append(ANNAHMEN)
    _blatt_ueberblick(ex, wb, quelle, blaetter)
    ex.active = 0
    for ws in ex.worksheets:   # Texte wie „= Ergebnis“ sind keine Formeln
        for zeile in ws.iter_rows():
            for c in zeile:
                if c.data_type == "f" and (ws.title, c.coordinate) not in ex.formelzellen:
                    c.data_type = "s"
    return ex


def _soffice():
    return shutil.which("soffice") or shutil.which("libreoffice")


def lade_werte(pfad: Path, arbeitsordner: Path):
    """Mappe mit Werten laden; nie durchgerechnet: per LibreOffice rechnen, falls vorhanden."""
    wb = load_workbook(pfad, data_only=True)
    if "aus_VermoegenNetto" not in wb.defined_names:
        sys.exit(f"{pfad} ist keine Mappe des Prognosemodells.")
    if _wert(wb, "aus_VermoegenNetto") is not None:
        return wb
    soffice = _soffice()
    if not soffice:
        sys.exit(f"{pfad} enthält keine berechneten Werte. Bitte in Excel öffnen, speichern "
                 "und den Export erneut starten.")
    print("Die Mappe enthält keine berechneten Werte, LibreOffice rechnet sie durch.")
    kopie = arbeitsordner / f"mappe{pfad.suffix}"
    shutil.copyfile(pfad, kopie)
    aus = arbeitsordner / "gerechnet"
    subprocess.run(
        [soffice, f"-env:UserInstallation={(arbeitsordner / 'lo-profil').as_uri()}",
         "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(aus), str(kopie)],
        check=True, capture_output=True, timeout=600)
    return load_workbook(aus / "mappe.xlsx", data_only=True)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("mappe", help="gerechnete Mappe (.xlsx oder .xlsm)")
    ap.add_argument("--ausgabe", metavar="PFAD",
                    help="Exportdatei; Standard: <Mappe>_Mandant.xlsx daneben")
    args = ap.parse_args()
    quelle = Path(args.mappe)
    if not quelle.is_file():
        sys.exit(f"{quelle} nicht gefunden.")
    name = f"{quelle.stem}_Mandant.xlsx"
    ziel = Path(args.ausgabe) if args.ausgabe else quelle.with_name(name)
    if ziel.is_dir():
        ziel = ziel / name
    ziel = ziel.with_suffix(".xlsx")
    if ziel.resolve() == quelle.resolve():
        sys.exit("Die Mappe würde überschrieben: bitte mit --ausgabe einen anderen Namen "
                 "angeben.")
    with tempfile.TemporaryDirectory() as tmp:
        wb = lade_werte(quelle, Path(tmp))
        formeln_wb = load_workbook(quelle, read_only=True)
        try:
            ex = erstelle_export(wb, formeln_wb, quelle.name)
        finally:
            formeln_wb.close()
    ziel.parent.mkdir(parents=True, exist_ok=True)
    try:
        ex.save(ziel)
    except PermissionError:
        sys.exit(f"Kann {ziel} nicht schreiben. Ist die Datei noch in Excel geöffnet?")
    print(f"geschrieben: {ziel} ({', '.join(ex.sheetnames)})")


if __name__ == "__main__":
    main()
