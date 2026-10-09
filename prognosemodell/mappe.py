"""Schreibt die Arbeitsmappe: Blätter, Kopfzeilen, benannte Bereiche, Validierung."""

import dataclasses

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.text import RichText
from openpyxl.comments import Comment
from openpyxl.drawing.text import (CharacterProperties, Paragraph, ParagraphProperties,
                                   RichTextProperties)
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from . import bwa, erklaerungen, formeln
from .modelle import (ANLAGE_FELDER, ANLAGE_JAHRE_NAME, ANLAGE_SPALTEN, ANLAGE_STATUS_NAME,
                      MIN_ANLAGEN, OBJEKT_ANLAGEN_SPALTEN,
                      CODENAME_MAPPE, CODENAMEN, FEHLER, OBJEKT_EINGELESEN, FMT_EURO, FMT_JAHR, FMT_PROZENT,
                      KST_VERKNUEPFT, NEUKAUF_POSITIONEN,
                      FMT_TEXT, FMT_ZAHL, HINWEIS, MAX_VARIANTEN, PRUEFUNGEN,
                      VARIANTEN_KOPF, WARNUNG,
                      NEU_FELDER, NEU_SPALTEN, NEU_STATUS_NAME, OBJEKT_FELDER,
                      DARLEHEN_BLOECKE, DARLEHEN_ERSTE, DARLEHEN_SPALTEN, NEU_DETAIL_ERSTE, QUELLEN,
                      PARAMETER, PROGNOSE_SPALTEN, Spalte, RUECKLAGE_JAHR_SPALTEN, RUECKLAGE_SPALTEN,
                      STATUS_NAME, STATUS_UEBERSCHRIFT, SZ_A, SZ_BASELINE, SZENARIEN,
                      VERGLEICH_KENNZAHLEN, VERKAUF_FELDER, VERKAUF_SPALTEN, VERKAUF_STATUS_NAME, Modell,
                      ZUORDNUNG_BEZEICHNUNG, ZUORDNUNG_MEHRDEUTIG, aus_spalten, liq_spalten,
                      prognosejahre)

# Felder, deren Wert aus dem Blatt Anlagen nur eine Vermutung ist: orange statt grün
ANLAGEN_PRUEFEN = ("baujahr",)

HINWEIS_FINANZIERUNG = ("Finanzierung: Bestandsobjekte mit Zinsaufwand, Restschuld, Zinssatz und "
                        "Rate aus dem Blatt Objekte (ohne Restschuld nur der Zins fortgeschrieben, "
                        "beim Verkauf abgelöst); Kredite der Neuobjekte im Blatt Neuobjekte, "
                        "Tilgungsplan im Blatt Darlehen.")

FONT_TITEL = Font(bold=True, size=14)
FONT_KOPF = Font(bold=True, color="FFFFFF")
FILL_KOPF = PatternFill("solid", fgColor="305496")
FILL_EINGABE = PatternFill("solid", fgColor="FFF2CC")   # gelb = hier wird getippt
FILL_BERECHNET = PatternFill("solid", fgColor="E7E6E6")  # grau = Formel
FILL_KOPF_PFLICHT = PatternFill("solid", fgColor="C00000")    # Pflichtfeld, Kopf mit *
FILL_KOPF_BERECHNET = PatternFill("solid", fgColor="595959")  # berechnete Spalte
FILL_KERN = PatternFill("solid", fgColor="FCE4A8")           # Kernparameter
FILL_FEHLER = PatternFill("solid", fgColor="F8CBAD")
FILL_WARNUNG = PatternFill("solid", fgColor="FFE699")
# Farblogik der Eingabezellen (Legende auf dem Startblatt)
FILL_PFLICHT = PatternFill("solid", fgColor="FF7C80")    # rot: Pflichtwert fehlt
FILL_KRITISCH = PatternFill("solid", fgColor="F4B183")   # orange: Annahme bei Verkauf
FILL_ANNAHME = PatternFill("solid", fgColor="BDD7EE")    # blau: Annahme (Formel)
FILL_EINGELESEN = PatternFill("solid", fgColor="C6E0B4")  # grün: aus der Buchhaltung
FARBEN = [
    (FILL_PFLICHT, "rot", "Pflichtwert fehlt oder Annahme gelöscht: bitte eintragen"),
    (FILL_KRITISCH, "orange", "Annahme bei einem verkauften Objekt: bestimmt Gewinn und "
     "§ 6b-Rücklage, möglichst durch echten Wert ersetzen. Baujahr aus dem AHK-Datum des "
     "Gebäudes bzw. ObjektID über die Inventarbezeichnung: prüfen. Im Blatt Objekte auch: "
     "Basiswert weicht vom Kostenstellenblatt ab"),
    (FILL_ANNAHME, "blau", "Annahme aus den zentralen Annahmen (Parameterblatt); "
     "Eintippen ersetzt sie"),
    (FILL_EINGELESEN, "grün", "aus der Buchhaltung eingelesen, unverändert, aus dem "
     "Kostenstellenblatt (Verknüpfung) oder aus dem Anlagenverzeichnis (Blatt Anlagen)"),
    (FILL_EINGABE, "gelb", "händisch eingetragen bzw. Eingabefeld"),
]
ISF = formeln.ISFORMEL
KOPF_LEGENDE = [
    (FILL_KOPF_PFLICHT, "Kopf rot *", "Pflichtfeld: ohne Wert rechnet die Zeile nicht"),
    (FILL_KOPF, "Kopf blau", "optionale Eingabe; leer gilt die Annahme bzw. der Standard"),
    (FILL_KOPF_BERECHNET, "Kopf grau", "berechnete Spalte, nicht überschreiben"),
]
TOOLTIP_HINWEIS = ("Kopfzeilen mit rotem Dreieck: Maus darüber zeigt Erläuterung bzw. "
                   "Herleitung der Spalte.")


def _name(wb, name: str, ref: str) -> None:
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def _kopf(ws, zeile: int, werte: list) -> None:
    for spalte, wert in enumerate(werte, start=1):
        c = ws.cell(row=zeile, column=spalte, value=wert)
        c.font = FONT_KOPF
        c.fill = FILL_KOPF
        c.alignment = Alignment(wrap_text=True, vertical="center")


def _tooltip(zelle, text: str) -> None:
    """Kommentar als Tooltip (Maus über der Zelle); Größe nach Textlänge."""
    if not text:
        return
    zeilen = sum(len(absatz) // 48 + 1 for absatz in text.split("\n"))
    zelle.comment = Comment(text, "Prognosemodell", height=min(40 + 15 * zeilen, 400), width=320)


def _feld_text(f) -> str:
    """Tooltip eines Eingabefelds: Pflicht oder optional, dann die Eingabehilfe."""
    if f.pflicht and f.annahme:
        art = "Pflichtfeld; leer steht eine Annahme (blau/orange)."
    elif f.pflicht:
        art = "Pflichtfeld."
    elif f.annahme:
        art = "Optional; leer steht eine Annahme (blau)."
    else:
        art = "Optional."
    teile = [art, f.hinweis]
    if f.auswahl:
        teile.append("Auswahl: " + ", ".join(f.auswahl))
    return "\n".join(t for t in teile if t)


def _kopf_felder(ws, zeile: int, felder, start: int = 1) -> None:
    """Kopf der Eingabefelder: Pflichtfelder rot mit *, Tooltip je Feld."""
    for i, f in enumerate(felder, start=start):
        c = ws.cell(row=zeile, column=i)
        if f.pflicht:
            c.value = f"{f.ueberschrift} *"
            c.fill = FILL_KOPF_PFLICHT
        _tooltip(c, _feld_text(f))


def _kopf_berechnet(ws, zeile: int, spalten, start: int = 1, texte: dict = None,
                    grau: bool = True) -> None:
    """Kopf berechneter Spalten: grau (in Blättern mit Eingaben), Herleitung als Tooltip."""
    texte = erklaerungen.SPALTEN if texte is None else texte
    for i, s_ in enumerate(spalten, start=start):
        c = ws.cell(row=zeile, column=i)
        if grau:
            c.fill = FILL_KOPF_BERECHNET
        text = texte.get(s_.name) or texte.get(s_.key, "")
        if text:
            _tooltip(c, "Berechnet: " + text)


def _validierung(ws, f, bereich: str, liste: str = None) -> None:
    """Genau eine Datenüberprüfung je Eingabespalte: Auswahlliste (auch aus einem benannten
    Bereich, liste) oder Zahlenbereich, dazu die Eingabehilfe des Felds beim Anklicken."""
    if liste:
        dv = DataValidation(type="list", formula1=liste, allow_blank=True)
    elif f.auswahl:
        dv = DataValidation(type="list", formula1='"' + ",".join(f.auswahl) + '"',
                            allow_blank=True)
    elif f.minimum is not None or f.maximum is not None:
        dv = DataValidation(
            type="whole" if f.ganzzahl else "decimal",
            operator="between",
            formula1=str(f.minimum if f.minimum is not None else -1e15),
            formula2=str(f.maximum if f.maximum is not None else 1e15),
            allow_blank=True,
        )
    elif f.hinweis:
        dv = DataValidation(allow_blank=True)
    else:
        return
    dv.showErrorMessage = bool(f.auswahl or f.minimum is not None or f.maximum is not None)
    dv.errorTitle = "Ungültiger Wert"
    dv.error = f"{f.ueberschrift}: Wert außerhalb des zulässigen Bereichs."
    if f.hinweis:
        dv.showInputMessage = True
        dv.promptTitle = f.ueberschrift[:32]
        dv.prompt = f.hinweis if len(f.hinweis) <= 255 else f.hinweis[:250].rsplit(" ", 1)[0] + " …"
    ws.add_data_validation(dv)
    dv.add(bereich)


def _farblogik(ws, spalte: str, erste: int, letzte: int, f, id_spalte: str = "A",
               verkauft: str = None, eingelesen: str = None, anlagen: str = None,
               anlagen_fill=FILL_EINGELESEN, kst: str = None) -> None:
    """Bedingte Formate einer Eingabespalte, in dieser Reihenfolge (erste Regel gewinnt):
    rot Pflicht/Annahme fehlt, grün aus dem Blatt Anlagen (anlagen_fill, orange = prüfen),
    grün verknüpft mit dem Kostenstellenblatt (kst = Zelle mit der Verknüpfung), orange davon
    abweichend, orange kritische Annahme, blau Annahme, grün eingelesen."""
    bereich = f"{spalte}{erste}:{spalte}{letzte}"
    z, id_ = f"{spalte}{erste}", f"${id_spalte}{erste}"
    regeln = []
    if f.pflicht or f.annahme:
        regeln.append((f'AND({id_}<>"",{z}="")', FILL_PFLICHT))
    if anlagen:
        regeln.append((f'AND({id_}<>"",{ISF}({z}),{anlagen})', anlagen_fill))
    if kst:
        regeln.append((f'AND({id_}<>"",{kst}<>"",{ISF}({z}))', FILL_EINGELESEN))
        regeln.append((f'AND({id_}<>"",{kst}<>"",{z}<>{kst})', FILL_KRITISCH))
    if f.annahme and f.kritisch:
        bedingung = f",{verkauft}" if verkauft else ""
        regeln.append((f'AND({id_}<>"",{ISF}({z}){bedingung})', FILL_KRITISCH))
    if f.annahme:
        regeln.append((f'AND({id_}<>"",{ISF}({z}))', FILL_ANNAHME))
    if eingelesen:
        regeln.append((f'AND({id_}<>"",{z}<>"",{z}={eingelesen})', FILL_EINGELESEN))
    for formel, fill in regeln:
        ws.conditional_formatting.add(
            bereich, FormulaRule(formula=[formel], fill=fill, stopIfTrue=True))


def _status_rot(ws, bereich: str, erste_zelle: str) -> None:
    """Status rot, sobald er gesetzt ist und nicht OK lautet."""
    ws.conditional_formatting.add(
        bereich,
        FormulaRule(formula=[f'AND({erste_zelle}<>"",{erste_zelle}<>"OK")'], fill=FILL_FEHLER),
    )


def _blatt_parameter(wb, modell: Modell) -> None:
    ws = wb.active
    ws.title = "Parameter"
    ws["A1"] = "Parameter"
    ws["A1"].font = FONT_TITEL
    ws["A2"] = HINWEIS_FINANZIERUNG
    _kopf(ws, 4, ["Bezeichnung", "Wert", "Name", "Erläuterung"])

    unbekannt = set(modell.parameter) - {p.name for p in PARAMETER}
    if unbekannt:
        raise ValueError(f"unbekannte Parameter: {sorted(unbekannt)}")
    ws["A3"] = ("★ = Kernparameter, zuerst prüfen. Abschnitte ohne Kernparameter sind "
                "eingeklappt: + am linken Rand öffnet sie. Maus über der Bezeichnung zeigt, was "
                "der Parameter beeinflusst.")
    ws["A3"].font = Font(italic=True)
    zeile = 4
    zeile_pruefung = None
    abschnitte = []   # [erste Zeile, letzte Zeile, mit Kernparameter]
    for p in PARAMETER:
        zeile += 1
        if p.abschnitt:
            zeile += 1
            ws.cell(row=zeile, column=1, value=p.abschnitt).font = Font(bold=True)
            zeile += 1
            abschnitte.append([zeile, zeile, False])
        elif abschnitte:
            abschnitte[-1][1] = zeile
        kern = p.name in erklaerungen.KERNPARAMETER
        if abschnitte and kern:
            abschnitte[-1][2] = True
        a = ws.cell(row=zeile, column=1, value=("★ " if kern else "") + p.bezeichnung)
        _tooltip(a, erklaerungen.PARAMETER_WIRKUNG.get(p.name, ""))
        if kern:
            a.font = Font(bold=True)
            a.fill = FILL_KERN
        c = ws.cell(row=zeile, column=2, value=modell.parameter.get(p.name, p.wert))
        c.number_format = p.format
        berechnet = isinstance(p.wert, str) and p.wert.startswith("=")
        c.fill = FILL_BERECHNET if berechnet else FILL_EINGABE
        ws.cell(row=zeile, column=3, value=p.name)
        ws.cell(row=zeile, column=4, value=p.erlaeuterung)
        _name(wb, p.name, f"Parameter!$B${zeile}")
        if p.auswahl:
            dv = DataValidation(type="list", formula1='"' + ",".join(p.auswahl) + '"',
                                allow_blank=False)
            ws.add_data_validation(dv)
            dv.add(c.coordinate)
        if p.name == "par_StatusPruefung":
            zeile_pruefung = zeile

    ws.sheet_properties.outlinePr.summaryBelow = False
    for von, bis, kern in abschnitte:
        ws.row_dimensions.group(von, bis, hidden=not kern, outline_level=1)
        ws.row_dimensions[von - 1].collapsed = not kern
    ws.conditional_formatting.add(
        f"B{zeile_pruefung}",
        FormulaRule(formula=[f'B{zeile_pruefung}<>"OK"'], fill=FILL_FEHLER))
    ws["F3"] = ("Steuerung per Makro (Datei als .xlsm, Makros aktivieren): die Schaltflächen "
                "unten. Gerechnet wird immer in den Formeln, auch ohne Makros.")
    ws["F3"].font = Font(italic=True)

    for spalte, breite in zip("ABCD", (36, 14, 22, 60)):
        ws.column_dimensions[spalte].width = breite
    ws.freeze_panes = "A5"


def _blatt_objekte(wb, modell: Modell) -> None:
    ws = wb.create_sheet("Objekte")
    erste, letzte = 2, modell.kapazitaet.objekte + 1
    status_spalte = len(OBJEKT_FELDER) + 1
    hilfe = [(STATUS_UEBERSCHRIFT, STATUS_NAME, 26, formeln.status_objekt),
             ("Annahmen (blau)", "obj_Annahmen", 11, formeln.annahmen_objekt),
             ("kritische Annahmen (orange)", "obj_Kritisch", 11,
              lambda z: formeln.annahmen_objekt(z, nur_kritisch=True))]
    # was das Blatt Anlagen je Objekt liefert (formeln.hspalte kennt die Lage)
    hilfe += [(s_.ueberschrift, s_.name, s_.breite,
               lambda z, k=s_.key: formeln.objekt_anlagen(k, z)) for s_ in OBJEKT_ANLAGEN_SPALTEN]
    formate = {s_.name: s_.format for s_ in OBJEKT_ANLAGEN_SPALTEN}
    # eingelesene Werte, ausgeblendet: grün, solange die Eingabe ihnen gleicht; bei
    # Objekten mit Kostenstellenblatt die Verknüpfung dorthin (bwa.verknuepfe_objekte)
    eingelesen = [f for f in OBJEKT_FELDER if f.key in OBJEKT_EINGELESEN]
    import_spalte = {f.key: formeln.ispalte(f.key) for f in eingelesen}
    assert formeln.OBJEKT_HILFSSPALTEN == len(hilfe)
    _kopf(ws, 1, [f.ueberschrift for f in OBJEKT_FELDER] + [h[0] for h in hilfe])
    _kopf_felder(ws, 1, OBJEKT_FELDER)
    _kopf_berechnet(ws, 1, [Spalte(name, u, name, FMT_ZAHL) for u, name, _, _ in hilfe],
                    start=status_spalte)
    for f in eingelesen:
        c = ws[f"{import_spalte[f.key]}1"]
        c.value = f"eingelesen: {f.ueberschrift}"
        c.font = Font(italic=True)
        if f.key in KST_VERKNUEPFT:
            bst = import_spalte[f.key]
            _name(wb, f"obk_{f.key}", f"Objekte!${bst}${erste}:${bst}${letzte}")
    kst_blatt = formeln.kst_blatt_spalte()
    ws[f"{kst_blatt}1"] = "Kostenstellenblatt"
    ws[f"{kst_blatt}1"].font = Font(italic=True)
    _name(wb, "obj_KStBlatt", f"Objekte!${kst_blatt}${erste}:${kst_blatt}${letzte}")
    ws.row_dimensions[1].height = 45

    for i, f in enumerate(OBJEKT_FELDER, start=1):
        bst = get_column_letter(i)
        ws.column_dimensions[bst].width = f.breite
        _name(wb, f.name, f"Objekte!${bst}${erste}:${bst}${letzte}")
        _validierung(ws, f, f"{bst}{erste}:{bst}{letzte}")
        _farblogik(ws, bst, erste, letzte, f, verkauft=f"COUNTIF(vk_ID,$A{erste})>0",
                   eingelesen=f"${import_spalte[f.key]}{erste}" if f.key in import_spalte
                   else None, anlagen=formeln.deckung_anlagen(f.key, erste),
                   anlagen_fill=FILL_KRITISCH if f.key in ANLAGEN_PRUEFEN else FILL_EINGELESEN,
                   kst=f"${import_spalte[f.key]}{erste}" if f.key in KST_VERKNUEPFT else None)

    for j, (_, name, breite, _) in enumerate(hilfe):
        bst = get_column_letter(status_spalte + j)
        ws.column_dimensions[bst].width = breite
        _name(wb, name, f"Objekte!${bst}${erste}:${bst}${letzte}")
    letzte_spalte = get_column_letter(status_spalte)
    _name(wb, "obj_Basis", f"Objekte!$A${erste}:${letzte_spalte}${letzte}")
    # nur die gelben Eingabefelder; die Makros in modObjekte schreiben nur hierhin
    eingabe_ende = get_column_letter(len(OBJEKT_FELDER))
    _name(wb, "obj_Eingabe", f"Objekte!$A${erste}:${eingabe_ende}${letzte}")

    def annahmen(zeile):
        for i, f in enumerate(OBJEKT_FELDER, start=1):
            c = ws.cell(row=zeile, column=i, value=formeln.annahme_objekt(f.key, zeile))
            c.number_format = f.format
            c.fill = FILL_EINGABE

    for zeile in range(erste, letzte + 1):
        annahmen(zeile)
        for j, (_, name, _, formel) in enumerate(hilfe):
            c = ws.cell(row=zeile, column=status_spalte + j, value=formel(zeile))
            c.fill = FILL_BERECHNET
            if name in formate:
                c.number_format = formate[name]
    st = get_column_letter(status_spalte)
    _status_rot(ws, f"{st}{erste}:{st}{letzte}", f"{st}{erste}")

    # Vorlagezeile mit den Annahmeformeln: „Annahmen wiederherstellen“ kopiert sie
    vorlage = letzte + 2
    annahmen(vorlage)
    ws.cell(row=vorlage, column=status_spalte, value="Vorlage der Annahmeformeln (Makro)")
    ws.row_dimensions[vorlage].hidden = True
    _name(wb, "obj_Vorlage", f"Objekte!$A${vorlage}:${eingabe_ende}${vorlage}")

    for zeile, obj in enumerate(modell.objekte, start=erste):
        for i, f in enumerate(OBJEKT_FELDER, start=1):
            wert = getattr(obj, f.key)
            if wert is not None:      # leer = Annahmeformel bleibt stehen
                ws.cell(row=zeile, column=i, value=wert)
        ist = modell.kostenstellen.get(obj.objekt_id)
        if ist is not None:
            quelle = {"name": ist.name, "miete": ist.miete, "erhaltung": ist.erhaltung,
                      "weitere_einnahmen": ist.weitere_einnahmen,
                      "weitere_ausgaben": ist.weitere_ausgaben, "afa_bwa": ist.abschreibung,
                      "zinsen": ist.zinsen}
            for key, bst in import_spalte.items():
                ws[f"{bst}{zeile}"] = quelle[key]

    for bst in list(import_spalte.values()) + [kst_blatt]:
        ws.column_dimensions[bst].hidden = True
    # Anlagen je Objekt als zuklappbare Gruppe
    anl_erste = status_spalte + len(hilfe) - len(OBJEKT_ANLAGEN_SPALTEN)
    ws.column_dimensions.group(get_column_letter(anl_erste),
                               get_column_letter(status_spalte + len(hilfe) - 1),
                               hidden=modell.schnellcheck, outline_level=1)
    # Steuerliche Stammdaten als Gruppe; im Schnellcheck zugeklappt
    details = [i for i, f in enumerate(OBJEKT_FELDER, start=1)
               if f.key not in ("objekt_id", "name", "miete", "erhaltung", "verkehrswert")
               and f.key not in OBJEKT_EINGELESEN]
    ws.column_dimensions.group(get_column_letter(details[0]), get_column_letter(details[-1]),
                               hidden=modell.schnellcheck, outline_level=1)
    ws.freeze_panes = "B2"
    ws["A" + str(vorlage + 2)] = ("Farben: rot = Pflicht fehlt, orange = Annahme bei verkauftem "
                                  "Objekt, blau = Annahme, grün = eingelesen oder aus dem Blatt "
                                  "Anlagen, gelb = händisch. Annahmen stellt das Parameterblatt "
                                  "ein.")
    ws["A" + str(vorlage + 3)] = (
        "Objekte mit Kostenstellenblatt: Miete, Erhaltung, weitere Einnahmen und Ausgaben, "
        "AfA lt. Buchhaltung und Zinsaufwand verweisen auf dessen Spalte Basisjahr (grün). Werte dort ändern, "
        "dann folgen Objekte, Prognose und BWA Alle Objekte. Hier überschrieben (orange) "
        "rechnet nur die Prognose damit; das Makro „Objekte → Kostenstellen“ schreibt solche "
        "Werte ins Kostenstellenblatt und stellt die Verknüpfung wieder her.")


def _blatt_anlagen(wb, modell: Modell) -> None:
    """Anlagenverzeichnis: je Anlage Buchwert und AfA je Jahr, Zuordnung über die ObjektID."""
    ws = wb.create_sheet("Anlagen")
    jahre = prognosejahre()
    erste = 3
    letzte = erste + max(MIN_ANLAGEN, len(modell.anlagen) + 50) - 1
    status_spalte = len(ANLAGE_FELDER) + len(ANLAGE_SPALTEN) + 1
    jahr_spalte = status_spalte + 1
    ws["A1"] = ("Anlagenverzeichnis: gelb Eingabe (beim Einlesen aus dem DATEV-Export bzw. der "
                "Aufschlüsselung im Kostenstellenblatt), grau Formel, orange ObjektID über die "
                "Bezeichnung gefunden (prüfen). Es zählen nur Zeilen mit Status OK. Buchwert "
                "Stand = Ende des Jahres")
    ws["A1"].font = Font(italic=True)
    c = ws.cell(row=1, column=12, value="=par_AnlStand")
    c.number_format, c.fill = FMT_JAHR, FILL_BERECHNET
    ws.cell(row=1, column=13, value="(Parameterblatt).").font = Font(italic=True)
    kopf = ([f.ueberschrift for f in ANLAGE_FELDER] + [s_.ueberschrift for s_ in ANLAGE_SPALTEN]
            + [STATUS_UEBERSCHRIFT])
    _kopf(ws, 2, kopf)
    _kopf_felder(ws, 2, ANLAGE_FELDER)
    _kopf_berechnet(ws, 2, ANLAGE_SPALTEN + [Spalte("status", "", ANLAGE_STATUS_NAME, FMT_TEXT)],
                    start=len(ANLAGE_FELDER) + 1)
    for i in range(jahre):
        c = ws.cell(row=2, column=jahr_spalte + i, value=f'="AfA "&(par_Startjahr+{i})')
        c.font, c.fill = FONT_KOPF, FILL_KOPF
    ws.row_dimensions[2].height = 45

    for i, f in enumerate(ANLAGE_FELDER, start=1):
        bst = get_column_letter(i)
        ws.column_dimensions[bst].width = f.breite
        _name(wb, f.name, f"Anlagen!${bst}${erste}:${bst}${letzte}")
        _validierung(ws, f, f"{bst}{erste}:{bst}{letzte}")
        _farblogik(ws, bst, erste, letzte, f)
    for i, s_ in enumerate(ANLAGE_SPALTEN, start=len(ANLAGE_FELDER) + 1):
        bst = get_column_letter(i)
        ws.column_dimensions[bst].width = s_.breite
        _name(wb, s_.name, f"Anlagen!${bst}${erste}:${bst}${letzte}")
    st = get_column_letter(status_spalte)
    ws.column_dimensions[st].width = 30
    _name(wb, ANLAGE_STATUS_NAME, f"Anlagen!${st}${erste}:${st}${letzte}")
    _name(wb, ANLAGE_JAHRE_NAME,
          f"Anlagen!${get_column_letter(jahr_spalte)}${erste}:"
          f"${get_column_letter(jahr_spalte + jahre - 1)}${letzte}")

    for zeile in range(erste, letzte + 1):
        for i, f in enumerate(ANLAGE_FELDER, start=1):
            c = ws.cell(row=zeile, column=i, value=formeln.annahme_anlage(f.key, zeile))
            c.number_format, c.fill = f.format, FILL_EINGABE
        berechnet = formeln.anlage_zeile(zeile)
        werte = [berechnet[s_.key] for s_ in ANLAGE_SPALTEN] + [berechnet["status"]] \
            + berechnet["jahre"]
        formate = [s_.format for s_ in ANLAGE_SPALTEN] + [FMT_TEXT] + [FMT_EURO] * jahre
        for j, (formel, fmt) in enumerate(zip(werte, formate)):
            c = ws.cell(row=zeile, column=len(ANLAGE_FELDER) + 1 + j, value=formel)
            c.number_format, c.fill = fmt, FILL_BERECHNET
    ws.conditional_formatting.add(
        f"{st}{erste}:{st}{letzte}",
        FormulaRule(formula=[f'AND({st}{erste}<>"",{st}{erste}<>"OK")'], fill=FILL_WARNUNG))
    # orange: ObjektID über die Bezeichnung gefunden oder mehrdeutig, bitte prüfen
    zu, oid = formeln.aspalte("zuordnung"), formeln.aspalte("objekt_id")
    ws.conditional_formatting.add(
        f"{oid}{erste}:{zu}{letzte}",
        FormulaRule(formula=[f'OR(LEFT(${zu}{erste},{len(ZUORDNUNG_BEZEICHNUNG)})='
                             f'"{ZUORDNUNG_BEZEICHNUNG}",LEFT(${zu}{erste},'
                             f'{len(ZUORDNUNG_MEHRDEUTIG)})="{ZUORDNUNG_MEHRDEUTIG}")'],
                    fill=FILL_KRITISCH, stopIfTrue=True))

    for zeile, anlage in enumerate(modell.anlagen, start=erste):
        for i, f in enumerate(ANLAGE_FELDER, start=1):
            wert = getattr(anlage, f.key)
            if wert is not None:      # leer = Formel bleibt stehen
                ws.cell(row=zeile, column=i, value=wert)
    ws.freeze_panes = "B3"
    ws.auto_filter.ref = f"A2:{get_column_letter(jahr_spalte + jahre - 1)}{letzte}"


def _blatt_afa_plan(wb, modell: Modell) -> None:
    """Schon geplante AfA je Objekt und Jahr (BWA 1240 der Kostenstellen-Datei): Jahre mit
    Wert ersetzen die Fortschreibung bei Halten, leere Jahre rechnet das Modell selbst."""
    ws = wb.create_sheet("AfA-Plan")
    jahre = prognosejahre()
    erste, letzte = 3, modell.kapazitaet.objekte + 2
    j0 = 4                       # erste Jahresspalte (D)
    import0 = j0 + jahre + 1     # ausgeblendete Kopie der eingelesenen Werte
    ws["A1"] = ("AfA-Plan: Abschreibungen je Objekt und Jahr, wie in der Kostenstellen-Datei "
                "(BWA 1240) schon geplant. Ein Jahr mit Wert ersetzt die AfA-Fortschreibung "
                "des Modells (bei Halten, höchstens bis zum Restbuchwert); leere Jahre "
                "schreibt das Modell fort. Grün = eingelesen, gelb = Eingabe.")
    ws["A1"].font = Font(italic=True)
    _kopf(ws, 2, ["ObjektID", "Objektname", "Jahre mit Wert"])
    for i in range(jahre):
        c = ws.cell(row=2, column=j0 + i, value=f"=par_Startjahr+{i}")
        c.font, c.fill, c.number_format = FONT_KOPF, FILL_KOPF, FMT_JAHR
    ende = get_column_letter(j0 + jahre - 1)
    _name(wb, "afp_ID", f"'AfA-Plan'!$A${erste}:$A${letzte}")
    _name(wb, "afp_Anzahl", f"'AfA-Plan'!$C${erste}:$C${letzte}")
    _name(wb, "afp_Jahre", f"'AfA-Plan'!${get_column_letter(j0)}${erste}:${ende}${letzte}")
    for zeile in range(erste, letzte + 1):
        ws.cell(row=zeile, column=1).fill = FILL_EINGABE
        c = ws.cell(row=zeile, column=2, value=f'=IF($A{zeile}="","",IFERROR(INDEX(obj_Name,'
                                               f'MATCH($A{zeile},obj_ID,0)),"ObjektID fehlt"))')
        c.fill = FILL_BERECHNET
        c = ws.cell(row=zeile, column=3,
                    value=f"=COUNT(${get_column_letter(j0)}{zeile}:${ende}{zeile})")
        c.fill = FILL_BERECHNET
        for i in range(jahre):
            c = ws.cell(row=zeile, column=j0 + i)
            c.fill, c.number_format = FILL_EINGABE, FMT_EURO
    zeile = erste
    for obj in modell.objekte:
        if not obj.objekt_id:
            continue
        ws.cell(row=zeile, column=1, value=obj.objekt_id)
        ist = modell.kostenstellen.get(obj.objekt_id)
        for jahr, wert in (ist.afa_plan if ist is not None else {}).items():
            i = jahr - bwa._basisjahr(modell) - 1
            if 0 <= i < jahre:
                ws.cell(row=zeile, column=j0 + i, value=wert)
                ws.cell(row=zeile, column=import0 + i, value=wert)
        zeile += 1
    erste_import = get_column_letter(import0)
    ws.conditional_formatting.add(
        f"{get_column_letter(j0)}{erste}:{ende}{letzte}",
        FormulaRule(formula=[f"AND(ISNUMBER({get_column_letter(j0)}{erste}),"
                             f"{get_column_letter(j0)}{erste}={erste_import}{erste})"],
                    fill=FILL_EINGELESEN))
    for i in range(jahre):
        ws.column_dimensions[get_column_letter(import0 + i)].hidden = True
        ws.column_dimensions[get_column_letter(j0 + i)].width = 12
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 10
    ws.freeze_panes = "D3"


def _blatt_neukauf(wb, modell: Modell) -> None:
    """Neukauf-Kostenstellen (in der Kostenstellen-Datei hinter „KSt 9999“): Planwerte je
    Jahr für Neuobjekte, die im Blatt Neuobjekte auf die Kostenstelle verweisen. Jahre ohne
    Wert schreibt das Blatt vom letzten Wert mit der Steigerung des Parameterblatts fort."""
    ws = wb.create_sheet("Neukauf-KSt")
    # Jahresspalten ab E (Basisjahr bis Endjahr), fortgeschrieben (Formeln), ausgeblendete
    # Kopie der eingelesenen Werte; mit eigenem Blatt verknüpft bwa.verknuepfe_neukauf
    j0, fort0, import0, n = formeln.nk_spalten()
    je = len(NEUKAUF_POSITIONEN)
    erste = formeln.NK_ERSTE
    max_nk = modell.kapazitaet.neukauf
    letzte = erste + max_nk * je - 1
    schluessel = get_column_letter(import0 + n)
    ws["A1"] = ("Neukauf-Kostenstellen: Planwerte je Jahr für Neuobjekte (Blatt Neuobjekte, "
                "Spalte Kostenstelle Neukauf). Gelb = Eingabe, grün = eingelesen (hinter "
                "KSt 9999); das Basisjahr verweist auf das Kostenstellenblatt der Mappe, dort "
                "ändern. Jahre ohne Wert schreibt das Modell vom letzten Wert mit der "
                "Steigerung fort (rechts), 0 vor dem ersten Wert gilt als leer; ab dem Jahr "
                "nach dem Kauf ersetzen sie Mietrendite und Erhaltungsquote des Neuobjekts.")
    ws["A1"].font = Font(italic=True)
    c = ws.cell(row=1, column=fort0, value="fortgeschrieben (rechnet so im Neuobjekt)")
    c.font = Font(bold=True)
    _kopf(ws, 2, ["Kostenstelle", "Bezeichnung", "Position", "BWA-Nr."])
    for i in range(n):
        for spalte in (j0 + i, fort0 + i):
            c = ws.cell(row=2, column=spalte, value=f"=par_Basisjahr+{i}")
            c.font, c.fill, c.number_format = FONT_KOPF, FILL_KOPF, FMT_JAHR
    _name(wb, "nk_ID", f"'Neukauf-KSt'!$A${erste}:$A${letzte}")
    _name(wb, "nk_Schluessel", f"'Neukauf-KSt'!${schluessel}${erste}:${schluessel}${letzte}")
    _name(wb, "nk_Fort", f"'Neukauf-KSt'!${get_column_letter(fort0)}${erste}:"
                         f"${get_column_letter(fort0 + n - 1)}${letzte}")
    # je Kostenstelle eine Zeile: Auswahlliste der Spalte Kostenstelle Neukauf (Neuobjekte)
    liste = get_column_letter(import0 + n + 1)
    for block in range(max_nk):
        ws[f"{liste}{erste + block}"] = f'=IF($A${erste + block * je}="","",$A${erste + block * je})'
    _name(wb, "nk_Liste", f"'Neukauf-KSt'!${liste}${erste}:${liste}${erste + max_nk - 1}")
    for block in range(max_nk):
        kopf = erste + block * je
        for k, (key, text, nr, satz) in enumerate(NEUKAUF_POSITIONEN):
            zeile = kopf + k
            for spalte in (1, 2):
                c = ws.cell(row=zeile, column=spalte)
                if k == 0:
                    c.fill = FILL_EINGABE
                else:
                    bst = get_column_letter(spalte)
                    c.value, c.fill = f'=IF({bst}{kopf}="","",{bst}{kopf})', FILL_BERECHNET
            ws.cell(row=zeile, column=3, value=text)
            ws.cell(row=zeile, column=4, value=nr)
            for i in range(n):
                c = ws.cell(row=zeile, column=j0 + i)
                c.fill, c.number_format = FILL_EINGABE, FMT_EURO
                wert = f"{get_column_letter(j0 + i)}{zeile}"
                vor = f"{get_column_letter(fort0 + i - 1)}{zeile}"
                # 0 vor dem ersten anderen Wert gilt als leer: Kostenstelle ohne Buchungen,
                # das Neuobjekt rechnet dann mit Mietrendite und Erhaltungsquote
                formel = (f'=IF(N({wert})<>0,{wert},"")' if i == 0 else
                          f'=IF(AND(ISNUMBER({wert}),OR(N({wert})<>0,ISNUMBER({vor}))),{wert},'
                          f'IF(ISNUMBER({vor}),{vor}*(1+{satz}),""))')
                c = ws.cell(row=zeile, column=fort0 + i, value=formel)
                c.fill, c.number_format = FILL_BERECHNET, FMT_EURO
            ws.cell(row=zeile, column=import0 + n,
                    value=f'=IF($A{zeile}="","",$A{zeile}&"|{key}")')
    basisjahr = bwa._basisjahr(modell)
    for block, lw in enumerate(list(modell.neukauf.values())[:modell.kapazitaet.neukauf]):
        kopf = erste + block * je
        ws.cell(row=kopf, column=1, value=lw.objekt_id)
        ws.cell(row=kopf, column=2, value=lw.name)
        for k, (key, _, _, _) in enumerate(NEUKAUF_POSITIONEN):
            for jahr, wert in lw.jahre.get(key, {}).items():
                i = jahr - basisjahr
                if 0 <= i < n:
                    ws.cell(row=kopf + k, column=j0 + i, value=wert)
                    ws.cell(row=kopf + k, column=import0 + i, value=wert)
    a, imp = get_column_letter(j0), get_column_letter(import0)
    ws.conditional_formatting.add(
        f"{a}{erste}:{get_column_letter(j0 + n - 1)}{letzte}",
        FormulaRule(formula=[f"AND(ISNUMBER({a}{erste}),{a}{erste}={imp}{erste})"],
                    fill=FILL_EINGELESEN))
    for i in range(n + 2):
        ws.column_dimensions[get_column_letter(import0 + i)].hidden = True
    for i in range(n):
        ws.column_dimensions[get_column_letter(j0 + i)].width = 11
        ws.column_dimensions[get_column_letter(fort0 + i)].width = 11
    for spalte, breite in zip("ABCD", (12, 22, 18, 14)):
        ws.column_dimensions[spalte].width = breite
    ws.freeze_panes = "E3"


def _blatt_prognose(wb, modell: Modell) -> None:
    """Je Zeile des Objektblatts ein Block mit einer Zeile je Prognosejahr, danach je Neuobjekt."""
    ws = wb.create_sheet("Prognose")
    jahre = prognosejahre()
    k = modell.kapazitaet
    erste, letzte = 2, (k.objekte + k.neuobjekte) * jahre + 1
    _kopf(ws, 1, [s.ueberschrift for s in PROGNOSE_SPALTEN])
    _kopf_berechnet(ws, 1, PROGNOSE_SPALTEN, grau=False)
    ws.row_dimensions[1].height = 32

    for i, s in enumerate(PROGNOSE_SPALTEN, start=1):
        bst = get_column_letter(i)
        ws.column_dimensions[bst].width = s.breite
        _name(wb, s.name, f"Prognose!${bst}${erste}:${bst}${letzte}")

    bloecke = ([(nr, formeln.prognose_zeile) for nr in range(1, k.objekte + 1)]
               + [(nr, formeln.prognose_zeile_neu) for nr in range(1, k.neuobjekte + 1)])
    for block, (nr, zeilenformeln) in enumerate(bloecke):
        for j in range(jahre):
            zeile = erste + block * jahre + j
            formeln_zeile = zeilenformeln(zeile, nr, erstes_jahr=j == 0)
            for i, s in enumerate(PROGNOSE_SPALTEN, start=1):
                c = ws.cell(row=zeile, column=i, value=formeln_zeile[s.key])
                c.number_format = s.format
                c.fill = FILL_BERECHNET

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(PROGNOSE_SPALTEN))}{letzte}"


def _blatt_verkaeufe(wb, modell: Modell) -> None:
    """Geplante Verkäufe: Eingaben, Aufteilung des Erlöses und Gewinn je Gebäude und G+B."""
    ws = wb.create_sheet("Verkäufe")
    erste, letzte = 2, modell.kapazitaet.verkaeufe + 1
    spalten = VERKAUF_FELDER + VERKAUF_SPALTEN
    status_spalte = len(spalten) + 1
    st = get_column_letter(status_spalte)
    _kopf(ws, 1, [s.ueberschrift for s in spalten] + [STATUS_UEBERSCHRIFT])
    _kopf_felder(ws, 1, VERKAUF_FELDER)
    _kopf_berechnet(ws, 1, VERKAUF_SPALTEN + [
        Spalte("status", "", VERKAUF_STATUS_NAME, FMT_TEXT),
        Spalte("preis_annahme", "", "vk_PreisAnnahme", FMT_ZAHL)], start=len(VERKAUF_FELDER) + 1)
    ws.row_dimensions[1].height = 45

    for i, s in enumerate(spalten, start=1):
        bst = get_column_letter(i)
        ws.column_dimensions[bst].width = s.breite
        _name(wb, s.name, f"'Verkäufe'!${bst}${erste}:${bst}${letzte}")
    for i, f in enumerate(VERKAUF_FELDER, start=1):
        bst = get_column_letter(i)
        # ObjektID als Auswahl aus dem Objektblatt, in derselben Überprüfung wie die
        # Eingabehilfe: zwei Überprüfungen auf einer Zelle verwirft Excel beim Öffnen
        _validierung(ws, f, f"{bst}{erste}:{bst}{letzte}",
                     liste="obj_ID" if f.key == "objekt_id" else None)
        # Verkaufspreis als Annahme ist immer kritisch; § 6b und Reinvestition nur blau
        if f.annahme:
            _farblogik(ws, bst, erste, letzte, f)
        elif modell.schnellcheck and f.key in ("nutzung_6b", "reinvest"):
            _farblogik(ws, bst, erste, letzte, dataclasses.replace(f, annahme=True))

    ws.column_dimensions[st].width = 34
    _name(wb, VERKAUF_STATUS_NAME, f"'Verkäufe'!${st}${erste}:${st}${letzte}")
    pa = get_column_letter(status_spalte + 1)
    ws.cell(row=1, column=status_spalte + 1, value="Preis angenommen").font = FONT_KOPF
    _name(wb, "vk_PreisAnnahme", f"'Verkäufe'!${pa}${erste}:${pa}${letzte}")
    for zeile in range(erste, letzte + 1):
        for i, f in enumerate(VERKAUF_FELDER, start=1):
            c = ws.cell(row=zeile, column=i,
                        value=formeln.annahme_verkauf(f.key, zeile, modell.schnellcheck))
            c.number_format = f.format
            c.fill = FILL_EINGABE
        c = ws.cell(row=zeile, column=status_spalte + 1, value=formeln.preis_annahme(zeile))
        c.fill = FILL_BERECHNET
        rechnung = formeln.verkauf_zeile(zeile)
        for i, s in enumerate(VERKAUF_SPALTEN, start=len(VERKAUF_FELDER) + 1):
            c = ws.cell(row=zeile, column=i, value=rechnung[s.key])
            c.number_format = s.format
            c.fill = FILL_BERECHNET
        c = ws.cell(row=zeile, column=status_spalte, value=formeln.status_verkauf(zeile))
        c.fill = FILL_BERECHNET
    _status_rot(ws, f"{st}{erste}:{st}{letzte}", f"{st}{erste}")

    for zeile, vk in enumerate(modell.verkaeufe, start=erste):
        for i, f in enumerate(VERKAUF_FELDER, start=1):
            wert = getattr(vk, f.key)
            if wert is not None:
                ws.cell(row=zeile, column=i, value=wert)

    hinweis = get_column_letter(status_spalte + 2)
    ws[f"{hinweis}1"] = ("Verkauf zum Jahresende: Miete und AfA laufen im Verkaufsjahr noch, "
                         "ab dem Folgejahr ist das Objekt inaktiv.")
    ws[f"{hinweis}2"] = ("Aufteilung des Nettoerlöses: Anteil G+B laut Kaufvertrag, "
                         "sonst 1 − Verkehrswertanteil Gebäude aus dem Objektblatt.")
    ws[f"{hinweis}3"] = ("§ 6b Neubau begonnen = ja: Mit dem Bau wurde vor Ende der Regelfrist "
                         "begonnen, die Rücklage läuft dann par_6bFristNeubau statt par_6bFrist Jahre.")
    ws[f"{hinweis}4"] = "Kopf rot mit * = Pflichtfeld, blau = optional, grau = berechnet. " \
        + TOOLTIP_HINWEIS
    ws.freeze_panes = "B2"


def _blatt_neuobjekte(wb, modell: Modell) -> None:
    """Reinvestitionsobjekte: Eingaben, Übertragung der Rücklage, AfA-Basis."""
    ws = wb.create_sheet("Neuobjekte")
    erste, letzte = 2, modell.kapazitaet.neuobjekte + 1
    spalten = NEU_FELDER + NEU_SPALTEN
    status_spalte = len(spalten) + 1
    st = get_column_letter(status_spalte)
    _kopf(ws, 1, [s.ueberschrift for s in spalten] + [STATUS_UEBERSCHRIFT])
    _kopf_felder(ws, 1, NEU_FELDER)
    _kopf_berechnet(ws, 1, NEU_SPALTEN + [Spalte("status", "", NEU_STATUS_NAME, FMT_TEXT)],
                    start=len(NEU_FELDER) + 1)
    ws.row_dimensions[1].height = 45

    for i, s in enumerate(spalten, start=1):
        bst = get_column_letter(i)
        ws.column_dimensions[bst].width = s.breite
        _name(wb, s.name, f"'Neuobjekte'!${bst}${erste}:${bst}${letzte}")
    # Auswahl aus den gebildeten Rücklagen bzw. den Neukauf-Kostenstellen
    listen = {feld: "rl_ID" for feld, _, _ in QUELLEN}
    listen["kst"] = "nk_Liste"
    for i, f in enumerate(NEU_FELDER, start=1):
        bst = get_column_letter(i)
        if f.key in listen:
            _validierung(ws, dataclasses.replace(f, auswahl=None), f"{bst}{erste}:{bst}{letzte}",
                         liste=listen[f.key])
        else:
            _validierung(ws, f, f"{bst}{erste}:{bst}{letzte}")
        # Formeln aus „reinvestieren = ja“ im Blatt Verkäufe sind Annahmen (blau)
        _farblogik(ws, bst, erste, letzte, dataclasses.replace(f, pflicht=False, annahme=True))

    ws.column_dimensions[st].width = 44
    _name(wb, NEU_STATUS_NAME, f"'Neuobjekte'!${st}${erste}:${st}${letzte}")
    for zeile in range(erste, letzte + 1):
        for i, f in enumerate(NEU_FELDER, start=1):
            c = ws.cell(row=zeile, column=i, value=formeln.annahme_neu(f.key, zeile))
            c.number_format = f.format
            c.fill = FILL_EINGABE
        rechnung = formeln.neu_zeile(zeile)
        for i, s in enumerate(NEU_SPALTEN, start=len(NEU_FELDER) + 1):
            c = ws.cell(row=zeile, column=i, value=rechnung[s.key])
            c.number_format = s.format
            c.fill = FILL_BERECHNET
        c = ws.cell(row=zeile, column=status_spalte, value=formeln.status_neu(zeile))
        c.fill = FILL_BERECHNET
    _status_rot(ws, f"{st}{erste}:{st}{letzte}", f"{st}{erste}")

    for zeile, neu in enumerate(modell.neuobjekte, start=erste):
        # erfasstes Neuobjekt ersetzt die ganze Annahmezeile, leere Felder bleiben leer;
        # die Aufteilung der AK bleibt ohne Eingabe vorbelegt
        for i, f in enumerate(NEU_FELDER, start=1):
            wert = getattr(neu, f.key)
            if wert is None and f.key in ("ak_gub_eingabe", "ak_geb_eingabe"):
                continue
            ws.cell(row=zeile, column=i).value = wert  # None leert die Zelle

    hinweis = get_column_letter(status_spalte + 2)
    ws[f"{hinweis}1"] = ("Kauf zum Jahresende: Übertragung und Bestand im Kaufjahr, "
                         "Miete, Erhaltung und AfA ab dem Folgejahr.")
    ws[f"{hinweis}2"] = ("Übertragung: ü1 Gebäudegewinn auf Gebäude, ü2 G+B-Gewinn auf G+B, "
                         "ü3 Rest des G+B-Gewinns auf Gebäude. AfA-Basis = AK Gebäude − ü1 − ü3.")
    ws[f"{hinweis}3"] = ("Nutzen mehrere Neuobjekte dieselbe Rücklage, gilt die Zeilenreihenfolge: "
                         "jede Zeile erhält, was die Zeilen darüber übrig lassen.")
    ws[f"{hinweis}4"] = ("Kaufnebenkosten werden im Verhältnis G+B zu Gebäude aktiviert; leer = 0. "
                         "AK G+B und AK Gebäude sind aus dem Anteil G+B vorbelegt (blau); ein "
                         "eingetippter Betrag ersetzt die Vorbelegung, z. B. G+B fest und "
                         "Gebäude = Rest.")
    ws[f"{hinweis}6"] = ("Blau: Neuobjekt aus „reinvestieren = ja“ in derselben Zeile des Blatts "
                         "Verkäufe, Werte aus den Annahmen des Parameterblatts. Eintippen ersetzt "
                         "die Annahme.")
    ws[f"{hinweis}5"] = ("AfA-Methode degressiv: par_AfADegressiv vom Restbuchwert, Wechsel zur "
                         "linearen AfA über die Restnutzungsdauer (1 / AfA-Satz), sobald höher.")
    ws[f"{hinweis}7"] = ("Bis zu drei Quell-Rücklagen: erst alle Gebäudegewinne auf das Gebäude "
                         "(ü1), dann die G+B-Gewinne auf G+B (ü2), der Rest auf das Gebäude (ü3). "
                         "Die Werte je Quelle stehen in den eingeklappten Spalten (+ über Status).")
    ws[f"{hinweis}8"] = ("Finanzierungsbedarf = Kaufpreis + Nebenkosten − Nettoerlös der "
                         "Quell-Verkäufe (in Zeilenreihenfolge, soweit nicht schon eingesetzt). "
                         "Finanzierung Rest = Kredit: Kreditbetrag leer = ganzer Bedarf; der Rest "
                         "kommt aus Eigenmitteln (Liquidität). Tilgungsplan im Blatt Darlehen.")
    ws[f"{hinweis}9"] = "Kopf rot mit * = Pflichtfeld, blau = optional, grau = berechnet. " \
        + TOOLTIP_HINWEIS
    # Hilfsspalten je Quelle eingeklappt, Schalter über der Statusspalte
    von = len(NEU_FELDER) + 1 + [s.key for s in NEU_SPALTEN].index(NEU_DETAIL_ERSTE)
    ws.column_dimensions.group(get_column_letter(von), get_column_letter(status_spalte - 1),
                               hidden=True)
    ws.sheet_properties.outlinePr.summaryRight = True
    ws.freeze_panes = "B2"


def _blatt_darlehen(wb, modell: Modell) -> None:
    """Tilgungsplan je Kredit eines Neuobjekts: Zinsen, Tilgung und Restschuld je Jahr."""
    ws = wb.create_sheet("Darlehen")
    erste = DARLEHEN_ERSTE
    letzte = erste + modell.kapazitaet.neuobjekte - 1
    n = prognosejahre()
    ws["A1"] = ("Kredite der Neuobjekte (Blatt Neuobjekte, Finanzierung Rest = Kredit). "
                "Auszahlung zum Ende des Kaufjahrs; Zins und Tilgung ab dem Folgejahr auf die "
                "Restschuld am Vorjahresende. Annuität: gleiche Rate (Zins + anfängliche "
                "Tilgung); linear: gleiche Tilgung; endfällig: Rückzahlung nach der Laufzeit. "
                "Ist die Restschuld getilgt, entfallen Zins und Tilgung. Die Darlehen der "
                "Bestandsobjekte rechnet die Prognose je Objekt (Blatt Objekte).")
    ws["A1"].font = Font(italic=True)
    _kopf(ws, 3, [s.ueberschrift for s in DARLEHEN_SPALTEN])
    _kopf_berechnet(ws, 3, DARLEHEN_SPALTEN, grau=False)
    ws.row_dimensions[3].height = 45
    for i, s in enumerate(DARLEHEN_SPALTEN, start=1):
        bst = get_column_letter(i)
        ws.column_dimensions[bst].width = s.breite
        _name(wb, s.name, f"Darlehen!${bst}${erste}:${bst}${letzte}")
    for block, (key, text, name) in enumerate(DARLEHEN_BLOECKE):
        a, z = formeln.dl_block_spalte(block, 0), formeln.dl_block_spalte(block, n - 1)
        ws[f"{a}2"] = text
        ws[f"{a}2"].font = Font(bold=True)
        _name(wb, name, f"Darlehen!${a}${erste}:${z}${letzte}")
        for i in range(n):
            sp = formeln.dl_block_spalte(block, i)
            c = ws[f"{sp}3"]
            c.value, c.number_format = f"=par_Startjahr+{i}", FMT_JAHR
            c.font, c.fill = FONT_KOPF, FILL_KOPF
            ws.column_dimensions[sp].width = 11
            # Summe aller Kredite je Jahr unter der Tabelle
            c = ws[f"{sp}{letzte + 1}"]
            c.value, c.number_format = f"=SUM({sp}{erste}:{sp}{letzte})", FMT_EURO
            c.font = Font(bold=True)
        ws.column_dimensions[formeln.dl_block_spalte(block, -1)].width = 3
    ws.cell(row=letzte + 1, column=1, value="Summe").font = Font(bold=True)
    for neu_nr, zeile in enumerate(range(erste, letzte + 1), start=1):
        rechnung = formeln.darlehen_zeile(zeile, neu_nr)
        for i, s in enumerate(DARLEHEN_SPALTEN, start=1):
            c = ws.cell(row=zeile, column=i, value=rechnung[s.key])
            c.number_format, c.fill = s.format, FILL_BERECHNET
        for block, (key, _, _) in enumerate(DARLEHEN_BLOECKE):
            for i in range(n):
                c = ws[f"{formeln.dl_block_spalte(block, i)}{zeile}"]
                c.value, c.number_format = rechnung[(key, i)], FMT_EURO
                c.fill = FILL_BERECHNET
    ws.freeze_panes = f"B{erste}"


def _blatt_ruecklagen(wb, modell: Modell) -> None:
    """§ 6b-Rücklagen: links je Verkauf, rechts der Spiegel je Jahr."""
    ws = wb.create_sheet("Rücklagen")
    erste = 2

    letzte = modell.kapazitaet.verkaeufe + 1
    _kopf(ws, 1, [s.ueberschrift for s in RUECKLAGE_SPALTEN])
    _kopf_berechnet(ws, 1, RUECKLAGE_SPALTEN, grau=False)
    for i, s in enumerate(RUECKLAGE_SPALTEN, start=1):
        bst = get_column_letter(i)
        ws.column_dimensions[bst].width = s.breite
        _name(wb, s.name, f"'Rücklagen'!${bst}${erste}:${bst}${letzte}")
    for verkauf_nr, zeile in enumerate(range(erste, letzte + 1), start=1):
        rechnung = formeln.ruecklage_zeile(zeile, verkauf_nr)
        for i, s in enumerate(RUECKLAGE_SPALTEN, start=1):
            c = ws.cell(row=zeile, column=i, value=rechnung[s.key])
            c.number_format = s.format
            c.fill = FILL_BERECHNET

    versatz = len(RUECKLAGE_SPALTEN) + 1  # eine Spalte Abstand
    ws.column_dimensions[get_column_letter(versatz)].width = 3
    letzte = erste + prognosejahre() - 1
    for i, s in enumerate(RUECKLAGE_JAHR_SPALTEN, start=versatz + 1):
        bst = get_column_letter(i)
        c = ws.cell(row=1, column=i, value=s.ueberschrift)
        c.font, c.fill = FONT_KOPF, FILL_KOPF
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[bst].width = s.breite
        _name(wb, s.name, f"'Rücklagen'!${bst}${erste}:${bst}${letzte}")
    for zeile in range(erste, letzte + 1):
        rechnung = formeln.ruecklage_jahr_zeile(zeile, erstes_jahr=zeile == erste)
        for i, s in enumerate(RUECKLAGE_JAHR_SPALTEN, start=versatz + 1):
            c = ws.cell(row=zeile, column=i, value=rechnung[s.key])
            c.number_format = s.format
            c.fill = FILL_BERECHNET
    _kopf_berechnet(ws, 1, RUECKLAGE_JAHR_SPALTEN, start=versatz + 1, grau=False)
    ws.row_dimensions[1].height = 45

    hinweis = letzte + 2
    ws.cell(row=hinweis, column=versatz + 1, value=(
        "Rücklage nur bei § 6b ja und Status OK im Blatt Verkäufe, gebildet zum Ende des "
        "Verkaufsjahrs aus den positiven Teilgewinnen Gebäude und G+B."))
    ws.cell(row=hinweis + 1, column=versatz + 1, value=(
        "Was bis zum Fristjahr nicht auf Neuobjekte übertragen ist, wird dort aufgelöst, "
        "plus Gewinnzuschlag je vollem Jahr."))
    ws.cell(row=hinweis + 2, column=versatz + 1, value=(
        "steuerwirksam = Veräußerungsgewinne − Einstellung + Auflösung + Zuschlag. "
        + HINWEIS_FINANZIERUNG))
    ws.freeze_panes = "B2"


TITEL_PLAN = SZ_A.kurz
TITEL_BASELINE = SZ_BASELINE.kurz


def _jahrestabellen(wb, ws, blatt: str, tabellen: list, hinweise: list, texte: dict) -> None:
    """Tabellen je Prognosejahr nebeneinander, je eine Spalte Abstand.

    tabellen: (Titel, Spalten, Formelfunktion(zeile, erstes_jahr)). Zeile 1 trägt
    den Titel, Zeile 2 die Kopfzeile, ab Zeile 3 je Jahr eine Zeile.
    """
    titel_zeile, kopf, erste = 1, 2, 3
    letzte = erste + prognosejahre() - 1
    versatz = 0
    for titel, spalten, formel in tabellen:
        ws.cell(row=titel_zeile, column=versatz + 1, value=titel).font = Font(bold=True)
        for i, s in enumerate(spalten, start=versatz + 1):
            bst = get_column_letter(i)
            c = ws.cell(row=kopf, column=i, value=s.ueberschrift)
            c.font, c.fill = FONT_KOPF, FILL_KOPF
            c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[bst].width = s.breite
            _name(wb, s.name, f"'{blatt}'!${bst}${erste}:${bst}${letzte}")
        _kopf_berechnet(ws, kopf, spalten, start=versatz + 1, texte=texte, grau=False)
        for zeile in range(erste, letzte + 1):
            rechnung = formel(zeile, erstes_jahr=zeile == erste)
            for i, s in enumerate(spalten, start=versatz + 1):
                c = ws.cell(row=zeile, column=i, value=rechnung[s.key])
                c.number_format = s.format
                c.fill = FILL_BERECHNET
        versatz += len(spalten) + 1
        ws.column_dimensions[get_column_letter(versatz)].width = 3
    ws.row_dimensions[kopf].height = 45
    for i, text in enumerate(hinweise + [HINWEIS_FINANZIERUNG], start=letzte + 2):
        ws.cell(row=i, column=1, value=text)
    ws.freeze_panes = f"B{erste}"


def _blatt_liquiditaet(wb) -> None:
    """Steuer und Geldfluss je Jahr, je Szenario eine Tabelle."""
    ws = wb.create_sheet("Liquidität")
    _jahrestabellen(wb, ws, "Liquidität", [
        (sz.titel, liq_spalten(sz), formeln.liquiditaet_zeile(sz)) for sz in SZENARIEN
    ], [
        "Steuer = Bemessungsgrundlage × Grenzsteuersatz. Verluste werden vorgetragen und mit "
        "späteren Gewinnen verrechnet (ohne Mindestbesteuerung, ohne Rücktrag).",
        "steuerwirksam aus Verkauf und Rücklage: in A aus dem Rücklagenspiegel (sofort versteuerte "
        "Gewinne, Auflösung und Gewinnzuschlag), in B und C jeder Veräußerungsgewinn sofort.",
        "Zinsertrag = Liquidität am Vorjahresende × Rendite Alternativanlage; er ist steuerpflichtig. "
        "Negative Liquidität kostet denselben Satz.",
        "freier Mittelzufluss = Einnahmen − Ausgaben + Zins − Kreditzinsen + Verkaufserlöse − "
        "Steuer − Kauf Neuobjekte + Kreditauszahlung − Tilgung. Kreditzinsen mindern das "
        "Ergebnis, die Tilgung nicht.",
        "B: Neuobjekte mit Quelle-Rücklage entfallen, ihr Geld bleibt in der Alternativanlage. "
        "C: alle Neuobjekte werden gekauft, die AfA läuft von den vollen AK.",
        "Baseline: alle Bestandsobjekte werden über das ganze Raster gehalten, ohne Verkäufe und "
        "Neuobjekte.",
        TOOLTIP_HINWEIS,
    ], erklaerungen.LIQUIDITAET)


def _blatt_auswertung(wb) -> None:
    """Gesamt-GuV, Steuer, stille Reserven und Gesamtvermögen je Jahr und Szenario."""
    ws = wb.create_sheet("Auswertung")
    _jahrestabellen(wb, ws, "Auswertung", [
        (sz.titel, aus_spalten(sz), formeln.auswertung_zeile(sz)) for sz in SZENARIEN
    ], [
        "Gesamt-GuV vor Steuern = laufendes Ergebnis + steuerwirksam aus Verkauf und Rücklage "
        "+ Zinsertrag − Zinsen Kredite.",
        "stille Reserven = Verkehrswert − Buchwert (Gebäude + G+B) der Objekte im Bestand am "
        "Jahresende. In B und C ohne § 6b-Kürzung der Neuobjekte.",
        "latente Steuer = (stille Reserven + Rücklagenbestand − Verlustvortrag) × Grenzsteuersatz, "
        "mindestens 0: die Steuer, wenn alle Objekte zum Verkehrswert verkauft würden.",
        "Gesamtvermögen = Verkehrswert Bestand + Liquidität kumuliert − Restschuld Kredite.",
        TOOLTIP_HINWEIS,
    ], erklaerungen.AUSWERTUNG)


def _blatt_vergleich(wb) -> None:
    """Kennzahlen der Szenarien am Ende des Rasters und Endvermögen je Jahr."""
    ws = wb.create_sheet("Vergleich", 1)
    ws["A1"] = "Vergleich der Szenarien"
    ws["A1"].font = FONT_TITEL
    ws["A2"] = ("A § 6b-Kette wie erfasst · B sofort versteuern, Kapital anlegen · C sofort "
                "versteuern, Neuobjekte trotzdem kaufen · Baseline alles halten.")
    ws["A3"] = ("Entscheidend ist das Endvermögen nach latenter Steuer. A − C zeigt die reine "
                "Wirkung von § 6b (Zins auf die gestundete Steuer), A − B die Frage Immobilie oder "
                "Geldanlage. " + HINWEIS_FINANZIERUNG)

    kopf = 5
    spalten = ([("Kennzahl", "vg_Kennzahl", 38)]
               + [(sz.kurz, f"vg_{sz.key}", 18) for sz in SZENARIEN]
               + [("A − B", "vg_DiffB", 16), ("A − C", "vg_DiffC", 16),
                  ("A − Baseline", "vg_DiffBaseline", 16), ("Erläuterung", None, 60)])
    _kopf(ws, kopf, [u for u, _, _ in spalten])
    ws.row_dimensions[kopf].height = 32
    erste = kopf + 1
    letzte = erste + len(VERGLEICH_KENNZAHLEN) - 1
    spalte_von = {sz.key: get_column_letter(2 + i) for i, sz in enumerate(SZENARIEN)}
    for zeile, (text, blatt, name, art, erlaeuterung) in enumerate(VERGLEICH_KENNZAHLEN,
                                                                   start=erste):
        werte = [text] + [formeln.vergleich_zeile(blatt, name, art, sz) for sz in SZENARIEN]
        a = spalte_von["A"]
        werte += [f"={a}{zeile}-{spalte_von[k]}{zeile}" for k in ("B", "C", "Baseline")]
        werte.append(erlaeuterung)
        for i, w in enumerate(werte, start=1):
            c = ws.cell(row=zeile, column=i, value=w)
            if 1 < i < len(werte):
                c.number_format = FMT_EURO
                c.fill = FILL_BERECHNET
        if zeile == erste:
            for i in range(1, len(werte)):
                ws.cell(row=zeile, column=i).font = Font(bold=True)
    for i, (_, name, breite) in enumerate(spalten, start=1):
        bst = get_column_letter(i)
        ws.column_dimensions[bst].width = breite
        if name:
            _name(wb, name, f"'Vergleich'!${bst}${erste}:${bst}${letzte}")

    # Endvermögen je Jahr und Szenario, mit Diagramm
    titel = letzte + 2
    ws.cell(row=titel, column=1,
            value="Endvermögen nach latenter Steuer je Jahr").font = Font(bold=True)
    jkopf = titel + 1
    _kopf(ws, jkopf, ["Jahr"] + [sz.kurz for sz in SZENARIEN])
    ws.row_dimensions[jkopf].height = 32
    jerste, jletzte = jkopf + 1, jkopf + prognosejahre()
    for zeile in range(jerste, jletzte + 1):
        c = ws.cell(row=zeile, column=1,
                    value="=par_Startjahr" if zeile == jerste else f"=A{zeile - 1}+1")
        c.number_format, c.fill = FMT_JAHR, FILL_BERECHNET
        for i, sz in enumerate(SZENARIEN, start=2):
            c = ws.cell(row=zeile, column=i,
                        value=f"=SUMIFS({sz.aus}_VermoegenNetto,{sz.aus}_Jahr,A{zeile})")
            c.number_format, c.fill = FMT_EURO, FILL_BERECHNET
    _name(wb, "vgj_Jahr", f"'Vergleich'!$A${jerste}:$A${jletzte}")
    for i, sz in enumerate(SZENARIEN, start=2):
        bst = get_column_letter(i)
        _name(wb, f"vgj_{sz.key}", f"'Vergleich'!${bst}${jerste}:${bst}${jletzte}")

    chart = _linien("Endvermögen nach latenter Steuer")
    chart.add_data(Reference(ws, min_col=2, max_col=1 + len(SZENARIEN), min_row=jkopf,
                             max_row=jletzte), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=jerste, max_row=jletzte))
    farben = {"A": "305496", "B": "C00000", "C": "70AD47", "Baseline": "A5A5A5"}
    for serie, sz in zip(chart.series, SZENARIEN):
        serie.smooth = False
        serie.graphicalProperties.line.solidFill = farben[sz.key]
        serie.graphicalProperties.line.width = 28000
        if sz == SZ_BASELINE:
            serie.graphicalProperties.line.dashStyle = "dash"
    ws.add_chart(chart, f"H{titel}")
    ws.freeze_panes = f"B{erste}"


def _blatt_uebersicht(wb) -> None:
    """Immobilienwert und Gesamtvermögen je Jahr: Plan mit Verkäufen gegen Halten ohne Verkauf."""
    ws = wb.create_sheet("Übersicht", 0)
    wb.active = 0
    ws["A1"] = "Übersicht: Immobilienwert und Gesamtvermögen"
    ws["A1"].font = FONT_TITEL
    ws["A2"] = (HINWEIS_FINANZIERUNG + " Wert = Verkehrswert aktuell aus dem Objektblatt, "
                "fortgeschrieben mit der Wertsteigerung vom Parameterblatt.")
    ws["A3"] = ("Gesamtvermögen = Verkehrswert + kumulierte Liquidität nach Steuern − latente "
                "Steuer auf stille Reserven und Rücklage (Blatt Auswertung).")

    kennzahlen = [
        ("Objekte im Modell", "ueb_Objekte", '=SUMPRODUCT(--(obj_ID<>""))', FMT_ZAHL),
        ("davon mit Annahmen (blau)", "ueb_MitAnnahmen", '=COUNTIF(obj_Annahmen,">0")',
         FMT_ZAHL),
        ("geplante Verkäufe", "ueb_Verkaeufe", '=SUMPRODUCT(--(vk_ID<>""))', FMT_ZAHL),
        ("Neuobjekte im Modell", "ueb_Neuobjekte", "=SUM(ne_Gueltig)", FMT_ZAHL),
        ("Wertsteigerung p. a.", None, "=par_Wertsteig", FMT_PROZENT),
        ("Plausibilitätsprüfung (Blatt Prüfung)", "ueb_Pruefung", "=pr_Gesamt", FMT_TEXT),
    ]
    for zeile, (text, name, formel, fmt) in enumerate(kennzahlen, start=4):
        ws.cell(row=zeile, column=1, value=text)  # läuft über die leeren Spalten B und C
        c = ws.cell(row=zeile, column=4, value=formel)
        c.fill = FILL_BERECHNET
        c.number_format = fmt
        if name:
            _name(wb, name, f"'Übersicht'!$D${zeile}")
    ws.conditional_formatting.add(
        "D5", FormulaRule(formula=["D5>0"], fill=FILL_ANNAHME))
    ws.conditional_formatting.add(
        "D9", FormulaRule(formula=['D9<>"OK"'], fill=FILL_FEHLER))

    kopf = 11
    ws.cell(row=kopf - 1, column=2, value="Immobilien (Verkehrswert)").font = Font(bold=True)
    ws.cell(row=kopf - 1, column=5,
            value="Gesamtvermögen nach latenter Steuer").font = Font(bold=True)
    _kopf(ws, kopf, ["Jahr", TITEL_BASELINE, TITEL_PLAN, "Differenz Plan − Baseline",
                     TITEL_BASELINE, TITEL_PLAN, "Differenz Plan − Baseline"])
    ws.row_dimensions[kopf].height = 32
    erste = kopf + 1
    letzte = erste + prognosejahre()  # Basisjahr plus Prognosejahre
    for zeile in range(erste, letzte + 1):
        if zeile == erste:
            # Ende Basisjahr: Stand laut Objektblatt, noch keine Liquidität
            buchwert = 'SUMIFS(obj_Restbuchwert,obj_ID,"<>")+SUMIFS(obj_AKGuB,obj_ID,"<>")'
            werte = ["=par_Basisjahr", "=SUM(obj_Verkehrswert)", f"=B{zeile}",
                     f"=C{zeile}-B{zeile}",
                     f"=B{zeile}-MAX(B{zeile}-({buchwert}),0)*par_Steuersatz", f"=E{zeile}"]
        else:
            werte = [f"=A{zeile - 1}+1",
                     f"=SUMIFS(prg_Verkehrswert,prg_Jahr,A{zeile},prg_Neu,0)",
                     f"=SUMIFS(prg_Verkehrswert,prg_Jahr,A{zeile},prg_Bestand,1)",
                     f"=C{zeile}-B{zeile}",
                     f"=SUMIFS(asb_VermoegenNetto,asb_Jahr,A{zeile})",
                     f"=SUMIFS(aus_VermoegenNetto,aus_Jahr,A{zeile})"]
        werte.append(f"=F{zeile}-E{zeile}")
        for spalte, formel in enumerate(werte, start=1):
            c = ws.cell(row=zeile, column=spalte, value=formel)
            c.number_format = FMT_JAHR if spalte == 1 else FMT_EURO
            c.fill = FILL_BERECHNET
    for name, bst in (("ueb_Jahr", "A"), ("ueb_Baseline", "B"), ("ueb_Plan", "C"),
                      ("ueb_Differenz", "D"), ("ueb_VermBaseline", "E"), ("ueb_VermPlan", "F"),
                      ("ueb_VermDifferenz", "G")):
        _name(wb, name, f"'Übersicht'!${bst}${erste}:${bst}${letzte}")

    for spalte, breite in zip("ABCDEFG", (10, 20, 20, 20, 20, 20, 20)):
        ws.column_dimensions[spalte].width = breite
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    for i, (titel, spalte_baseline) in enumerate(
            (("Verkehrswert am Jahresende", 2),
             ("Gesamtvermögen nach latenter Steuer", 5))):
        ws.add_chart(_diagramm(ws, titel, spalte_baseline, kopf, erste, letzte), f"I{4 + i * 26}")


def _diagramm(ws, titel: str, spalte_baseline: int, kopf: int, erste: int, letzte: int):
    """Liniendiagramm Baseline (gestrichelt) gegen Plan aus zwei benachbarten Spalten."""
    chart = _linien(titel)
    chart.add_data(Reference(ws, min_col=spalte_baseline, max_col=spalte_baseline + 1,
                             min_row=kopf, max_row=letzte), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=erste, max_row=letzte))
    baseline, plan = chart.series
    baseline.graphicalProperties.line.solidFill = "A5A5A5"
    baseline.graphicalProperties.line.dashStyle = "dash"
    plan.graphicalProperties.line.solidFill = "305496"
    for serie in chart.series:
        serie.smooth = False
        serie.graphicalProperties.line.width = 28000
    return chart


def _schrift(groesse: int, drehung: int = None) -> RichText:
    """Achsenschrift in Punkt × 100, optional gedreht (Grad × 60.000, negativ = schräg)."""
    zeichen = CharacterProperties(sz=groesse)
    return RichText(bodyPr=RichTextProperties(rot=drehung, vert="horz"),
                    p=[Paragraph(pPr=ParagraphProperties(defRPr=zeichen),
                                 endParaRPr=zeichen)])


def _linien(titel: str) -> LineChart:
    """Liniendiagramm mit Achsen, die sich in Excel und LibreOffice nicht überlagern.

    Beträge in Tsd. € (Einheit im Titel statt Achsentitel), Jahre schräg und immer am
    unteren Rand, auch bei negativen Werten; Titel und Legende außerhalb der Zeichenfläche.
    """
    chart = LineChart()
    chart.title = f"{titel} (Tsd. €)"
    chart.title.overlay = False
    chart.y_axis.number_format = '#,##0,'
    chart.y_axis.majorGridlines = None
    chart.y_axis.txPr = _schrift(900)
    chart.x_axis.number_format = "0"
    chart.x_axis.tickLblPos = "low"
    chart.x_axis.txPr = _schrift(800, -2700000)
    # openpyxl blendet die Achsen sonst in neueren Excel-Versionen aus
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.legend.position = "b"
    chart.legend.overlay = False
    chart.height, chart.width = 12, 22
    return chart


def _blatt_pruefung(wb) -> None:
    """Plausibilitätsprüfungen als Formeln: je Prüfung Anzahl betroffener Zeilen und Ergebnis.

    Das Makro modPruefung liest nur das Ergebnis; geprüft wird im Blatt, damit Fehler
    auch ohne Makros sichtbar bleiben.
    """
    ws = wb.create_sheet("Prüfung")
    ws["A1"] = "Plausibilitätsprüfung"
    ws["A1"].font = FONT_TITEL
    ws["A2"] = HINWEIS_FINANZIERUNG
    summen = formeln.pruefung_summen()
    for zeile, (name, text) in enumerate([("pr_Gesamt", "Gesamtergebnis"),
                                          ("pr_Fehler", "Fehler"),
                                          ("pr_Warnungen", "Warnungen")], start=3):
        ws.cell(row=zeile, column=1, value=text).font = Font(bold=True)
        c = ws.cell(row=zeile, column=2, value=summen[name])
        c.fill = FILL_BERECHNET
        c.number_format = FMT_TEXT if name == "pr_Gesamt" else FMT_ZAHL
        _name(wb, name, f"'Prüfung'!$B${zeile}")
    ws.conditional_formatting.add("B3", FormulaRule(formula=['B3<>"OK"'], fill=FILL_FEHLER))

    _kopf(ws, 7, ["Nr", "Prüfung", "Art", "betroffen", "Ergebnis", "wo nachsehen"])
    erste = 8
    letzte = erste + len(PRUEFUNGEN) - 1
    anzahl = formeln.pruefung_anzahl()
    for nr, (zeile, p) in enumerate(zip(range(erste, letzte + 1), PRUEFUNGEN), start=1):
        ws.cell(row=zeile, column=1, value=nr)
        ws.cell(row=zeile, column=2, value=p.bezeichnung).alignment = Alignment(wrap_text=True)
        ws.cell(row=zeile, column=3, value=p.art)
        for spalte, formel in ((4, anzahl[p.key]), (5, formeln.pruefung_ergebnis(zeile))):
            c = ws.cell(row=zeile, column=spalte, value=formel)
            c.fill = FILL_BERECHNET
        ws.cell(row=zeile, column=6, value=p.wo).alignment = Alignment(wrap_text=True)
    for name, bst in (("pr_Bezeichnung", "B"), ("pr_Art", "C"), ("pr_Anzahl", "D"),
                      ("pr_Ergebnis", "E")):
        _name(wb, name, f"'Prüfung'!${bst}${erste}:${bst}${letzte}")
    for art, fill in ((FEHLER, FILL_FEHLER), (WARNUNG, FILL_WARNUNG)):
        ws.conditional_formatting.add(
            f"E{erste}:E{letzte}", FormulaRule(formula=[f'E{erste}="{art}"'], fill=fill))

    hinweise = [
        f"{FEHLER}: Eingabe unvollständig oder unzulässig; die Zeile rechnet nicht oder nur "
        f"teilweise mit. {WARNUNG}: rechnet, ist aber steuerlich ungünstig oder fachlich zu "
        f"prüfen. {HINWEIS}: zur Kenntnis, zählt nicht ins Gesamtergebnis.",
        "Das Gesamtergebnis steht auch auf dem Parameterblatt und in der Übersicht. Das Makro "
        "„Plausibilität prüfen“ rechnet neu und listet die auffälligen Prüfungen auf.",
        "Je Objekt, Verkauf und Neuobjekt steht der Grund in der Statusspalte des Blatts. "
        "AK Gebäude und AK G+B werden getrennt erfasst, ihre Summe ist der Kaufpreis; beim "
        "Neuobjekt ergibt der Anteil G+B mit dem Rest Gebäude immer den Kaufpreis.",
    ]
    for i, text in enumerate(hinweise, start=letzte + 2):
        ws.cell(row=i, column=1, value=text)
    for spalte, breite in zip("ABCDEF", (16, 70, 10, 10, 10, 50)):
        ws.column_dimensions[spalte].width = breite
    ws.freeze_panes = "A8"


def _blatt_varianten(wb) -> None:
    """Festgehaltene Ergebnisse: das Makro „Variante festhalten“ schreibt je Lauf eine Zeile."""
    ws = wb.create_sheet("Varianten")
    ws["A1"] = "Varianten: festgehaltene Ergebnisse"
    ws["A1"].font = FONT_TITEL
    ws["A2"] = (HINWEIS_FINANZIERUNG + " Feste Werte, keine Formeln: Eingaben ändern (etwa "
                "Verkaufsjahr oder Preis), dann erneut festhalten und die Zeilen vergleichen.")
    kopf = 4
    _kopf(ws, kopf, VARIANTEN_KOPF)
    ws.row_dimensions[kopf].height = 32
    erste, letzte = kopf + 1, kopf + MAX_VARIANTEN
    _name(wb, "var_Tabelle",
          f"Varianten!$A${erste}:${get_column_letter(len(VARIANTEN_KOPF))}${letzte}")
    for name, bst in zip(("var_Bezeichnung", "var_Datum", "var_A", "var_B", "var_C",
                          "var_Baseline", "var_DiffB", "var_DiffC", "var_Steuer",
                          "var_Pruefung"), "ABCDEFGHIJ"):
        _name(wb, name, f"Varianten!${bst}${erste}:${bst}${letzte}")
    for zeile in range(erste, letzte + 1):
        ws.cell(row=zeile, column=2).number_format = "DD.MM.YYYY HH:MM"
        for spalte in range(3, len(VARIANTEN_KOPF)):
            ws.cell(row=zeile, column=spalte).number_format = FMT_EURO
    for spalte, breite in zip("ABCDEFGHIJ", (28, 16, 16, 16, 16, 18, 14, 14, 16, 16)):
        ws.column_dimensions[spalte].width = breite
    ws.freeze_panes = f"B{erste}"


# Optionen des Startblatts: Text und Szenario (Kennzahlen im Blatt Vergleich)
OPTIONEN = [
    ("Halten (nichts verkaufen)", "vg_Baseline"),
    ("Verkaufen und mit § 6b-Rücklage reinvestieren (Plan wie erfasst)", "vg_A"),
    ("Verkaufen, sofort versteuern, trotzdem reinvestieren", "vg_C"),
    ("Verkaufen, sofort versteuern, Erlös anlegen", "vg_B"),
]
SCHRITTE = [
    ("Parameter", "Kernparameter (★) prüfen: Steuersatz, Steigerungen, Alternativrendite und "
     "die Annahmen bei fehlenden Daten (blau); weitere Abschnitte mit + aufklappen."),
    ("Objekte", "Je Objekt mindestens ObjektID und Miete (Pflicht, rot wenn leer). Alles "
     "Weitere füllen die Annahmen; echte Werte einfach darübertippen."),
    ("Anlagen", "Anlagenverzeichnis: AK, Buchwert und AfA je Anlage ersetzen die Annahmen zu "
     "AK, Kaufjahr und Restbuchwert (grün). Anlagen ohne Objekt dort zuordnen."),
    ("Verkäufe", "Geplanten Verkauf erfassen: Objekt, Verkaufsjahr; Preis leer = Verkehrswert; "
     "§ 6b nutzen und reinvestieren mit ja/nein."),
    ("Neuobjekte", "Reinvestitionen: entstehen bei „reinvestieren = ja“ automatisch (blau), "
     "oder hier händisch erfassen."),
    ("Prüfung", "Plausibilitätsprüfung: Fehler beheben, Warnungen (orange Annahmen) prüfen."),
    ("Vergleich", "Ergebnis: Szenarien über 20 Jahre; Details in Verkauf und Kauf."),
]
# Blattreiter: gelb Eingabe, grau Rechnung, blau Ausgabe, grün Kontrolle
REITER = {"Start": "305496", "Parameter": "FFC000", "Objekte": "FFC000", "Anlagen": "FFC000",
          "AfA-Plan": "FFC000", "Neukauf-KSt": "FFC000",
          "Verkäufe": "FFC000",
          "Neuobjekte": "FFC000", "Prognose": "A5A5A5", "Rücklagen": "A5A5A5",
          "Liquidität": "A5A5A5", "Darlehen": "A5A5A5", "Prüfung": "70AD47", "Varianten": "70AD47",
          "BWA-Zuordnung": "FFC000"}
FARBE_AUSGABE = "5B9BD5"
# im Schnellcheck ausgeblendet; über Rechtsklick auf einen Reiter wieder einblendbar
SCHNELL_AUSGEBLENDET = ("AfA-Plan", "Neukauf-KSt", "Darlehen", "Prognose", "Rücklagen", "Liquidität",
                        "Auswertung")


def _link(zelle, blatt: str) -> None:
    from openpyxl.worksheet.hyperlink import Hyperlink
    zelle.hyperlink = Hyperlink(ref=zelle.coordinate, location=f"'{blatt}'!A1",
                                display=str(zelle.value))
    zelle.font = Font(color="0563C1", underline="single")


def _blatt_start(wb, modell: Modell) -> None:
    """Startblatt: Handlungsempfehlung, Datenlage, Anleitung und Farblegende."""
    ws = wb.create_sheet("Start", 0)
    ws["A1"] = "Prognosemodell V+V" + (" – Schnellcheck" if modell.schnellcheck else "")
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = (HINWEIS_FINANZIERUNG + " Steuersätze und Fristen vor dem Echteinsatz mit dem "
                "zuständigen Berufsträger prüfen.")
    ws["A2"].font = Font(italic=True)

    ws["A4"] = "Handlungsempfehlung"
    ws["A4"].font = Font(bold=True, size=13)
    erste = 10
    letzte = erste + len(OPTIONEN) - 1
    werte, namen = f"$B${erste}:$B${letzte}", f"$A${erste}:$A${letzte}"
    beste = f"INDEX({namen},MATCH(MAX({werte}),{werte},0))"
    ws["A5"] = (f'=IF(ueb_Verkaeufe=0,"Noch kein Verkauf erfasst: im Blatt Verkäufe Objekt und '
                f'Verkaufsjahr eintragen, dann zeigt dieses Blatt die beste Option.",'
                f'IF(MAX($B${erste + 1}:$B${letzte})<=$B${erste}+1,"Halten: kein '
                f'Verkaufsszenario erreicht ein höheres Endvermögen.","Empfehlung: "&{beste}))')
    ws["A5"].font = Font(bold=True, size=12, color="305496")
    _name(wb, "start_Empfehlung", "Start!$A$5")
    ws["A6"] = "Vorsprung der besten Option gegenüber Halten"
    ws["D6"] = f"=MAX({werte})-$B${erste}"
    ws["D6"].number_format = FMT_EURO
    _name(wb, "start_Vorsprung", "Start!$D$6")
    kritisch = 'COUNTIF(obj_Kritisch,">0")+SUM(vk_PreisAnnahme)'
    ws["A7"] = (f'=IF(pr_Fehler>0,"Nicht belastbar: "&pr_Fehler&" Fehler, siehe Blatt Prüfung.",'
                f'IF({kritisch}>0,"Vorläufig: "&({kritisch})&" kritische Annahme(n) bei '
                f'Verkäufen (orange) durch echte Werte ersetzen.",'
                f'"Belastbar im Rahmen der zentralen Annahmen."))')
    _name(wb, "start_Belastbarkeit", "Start!$A$7")

    _kopf_start = ["Option", "Endvermögen nach latenter Steuer", "Differenz zu Halten", "Rang"]
    for spalte, text in enumerate(_kopf_start, start=1):
        c = ws.cell(row=erste - 1, column=spalte, value=text)
        c.font, c.fill = FONT_KOPF, FILL_KOPF
        c.alignment = Alignment(wrap_text=True)
    ws.cell(row=erste - 2, column=1, value="=\"Endvermögen am Ende \"&par_Endjahr&\" je Option\"")
    for zeile, (text, name) in zip(range(erste, letzte + 1), OPTIONEN):
        ws.cell(row=zeile, column=1, value=text)
        for spalte, formel, fmt in ((2, f"=INDEX({name},1)", FMT_EURO),
                                    (3, f"=B{zeile}-$B${erste}", FMT_EURO),
                                    (4, f'=COUNTIF({werte},">"&B{zeile})+1', FMT_ZAHL)):
            c = ws.cell(row=zeile, column=spalte, value=formel)
            c.number_format, c.fill = fmt, FILL_BERECHNET
    ws.conditional_formatting.add(
        f"A{erste}:D{letzte}", FormulaRule(formula=[f"$D{erste}=1"], fill=FILL_EINGELESEN))
    _name(wb, "start_Optionen", f"Start!{namen}")
    _name(wb, "start_Werte", f"Start!{werte}")

    zeile = letzte + 2
    kennzahlen = [
        ("Wert der § 6b-Kette (A − C): Zins auf die gestundete Steuer", "=INDEX(vg_DiffC,1)",
         FMT_EURO),
        ("Steuer gesamt im Plan über alle Jahre", "=SUM(liq_Steuer)", FMT_EURO),
        ("tiefster Liquiditätsstand im Plan (negativ = Finanzierungsbedarf)", "=MIN(liq_Kum)",
         FMT_EURO),
        ("im Jahr", "=INDEX(liq_Jahr,MATCH(MIN(liq_Kum),liq_Kum,0))", FMT_JAHR),
    ]
    for text, formel, fmt in kennzahlen:
        ws.cell(row=zeile, column=1, value=text)
        c = ws.cell(row=zeile, column=4, value=formel)
        c.number_format, c.fill = fmt, FILL_BERECHNET
        zeile += 1

    zeile += 1
    ws.cell(row=zeile, column=1, value="Datenlage").font = Font(bold=True, size=13)
    zeile += 1
    for text, formel, fmt, name in [
        ("Objekte im Modell", "=ueb_Objekte", FMT_ZAHL, None),
        ("davon mit Annahmen (blau)", '=COUNTIF(obj_Annahmen,">0")', FMT_ZAHL, None),
        ("davon mit Anlagen aus dem Anlagenverzeichnis", '=COUNTIF(obj_AnlAbn,">0")', FMT_ZAHL,
         None),
        ("Anlagen ohne Objekt oder unvollständig (Blatt Anlagen)",
         formeln.pruefung_anzahl()["anlagen"], FMT_ZAHL, None),
        ("kritische Annahmen bei Verkäufen (orange)", f"={kritisch}", FMT_ZAHL,
         "start_Kritisch"),
        ("Plausibilitätsprüfung", "=pr_Gesamt", FMT_TEXT, None),
    ]:
        ws.cell(row=zeile, column=1, value=text)
        c = ws.cell(row=zeile, column=4, value=formel)
        c.number_format, c.fill = fmt, FILL_BERECHNET
        if name:
            _name(wb, name, f"Start!$D${zeile}")
        zeile += 1

    zeile += 1
    ws.cell(row=zeile, column=1, value="So geht's").font = Font(bold=True, size=13)
    zeile += 1
    for nr, (blatt, text) in enumerate(SCHRITTE, start=1):
        _link(ws.cell(row=zeile, column=1, value=f"{nr}. {blatt}"), blatt)
        ws.cell(row=zeile, column=2, value=text)
        zeile += 1

    zeile += 1
    ws.cell(row=zeile, column=1, value="Farben der Eingabezellen").font = Font(bold=True, size=13)
    zeile += 1
    for fill, farbe, text in FARBEN:
        c = ws.cell(row=zeile, column=1, value=farbe)
        c.fill = fill
        ws.cell(row=zeile, column=2, value=text)
        zeile += 1
    for fill, farbe, text in ((FILL_BERECHNET, "grau", "Formel, nicht überschreiben"),):
        ws.cell(row=zeile, column=1, value=farbe).fill = fill
        ws.cell(row=zeile, column=2, value=text)
        zeile += 1
    for fill, farbe, text in KOPF_LEGENDE:
        c = ws.cell(row=zeile, column=1, value=farbe)
        c.fill, c.font = fill, FONT_KOPF
        ws.cell(row=zeile, column=2, value=text)
        zeile += 1
    ws.cell(row=zeile, column=1, value=TOOLTIP_HINWEIS).font = Font(italic=True)
    zeile += 1

    zeile += 1
    ws.cell(row=zeile, column=1, value="Zentrale Annahmen").font = Font(bold=True, size=13)
    _link(ws.cell(row=zeile, column=2, value="ändern im Blatt Parameter"), "Parameter")
    zeile += 1
    for p in PARAMETER:
        if p.name.startswith(("par_Ann", "par_ErhAlterung", "par_NeuErh", "par_San")) \
                or p.name in ("par_Steuersatz", "par_Mietsteig", "par_Erhaltsteig",
                                                      "par_Wertsteig", "par_Alternativrendite"):
            kern = p.name in erklaerungen.KERNPARAMETER
            a = ws.cell(row=zeile, column=1, value=("★ " if kern else "") + p.bezeichnung)
            if kern:
                a.font = Font(bold=True)
            c = ws.cell(row=zeile, column=4, value=f"={p.name}")
            c.number_format, c.fill = p.format, FILL_BERECHNET
            zeile += 1

    for spalte, breite in zip("ABCD", (62, 22, 20, 16)):
        ws.column_dimensions[spalte].width = breite


def erstelle_mappe(modell: Modell) -> Workbook:
    k = modell.kapazitaet
    for art, anzahl, platz in (("Objekte", len(modell.objekte), k.objekte),
                               ("Verkäufe", len(modell.verkaeufe), k.verkaeufe),
                               ("Neuobjekte", len(modell.neuobjekte), k.neuobjekte)):
        if anzahl > platz:
            raise ValueError(f"{anzahl} {art}, die Mappe hat Platz für {platz}.")
    wb = Workbook()
    _blatt_parameter(wb, modell)
    _blatt_objekte(wb, modell)
    _blatt_anlagen(wb, modell)
    _blatt_afa_plan(wb, modell)
    _blatt_neukauf(wb, modell)
    _blatt_verkaeufe(wb, modell)
    _blatt_neuobjekte(wb, modell)
    _blatt_darlehen(wb, modell)
    _blatt_prognose(wb, modell)
    _blatt_ruecklagen(wb, modell)
    _blatt_liquiditaet(wb)
    _blatt_auswertung(wb)
    _blatt_uebersicht(wb)
    _blatt_vergleich(wb)
    _blatt_pruefung(wb)
    _blatt_varianten(wb)
    bwa_blaetter = bwa.blaetter_bwa(wb, modell)
    bwa.blatt_sonderbereich(wb, modell)
    _blatt_start(wb, modell)
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = REITER.get(ws.title, FARBE_AUSGABE)
        if modell.schnellcheck and (ws.title in SCHNELL_AUSGEBLENDET or ws.title in bwa_blaetter):
            ws.sheet_state = "hidden"
    wb.code_name = CODENAME_MAPPE
    for i, ws in enumerate(wb.worksheets, start=1):
        # BWA-Blätter heißen nach der Kostenstelle; Codename für VBA dann wsBWA<n>
        ws.sheet_properties.codeName = CODENAMEN.get(ws.title, f"wsBWA{i}")
    return wb
