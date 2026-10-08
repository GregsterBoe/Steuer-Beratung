"""Ausgabe im DATEV-BWA-Format und Sonderbereich Verkauf und Kauf (Projektplan Abschnitt 18).

Je Kostenstelle ein Blatt im Layout der Kostenstellenblätter (vorlagen.bwa_kopf):
links die Ist-Werte, rechts die Planjahre als Formeln aus der Prognose (Szenario A).
Das Summenblatt „Alle Objekte“ rechnet über alle Objekte, auch über solche, die
erst in Excel angelegt werden und kein eigenes Blatt haben, und stimmt mit dem
Blatt Liquidität überein: Ergebnis vor Steuern = Ergebnis vor Verlustvortrag.

Das Blatt „Verkauf und Kauf“ zeigt je Verkauf eine Ergebnissicht und eine
Detailsicht, darunter je Neuobjekt die Detailsicht des Kaufs.
"""

import re

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .modelle import (ARTEN_ABNUTZBAR, FMT_EURO, FMT_JAHR, FMT_PROZENT, QUELLEN,
                      PARAMETER, STATUS_6B_UNZULAESSIG, STATUS_OK, Modell,
                      prognosejahre)
from .vorlagen import (BWA_ZEILEN, ERSTE_ZEILE, SPALTE_JAHR, SPALTE_PLAN, SPALTE_VORJAHRE,
                       bwa_kopf)

SUMMENBLATT = "Alle Objekte"
SONDERBEREICH = "Verkauf und Kauf"
ZEILE_JAHR = ERSTE_ZEILE - 1   # Hilfszeile mit dem Planjahr als Zahl, für die Formeln
ZEILEN = {nr: ERSTE_ZEILE + i for i, (nr, _) in enumerate(BWA_ZEILEN)}   # wie bwa_kopf

# Kostenarten aus „weitere Ausgaben“; 1260 Sonstige Kosten nimmt den Rest auf
AUSGABEN_ZEILEN = (1100, 1120, 1140, 1150, 1160, 1180, 1200, 1220)
KOSTENARTEN = AUSGABEN_ZEILEN + (1240, 1250, 1260)
# Kostenstellenblatt führend: Spalte Basisjahr ist Eingabe, das Blatt Objekte verweist darauf
# (Feld -> BWA-Zeilen). Ohne eingelesene Kostenstelle umgekehrt: Basisjahr aus dem
# Objektblatt (IST_AUS_OBJEKT), die Kostenarten 1100–1220 bleiben leer, 1260 nimmt den Rest auf
OBJEKT_AUS_KST = {"miete": (1020,), "weitere_einnahmen": (1090,), "erhaltung": (1250,),
                  "afa_bwa": (1240,), "weitere_ausgaben": (1260,) + AUSGABEN_ZEILEN,
                  "zinsen": (1310,)}
IST_AUS_OBJEKT = {1020: "obj_MieteBasis", 1090: "obj_EinnBasis", 1240: "obj_AfABWA",
                  1250: "obj_ErhBasis", 1310: "obj_ZinsBasis"}
AUSGABEN_AUS_OBJEKT = "obj_AusgBasis"
# mindestens so viele Vorgänge und Neuobjekte im Sonderbereich, damit in Excel ergänzte passen
MIN_VORGAENGE = 3

FMT_BWA = '#,##0.00;-#,##0.00;""'   # 0 bleibt leer, etwa nach dem Verkauf
FONT_TITEL = Font(bold=True, size=14)
FONT_ABSCHNITT = Font(bold=True, size=12)
FONT_HILFE = Font(italic=True, color="808080", size=8)
FONT_HILFE_GROSS = Font(italic=True, color="595959")
FILL_BERECHNET = PatternFill("solid", fgColor="E7E6E6")
FILL_KOPF = PatternFill("solid", fgColor="D9E1F2")
FILL_EINGABE = PatternFill("solid", fgColor="FFF2CC")
FILL_GEAENDERT = PatternFill("solid", fgColor="FFE699")   # abweichend vom Standard
FILL_FEHLER = PatternFill("solid", fgColor="F8CBAD")
FILL_AUS_OBJEKT = PatternFill("solid", fgColor="DDEBF7")  # Basisjahr aus dem Objektblatt
SUMMENZEILEN = (1051, 1080, 1092, 1280, 1300, 1320, 1330, 1345, 1353, 1380)
GUELTIG = (STATUS_OK, STATUS_6B_UNZULAESSIG)

UNGUELTIGE_ZEICHEN = re.compile(r"[\[\]:*?/\\]")


def _basisjahr(modell: Modell) -> int:
    return modell.parameter.get(
        "par_Basisjahr", next(p.wert for p in PARAMETER if p.name == "par_Basisjahr"))


def blattname(wunsch: str, vorhanden) -> str:
    """Gültiger, eindeutiger Blattname (höchstens 31 Zeichen, ohne []:*?/\\)."""
    basis = UNGUELTIGE_ZEICHEN.sub("-", str(wunsch)).strip("'")[:31] or "BWA"
    name, i = basis, 2
    belegt = {t.lower() for t in vorhanden}
    while name.lower() in belegt:
        name = f"BWA {basis}"[:31] if i == 2 else f"{basis[:27]} ({i})"
        i += 1
    return name


def _summenformeln(z: dict, sp: str) -> dict:
    """Summenzeilen der BWA Form 01 als Formeln, wie im Kostenstellenblatt der Kanzlei."""
    def c(nr):
        return f"{sp}{z[nr]}"

    return {
        1051: f"={c(1020)}",
        1080: f"={c(1051)}",
        1092: f"={c(1080)}+{c(1090)}",
        1280: "=" + "+".join(c(nr) for nr in KOSTENARTEN),
        1300: f"={c(1092)}-{c(1280)}",
        1320: f"={c(1310)}+{c(1312)}",
        1330: f"={c(1322)}+{c(1323)}",
        1345: f"={c(1300)}-{c(1320)}+{c(1330)}",
        1353: f"={c(1345)}+{c(1351)}-{c(1352)}",
        1380: f"={c(1353)}-{c(1355)}",
    }


def _gewinne(jahr: str, vergleich: str, *bedingung: str) -> str:
    """Veräußerungsgewinne gültiger Verkäufe im Jahr; vergleich ">0" oder "<0"."""
    extra = "".join(f",{b}" for b in bedingung)
    return "+".join(f'SUMIFS(vk_Gewinn,vk_Jahr,{jahr},vk_Status,"{st}",vk_Gewinn,"{vergleich}"{extra})'
                    for st in GUELTIG)


def _verkauf(feld: str, jahr: str, *bedingung: str) -> str:
    """Summe eines Verkaufsfelds über die gültigen Verkäufe im Jahr."""
    extra = "".join(f",{b}" for b in bedingung)
    return "(" + "+".join(f'SUMIFS({feld},vk_Jahr,{jahr},vk_Status,"{st}"{extra})'
                          for st in GUELTIG) + ")"


# --- Zuordnung der Sonderposten zu BWA-Zeilen (Blatt BWA-Zuordnung) ---
#
# Verkauf, Rücklage, Neuobjekte und Zins stehen je BWA-Blatt in einem Herleitungsblock
# unter der BWA, ergebniswirksam gerechnet (Ertrag positiv, Aufwand negativ). Die BWA-Zeile
# je Posten wählt der Anwender im Blatt BWA-Zuordnung; die Zeile nimmt den Posten mit dem
# Vorzeichen ihrer Art auf (Aufwandszeilen negativ), das Ergebnis bleibt daher gleich.

ZUORDNUNG = "BWA-Zuordnung"
ERTRAGSZEILEN = (1020, 1090, 1322, 1323, 1351)
AUFWANDSZEILEN = (1240, 1250, 1260, 1310, 1312, 1352)
# Kostenarten 1100–1220 bleiben den eingelesenen Kosten vorbehalten
ZIELZEILEN = (1020, 1090, 1240, 1250, 1260, 1310, 1312, 1322, 1323, 1351, 1352)
NETTO, BRUTTO = "netto", "brutto"


def _netto(ausdruck: str) -> str:
    return f'IF(zuo_Verkauf="{BRUTTO}",0,{ausdruck})'


def _brutto(ausdruck: str) -> str:
    return f'IF(zuo_Verkauf="{BRUTTO}",{ausdruck},0)'


def _posten() -> list:
    """(Schlüssel, Bezeichnung, Standard-Nr., Erläuterung, Formel(jahr, vk, rl, prg, summe)).

    vk, rl, prg sind die Filter auf das Objekt des Blatts (leer im Summenblatt);
    summe ist True im Summenblatt. Ergebnis: ergebniswirksamer Betrag ohne „=“.
    """
    def x(*b):
        return tuple(f for f in b if f)

    def prg(name, jahr, f, vz=""):
        return f"{vz}SUMIFS({name},prg_Jahr,{jahr},prg_Neu,1{f})"

    zins = "SUMIFS(liq_Zins,liq_Jahr,{j})"
    return [
        ("gewinn", "Veräußerungsgewinn (netto)", 1323,
         "Gewinn gültiger Verkäufe im Verkaufsjahr; nur bei Ausweis netto",
         lambda j, vk, rl, pf, s: _netto(_gewinne(j, ">0", *x(vk)))),
        ("verlust", "Veräußerungsverlust (netto)", 1312,
         "Verlust gültiger Verkäufe im Verkaufsjahr; nur bei Ausweis netto",
         lambda j, vk, rl, pf, s: _netto(_gewinne(j, "<0", *x(vk)))),
        ("erloes", "Veräußerungspreis (brutto)", 1323,
         "Verkaufspreis gültiger Verkäufe; nur bei Ausweis brutto",
         lambda j, vk, rl, pf, s: _brutto(_verkauf("vk_Preis", j, *x(vk)))),
        ("kosten", "Veräußerungskosten (brutto)", 1312,
         "Kosten des Verkaufs; nur bei Ausweis brutto",
         lambda j, vk, rl, pf, s: _brutto("-" + _verkauf("vk_Kosten", j, *x(vk)))),
        ("abgang", "Buchwertabgang Gebäude und G+B (brutto)", 1312,
         "Buchwert Gebäude Ende Verkaufsjahr + AK G+B; nur bei Ausweis brutto",
         lambda j, vk, rl, pf, s: _brutto(f"-{_verkauf('vk_BuchwertGeb', j, *x(vk))}"
                                          f"-{_verkauf('vk_AKGuB', j, *x(vk))}")),
        ("bildung", "Einstellung § 6b-Rücklage", 1312, "im Verkaufsjahr",
         lambda j, vk, rl, pf, s: f"-SUMIFS(rl_Betrag,rl_Jahr,{j}{rl})"),
        ("aufloesung", "Auflösung § 6b-Rücklage", 1323,
         "nicht übertragener Rest im Fristjahr",
         lambda j, vk, rl, pf, s: f"SUMIFS(rl_Aufloesung,rl_Fristjahr,{j}{rl})"),
        ("zuschlag", "Gewinnzuschlag § 6b Abs. 7", 1323, "im Fristjahr",
         lambda j, vk, rl, pf, s: f"SUMIFS(rl_Zuschlag,rl_Fristjahr,{j}{rl})"),
        ("neu_miete", "Neuobjekte: Mieten", 1020, "ab dem Jahr nach dem Kauf",
         lambda j, vk, rl, pf, s: prg("prg_Miete", j, pf)),
        ("neu_einnahmen", "Neuobjekte: weitere Einnahmen", 1090, "",
         lambda j, vk, rl, pf, s: prg("prg_Einnahmen", j, pf)),
        ("neu_erhaltung", "Neuobjekte: Erhaltung", 1250, "",
         lambda j, vk, rl, pf, s: prg("prg_Erhaltung", j, pf, "-")),
        ("neu_ausgaben", "Neuobjekte: weitere Ausgaben", 1260, "",
         lambda j, vk, rl, pf, s: prg("prg_Ausgaben", j, pf, "-")),
        ("neu_afa", "Neuobjekte: Abschreibungen", 1240, "AfA nach Übertragung § 6b",
         lambda j, vk, rl, pf, s: prg("prg_AfA", j, pf, "-")),
        ("neu_zins", "Neuobjekte: Zinsen Kredite", 1310,
         "Kredit lt. Blatt Neuobjekte, Tilgungsplan im Blatt Darlehen",
         lambda j, vk, rl, pf, s: prg("prg_KreditZins", j, pf, "-")),
        ("zinsertrag", "Zinsertrag Alternativanlage", 1322, "nur im Blatt Alle Objekte",
         lambda j, vk, rl, pf, s: f"MAX({zins.format(j=j)},0)" if s else "0"),
        ("zinsaufwand", "Zinsaufwand bei negativer Liquidität", 1310,
         "nur im Blatt Alle Objekte",
         lambda j, vk, rl, pf, s: f"MIN({zins.format(j=j)},0)" if s else "0"),
    ]


POSTEN = _posten()
# Herleitungsblock unter der BWA: Titel, dann je Posten eine Zeile, dann die Summe
HERLEITUNG_TITEL = ERSTE_ZEILE + len(BWA_ZEILEN) + 1
HERLEITUNG_ERSTE = HERLEITUNG_TITEL + 1
HERLEITUNG_LETZTE = HERLEITUNG_ERSTE + len(POSTEN) - 1
SPALTE_ZIEL = 4   # D: BWA-Nr. des Postens (Spalte B bleibt den BWA-Zeilen vorbehalten)
# Finanzierung nachrichtlich unter der Herleitung (nicht ergebniswirksam): je Posten
# (Bezeichnung, Formel je Objekt mit {id} und {j}, Formel im Summenblatt mit {j})
FINANZIERUNG = [
    ("Kreditauszahlung (Kauf zum Jahresende)",
     "SUMIFS(ne_Kredit,ne_ID,{id},ne_Kaufjahr,{j},ne_Gueltig,1)", "SUMIFS(liq_Kredit,liq_Jahr,{j})"),
    ("Tilgung (im Verkaufsjahr mit Ablösung)", "SUMIFS(prg_Tilgung,prg_ID,{id},prg_Jahr,{j})", "SUMIFS(liq_Tilgung,liq_Jahr,{j})"),
    ("Restschuld Ende", "SUMIFS(prg_Restschuld,prg_ID,{id},prg_Jahr,{j})",
     "SUMIFS(liq_Restschuld,liq_Jahr,{j})"),
]
FINANZ_TITEL = HERLEITUNG_LETZTE + 2
FINANZ_LETZTE = FINANZ_TITEL + len(FINANZIERUNG)


def _zugeordnet(sp: str, nr: int) -> str:
    """Summe der Posten, die auf BWA-Zeile nr gebucht werden, mit deren Vorzeichen."""
    d = get_column_letter(SPALTE_ZIEL)
    vz = "+" if nr in ERTRAGSZEILEN else "-"
    return (f"{vz}SUMIFS({sp}${HERLEITUNG_ERSTE}:{sp}${HERLEITUNG_LETZTE},"
            f"${d}${HERLEITUNG_ERSTE}:${d}${HERLEITUNG_LETZTE},{nr})")


def _mit_zuordnung(z: dict, sp: str, basis: dict) -> dict:
    """Planformeln: Basiswerte plus zugeordnete Posten in den Zielzeilen, dann Summenzeilen."""
    formeln = {nr: f"={f}" for nr, f in basis.items()}
    for nr in ZIELZEILEN:
        formeln[nr] = f"={basis.get(nr, '0')}{_zugeordnet(sp, nr)}"
    formeln.update(_summenformeln(z, sp))
    return formeln


def _herleitung(ws, planjahre: int, objekt: bool) -> None:
    """Herleitungsblock unter der BWA: je Posten Ziel-Nr. und ergebniswirksamer Betrag."""
    c = ws.cell(row=HERLEITUNG_TITEL, column=3,
                value="Herleitung Sonderposten (+ Ertrag, − Aufwand); Zuordnung im Blatt "
                      f"{ZUORDNUNG}")
    c.font = Font(bold=True)
    ws.cell(row=HERLEITUNG_TITEL, column=SPALTE_ZIEL, value="→ Nr.").font = Font(bold=True)
    vk = "vk_ID,$B$2" if objekt else ""
    rl = ",rl_ObjektID,$B$2" if objekt else ""
    pf = ",prg_ID,$B$2" if objekt else ""
    for zeile, (key, text, _, _, formel) in zip(range(HERLEITUNG_ERSTE, HERLEITUNG_LETZTE + 1),
                                                POSTEN):
        ws.cell(row=zeile, column=3, value=text).font = FONT_HILFE_GROSS
        ziel = ws.cell(row=zeile, column=SPALTE_ZIEL, value=f"=zuo_{key}")
        ziel.font, ziel.fill = FONT_HILFE_GROSS, FILL_BERECHNET
        for i in range(planjahre):
            jahr = f"{get_column_letter(SPALTE_PLAN + i)}${ZEILE_JAHR}"
            zelle = ws.cell(row=zeile, column=SPALTE_PLAN + i,
                            value="=" + formel(jahr, vk, rl, pf, not objekt))
            zelle.number_format, zelle.font = FMT_BWA, FONT_HILFE_GROSS


def _finanzierung(ws, planjahre: int, id_: str = None) -> None:
    """Kredit nachrichtlich unter der Herleitung: Auszahlung, Tilgung, Restschuld je Planjahr.
    id_: Zelle mit der NeuID; ohne: Summe aller Kredite (Summenblatt)."""
    ws.cell(row=FINANZ_TITEL, column=3, value="Finanzierung (nachrichtlich, Zinsen in 1310)"
            ).font = Font(bold=True)
    for zeile, (text, objekt, summe) in enumerate(FINANZIERUNG, start=FINANZ_TITEL + 1):
        ws.cell(row=zeile, column=3, value=text).font = FONT_HILFE_GROSS
        for i in range(planjahre):
            jahr = f"{get_column_letter(SPALTE_PLAN + i)}${ZEILE_JAHR}"
            formel = (summe.format(j=jahr) if id_ is None else
                      f'IF({id_}="",0,{objekt.format(id=id_, j=jahr)})')
            c = ws.cell(row=zeile, column=SPALTE_PLAN + i, value=f"={formel}")
            c.number_format, c.font = FMT_BWA, FONT_HILFE_GROSS


def _planformeln_objekt(z: dict, sp: str) -> dict:
    """Planspalte eines Kostenstellenblatts; ObjektID in $B$2, Planjahr in Zeile 5.

    Bestandswerte (prg_Neu = 0) direkt, Neuobjekt, Verkauf und Rücklage über die Zuordnung.
    """
    jahr = f"{sp}${ZEILE_JAHR}"

    def p(name):
        return f"SUMIFS({name},prg_ID,$B$2,prg_Jahr,{jahr},prg_Neu,0)"

    basis = {
        1020: p("prg_Miete"),
        1090: p("prg_Einnahmen"),
        1240: p("prg_AfA"),
        1250: p("prg_Erhaltung"),
        1310: p("prg_KreditZins"),   # Bestandsdarlehen; Kredite der Neuobjekte über die Zuordnung
    }
    # Aufteilung der weiteren Ausgaben nach dem Anteil der Kostenart im Basisjahr
    ausg = "IFERROR(INDEX(obj_AusgBasis,MATCH($B$2,obj_ID,0)),0)"
    for nr in AUSGABEN_ZEILEN:
        basis[nr] = (f"IF(N({ausg})=0,0,{p('prg_Ausgaben')}"
                     f"*N(${get_column_letter(SPALTE_JAHR)}${z[nr]})/{ausg})")
    andere = "+".join(f"{sp}{z[nr]}" for nr in AUSGABEN_ZEILEN)
    basis[1260] = f"{p('prg_Ausgaben')}-({andere})"
    return _mit_zuordnung(z, sp, basis)


def _planformeln_summe(z: dict, sp: str, bestand: list) -> dict:
    """Planspalte des Summenblatts: Szenario A über alle Objekte, auch ohne eigenes Blatt."""
    jahr = f"{sp}${ZEILE_JAHR}"

    def p(name):
        return f"SUMIFS({name},prg_Jahr,{jahr},prg_Neu,0)"

    basis = {
        1020: p("prg_Miete"),
        1090: p("prg_Einnahmen"),
        1240: p("prg_AfA"),
        1250: p("prg_Erhaltung"),
        1310: p("prg_KreditZins"),
        1355: f"SUMIFS(liq_Steuer,liq_Jahr,{jahr})",
    }
    for nr in AUSGABEN_ZEILEN:
        basis[nr] = "+".join(f"'{b}'!{sp}{z[nr]}" for b in bestand) or "0"
    andere = "+".join(f"{sp}{z[nr]}" for nr in AUSGABEN_ZEILEN)
    basis[1260] = f"{p('prg_Ausgaben')}-({andere})"
    return _mit_zuordnung(z, sp, basis)


def _jahreszeile(ws, planjahre: int) -> None:
    ws.cell(row=ZEILE_JAHR, column=3, value="Jahr (für Formeln)").font = FONT_HILFE
    ws.cell(row=ZEILE_JAHR, column=SPALTE_JAHR, value="=par_Basisjahr").font = FONT_HILFE
    for i in range(planjahre):
        c = ws.cell(row=ZEILE_JAHR, column=SPALTE_PLAN + i, value=f"=par_Startjahr+{i}")
        c.font = FONT_HILFE
        c.number_format = FMT_JAHR


def _planspalten(ws, z: dict, planjahre: int, formeln_je_spalte) -> None:
    for i in range(planjahre):
        sp = get_column_letter(SPALTE_PLAN + i)
        for nr, formel in formeln_je_spalte(sp).items():
            c = ws.cell(row=z[nr], column=SPALTE_PLAN + i, value=formel)
            c.number_format = FMT_BWA
            c.fill = FILL_BERECHNET


# Aufschlüsselung der Abschreibungen je Anlage unter der Herleitung, im Layout des
# Kostenstellenblatts der Kanzlei: Block „Buchwert, JE“, dann „Abschreibungen JW“. Die
# Inventar-Nr. steht in Spalte D, die Werte kommen per MATCH aus dem Blatt Anlagen.
# Werte bei Halten, also auch nach einem Verkauf weiter fortgeschrieben.
SPALTE_ANLAGE = 4   # D


def _abschreibungen(ws, anlagen: list, planjahre: int) -> None:
    erste_zeile = FINANZ_LETZTE + 3
    vorjahr = get_column_letter(SPALTE_VORJAHRE + 1)   # Basisjahr − 1
    basis = get_column_letter(SPALTE_JAHR)
    plan = [get_column_letter(SPALTE_PLAN + i) for i in range(planjahre)]
    n = len(anlagen)
    titel = ws.cell(row=erste_zeile, column=3,
                    value="Aufschlüsselung Abschreibungen je Anlage (Blatt Anlagen, bei Halten)")
    titel.font = Font(bold=True)
    bw_kopf = erste_zeile + 1
    afa_kopf = bw_kopf + n + 2

    def match(zeile):
        return f"MATCH(${get_column_letter(SPALTE_ANLAGE)}{zeile},anl_Nr,0)"

    def zelle(zeile, sp, formel, fett=False):
        c = ws[f"{sp}{zeile}"]
        c.value = formel
        c.number_format = FMT_BWA
        c.font = Font(bold=fett)
        if not fett:
            c.fill = FILL_BERECHNET

    for kopf, text in ((bw_kopf, "Buchwert, JE"), (afa_kopf, "Abschreibungen JW")):
        ws.cell(row=kopf, column=3, value=text).font = Font(bold=True)
        for sp in [vorjahr, basis] + plan:
            zelle(kopf, sp, f"=SUM({sp}{kopf + 1}:{sp}{kopf + n})", fett=True)
    for k, anlage in enumerate(anlagen, start=1):
        for kopf in (bw_kopf, afa_kopf):
            zeile = kopf + k
            ws.cell(row=zeile, column=3, value=anlage.bezeichnung or anlage.nr)
            ws.cell(row=zeile, column=SPALTE_ANLAGE, value=anlage.nr).font = FONT_HILFE
        bw, afa = bw_kopf + k, afa_kopf + k
        zelle(bw, vorjahr, f'=IFERROR(IF(par_AnlStand=par_Basisjahr-1,'
                           f'INDEX(anl_BWStand,{match(bw)}),""),"")')
        zelle(bw, basis, f'=IFERROR(INDEX(anl_BWBasis,{match(bw)}),"")')
        zelle(afa, vorjahr, f'=IFERROR(IF(par_AnlStand=par_Basisjahr-1,'
                            f'INDEX(anl_AfAStand,{match(afa)}),""),"")')
        zelle(afa, basis, f'=IFERROR(INDEX(anl_AfABasis,{match(afa)}),"")')
        vor = basis
        for i, sp in enumerate(plan):
            zelle(afa, sp, f'=IFERROR(INDEX(anl_AfAJahre,{match(afa)},{i + 1}),"")')
            zelle(bw, sp, f"=N({vor}{bw})-N({sp}{afa})")
            vor = sp


def _blatt_objekt(wb, titel: str, objekt_id: str, name, ist, basisjahr: int,
                  planjahre: int, neu: bool, anlagen: list = (), kst: bool = False):
    """kst: Ist-Werte aus der Kostenstellen-Datei; dann ist die Spalte Basisjahr Eingabe
    (gelb) und führt, sonst kommt sie aus dem Blatt Objekte (hellblau)."""
    ws = wb.create_sheet(titel)
    z = bwa_kopf(ws, objekt_id, name or "", basisjahr, planjahre, monate_einklappen=True)
    _jahreszeile(ws, planjahre)
    for nr, werte in (ist or {}).items():
        if nr in z:
            for spalte, wert in werte.items():
                ws.cell(row=z[nr], column=spalte, value=wert)
    t = get_column_letter(SPALTE_JAHR)
    if kst:
        # Basisjahr führt: Einzelzeilen gelb zum Ändern, Summenzeilen als Formel
        for nr, bez in BWA_ZEILEN:
            if bez and nr not in SUMMENZEILEN and nr != 1094:
                ws.cell(row=z[nr], column=SPALTE_JAHR).fill = FILL_EINGABE
        for nr, formel in _summenformeln(z, t).items():
            ws.cell(row=z[nr], column=SPALTE_JAHR, value=formel)
    elif not neu:
        # ohne Kostenstellen-Datei: Basisjahr aus dem Objektblatt
        def obj(feld):
            return f"IFERROR(N(INDEX({feld},MATCH($B$2,obj_ID,0))),0)"
        _basisspalte(ws, z, obj)
    _planspalten(ws, z, planjahre, lambda sp: _planformeln_objekt(z, sp))
    _herleitung(ws, planjahre, objekt=True)
    _finanzierung(ws, planjahre, "$B$2")
    if anlagen:
        _abschreibungen(ws, anlagen, planjahre)
    if neu:
        basis = ""
    elif kst:
        basis = (" Basisjahr gelb: hier ändern, das Blatt Objekte und die Summe in "
                 f"„{SUMMENBLATT}“ übernehmen den Wert.")
    else:
        basis = (" Basisjahr hellblau: aus dem Blatt Objekte (1260 = weitere Ausgaben "
                 "abzüglich 1100–1220).")
    ws["E2"] = (("Neuobjekt: Werte ab dem Jahr nach dem Kauf, Kreditzinsen in 1310. " if neu
                 else "Zinsen der Darlehen in 1310 bis zum Verkauf, Tilgung und Restschuld "
                      "unter der Herleitung. ") + "Planspalten = Szenario A; Steuer und Zins der Alternativanlage nur "
                f"im Blatt „{SUMMENBLATT}“." + basis)
    ws["E2"].font = Font(italic=True)
    return ws


def _blatt_neukauf_kst(wb, titel: str, lw, basisjahr: int, planjahre: int):
    """Neukauf-Kostenstelle (hinter „KSt 9999“) als eigenes Blatt: Datenbasis der Neuobjekte,
    die auf sie verweisen. Basisjahr und Planjahre sind Eingabe (gelb), Summen als Formel."""
    ws = wb.create_sheet(titel)
    z = bwa_kopf(ws, lw.objekt_id, lw.name or "", basisjahr, planjahre, monate_einklappen=True)
    _jahreszeile(ws, planjahre)
    for nr, werte in lw.ist.items():
        if nr in z:
            for spalte, wert in werte.items():
                ws.cell(row=z[nr], column=spalte, value=wert)
    for nr, jahre in lw.plan.items():
        if nr in z and nr not in SUMMENZEILEN:
            for jahr, wert in jahre.items():
                i = jahr - basisjahr - 1
                if 0 <= i < planjahre:
                    ws.cell(row=z[nr], column=SPALTE_PLAN + i, value=wert)
    for spalte in range(SPALTE_JAHR, SPALTE_PLAN + planjahre):
        sp = get_column_letter(spalte)
        for nr, bez in BWA_ZEILEN:
            if bez and nr not in SUMMENZEILEN and nr != 1094:
                ws.cell(row=z[nr], column=spalte).fill = FILL_EINGABE
        for nr, formel in _summenformeln(z, sp).items():
            ws.cell(row=z[nr], column=spalte, value=formel)
    # Objektblatt des Neuobjekts, das auf die Kostenstelle verweist: AfA und Kreditzinsen
    # aus dem Modell; ohne Verweis bleibt der eingelesene Wert
    ws["C3"] = "Neuobjekt:"
    ws["C3"].font = FONT_HILFE_GROSS
    ws["D3"] = '=IFERROR(INDEX(ne_ID,MATCH($B$2,ne_KSt,0)),"")'
    ws["D3"].font, ws["D3"].fill = Font(bold=True), FILL_BERECHNET
    for nr, name in ((1240, "prg_AfA"), (1310, "prg_KreditZins")):
        for i in range(planjahre):
            c = ws.cell(row=z[nr], column=SPALTE_PLAN + i)
            eingelesen = c.value if isinstance(c.value, (int, float)) else '""'
            jahr = f"{get_column_letter(SPALTE_PLAN + i)}${ZEILE_JAHR}"
            c.value = (f'=IF($D$3="",{eingelesen},'
                       f"SUMIFS({name},prg_ID,$D$3,prg_Jahr,{jahr}))")
            c.fill = FILL_BERECHNET
    _finanzierung(ws, planjahre, "$D$3")
    ws["E2"] = ("Neukauf-Kostenstelle, zugleich Blatt des Neuobjekts in D3 (Blatt Neuobjekte, "
                "Spalte Kostenstelle Neukauf). Gelb ändern: Miete 1020, weitere Einnahmen 1090, "
                "Erhaltung 1250 und weitere Ausgaben 1100–1220, 1260 gehen über das Blatt "
                "Neukauf-KSt in die Prognose ab dem Jahr nach dem Kauf; leere Jahre werden "
                "fortgeschrieben. AfA 1240 und Kreditzinsen 1310 rechnet das Modell (grau). "
                f"Das Neuobjekt zählt in „{SUMMENBLATT}“.")
    ws["E2"].font = Font(italic=True)
    return ws


def verknuepfe_neukauf(wb, block: int, titel: str) -> None:
    """Zeilen der Neukauf-Kostenstelle im Blatt Neukauf-KSt verweisen auf ihr Blatt:
    Jahr mit mindestens einem Wert = Summe der BWA-Zeilen, sonst leer (Fortschreibung)."""
    from . import formeln
    from .einlesen import NEUKAUF_POSITIONEN as ZEILEN_JE_POSITION
    from .modelle import NEUKAUF_POSITIONEN
    ws = wb["Neukauf-KSt"]
    j0, _, import0, n = formeln.nk_spalten()
    blatt = titel.replace("'", "''")
    kopf = formeln.NK_ERSTE + block * len(NEUKAUF_POSITIONEN)
    for k, (key, *_) in enumerate(NEUKAUF_POSITIONEN):
        for i in range(n):
            sp = get_column_letter(SPALTE_JAHR if i == 0 else SPALTE_PLAN + i - 1)
            refs = ",".join(f"'{blatt}'!${sp}${ZEILEN[nr]}" for nr in ZEILEN_JE_POSITION[key]
                            if nr in ZEILEN)
            formel = f'=IF(COUNT({refs})=0,"",SUM({refs}))'
            ws.cell(row=kopf + k, column=j0 + i, value=formel)
            ws.cell(row=kopf + k, column=import0 + i, value=formel)


def _bezug(titel: str, nr: int, z: dict) -> str:
    blatt = titel.replace("'", "''")
    return f"'{blatt}'!${get_column_letter(SPALTE_JAHR)}${z[nr]}"


def verknuepfe_objekt(wb, zeile: int, titel: str, z: dict) -> None:
    """Basiswerte der Objektzeile verweisen auf die Spalte Basisjahr des Kostenstellenblatts;
    die ausgeblendete Spalte „eingelesen“ hält dieselbe Verknüpfung zum Vergleich."""
    from . import formeln
    ws = wb["Objekte"]
    for key, nummern in OBJEKT_AUS_KST.items():
        formel = "=" + "+".join(_bezug(titel, nr, z) for nr in nummern)
        ws[f"{formeln.spalte(key)}{zeile}"] = formel
        ws[f"{formeln.ispalte(key)}{zeile}"] = formel
    ws[f"{formeln.kst_blatt_spalte()}{zeile}"] = titel


def _basisspalte(ws, z: dict, obj) -> None:
    """Spalte Basisjahr: Werte des Objektblatts (obj(Bereich) liefert den Ausdruck), 1260 =
    weitere Ausgaben abzüglich der Kostenarten 1100–1220 (wie eingetragen), Summen als Formel."""
    t = get_column_letter(SPALTE_JAHR)
    for nr, feld in IST_AUS_OBJEKT.items():
        ws.cell(row=z[nr], column=SPALTE_JAHR, value=f"={obj(feld)}")
    andere = "+".join(f"N({t}{z[nr]})" for nr in AUSGABEN_ZEILEN)
    ws.cell(row=z[1260], column=SPALTE_JAHR, value=f"={obj(AUSGABEN_AUS_OBJEKT)}-({andere})")
    for nr, formel in _summenformeln(z, t).items():
        ws.cell(row=z[nr], column=SPALTE_JAHR, value=formel)
    for nr in list(IST_AUS_OBJEKT) + [1260]:
        ws.cell(row=z[nr], column=SPALTE_JAHR).fill = FILL_AUS_OBJEKT


def blaetter_bwa(wb, modell: Modell) -> list:
    """Summenblatt und je Objekt und Neuobjekt ein BWA-Blatt; liefert die Blattnamen."""
    basisjahr = _basisjahr(modell)
    planjahre = prognosejahre()
    summe = wb.create_sheet(blattname(SUMMENBLATT, wb.sheetnames))
    eintraege = []
    for zeile, obj in enumerate(modell.objekte, start=2):   # Zeile im Blatt Objekte
        if not obj.objekt_id:
            continue
        ist = modell.kostenstellen.get(obj.objekt_id)
        wunsch = ist.blatt if ist is not None else obj.objekt_id
        eintraege.append((wunsch, obj.objekt_id, obj.name, ist.ist if ist else None, False,
                          zeile if ist is not None else None))
    for ne in modell.neuobjekte:
        # mit Neukauf-Kostenstelle ist deren Blatt das Blatt des Neuobjekts
        if ne.neu_id and ne.kst not in modell.neukauf:
            eintraege.append((ne.neu_id, ne.neu_id, ne.name, None, True, None))
    titel = []
    for wunsch, objekt_id, name, ist, neu, obj_zeile in eintraege:
        t = blattname(wunsch, wb.sheetnames)
        anlagen = [a for a in modell.anlagen
                   if not neu and a.objekt_id == objekt_id and a.art in ARTEN_ABNUTZBAR]
        _blatt_objekt(wb, t, objekt_id, name, ist, basisjahr, planjahre, neu, anlagen,
                      kst=obj_zeile is not None)
        if obj_zeile is not None:
            verknuepfe_objekt(wb, obj_zeile, t, ZEILEN)
        titel.append(t)
    neukauf = []
    for block, lw in enumerate(list(modell.neukauf.values())[:modell.kapazitaet.neukauf]):
        t = blattname(lw.blatt or lw.objekt_id, wb.sheetnames)
        _blatt_neukauf_kst(wb, t, lw, basisjahr, planjahre)
        verknuepfe_neukauf(wb, block, t)
        neukauf.append(t)

    z = bwa_kopf(summe, "KSt", SUMMENBLATT, basisjahr, planjahre, monate_einklappen=True)
    _jahreszeile(summe, planjahre)
    bestand = [t for t, e in zip(titel, eintraege) if not e[4]]
    # Ist-Werte und Basisjahr: Summe der Kostenstellenblätter
    for spalte in range(SPALTE_VORJAHRE, SPALTE_JAHR + 1):
        sp = get_column_letter(spalte)
        for nr, bez in BWA_ZEILEN:
            if bez and bestand:
                summe.cell(row=z[nr], column=spalte,
                           value="=" + "+".join(f"N('{b}'!{sp}{z[nr]})" for b in bestand))
    _planspalten(summe, z, planjahre, lambda sp: _planformeln_summe(z, sp, bestand))
    _herleitung(summe, planjahre, objekt=False)
    _finanzierung(summe, planjahre)
    summe["E2"] = ("Planspalten = Szenario A über alle Objekte und Neuobjekte, mit den "
                   "Krediten der Neuobjekte. Nr. 1353 = Ergebnis vor Verlustvortrag im Blatt "
                   f"Liquidität (Kontrolle im Blatt {ZUORDNUNG}), Steuer mit Verlustvortrag. "
                   "Ist-Spalten und Basisjahr: Summe der Kostenstellenblätter; Objekte ohne "
                   "eigenes Blatt nur in den Planspalten.")
    summe["E2"].font = Font(italic=True)
    ende = get_column_letter(SPALTE_PLAN + planjahre - 1)
    anfang = get_column_letter(SPALTE_PLAN)
    for nr, bez in BWA_ZEILEN:
        if bez:
            _name(wb, f"bwa_{nr}", f"'{summe.title}'!${anfang}${z[nr]}:${ende}${z[nr]}")
    _name(wb, "bwa_Jahr", f"'{summe.title}'!${anfang}${ZEILE_JAHR}:${ende}${ZEILE_JAHR}")
    for zeile, posten in zip(range(HERLEITUNG_ERSTE, HERLEITUNG_LETZTE + 1), POSTEN):
        _name(wb, f"bwah_{posten[0]}", f"'{summe.title}'!${anfang}${zeile}:${ende}${zeile}")
    zuordnung = blatt_zuordnung(wb, summe, z, planjahre, modell.bwa_zuordnung)
    return [zuordnung, summe.title] + titel + neukauf


def blatt_zuordnung(wb, summe, z: dict, planjahre: int, vorgabe: dict) -> str:
    """Steuerblatt: BWA-Zeile je Sonderposten, Ausweis des Verkaufs, Kontrolle.

    Steht vor dem Summenblatt. Die Kontrolle vergleicht je Planjahr das Ergebnis des
    Summenblatts (Nr. 1353) mit dem Ergebnis vor Verlustvortrag im Blatt Liquidität.
    """
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.worksheet.datavalidation import DataValidation

    ws = wb.create_sheet(ZUORDNUNG, wb.sheetnames.index(summe.title))
    ws["A1"] = "BWA-Zuordnung: auf welche BWA-Zeile Verkauf, Rücklage, Kauf und Zins gehen"
    ws["A1"].font = FONT_TITEL
    ws["A2"] = ("Gelbe Zellen ändern. Jede BWA-Zeile nimmt den Posten mit dem Vorzeichen "
                "ihrer Art auf (Ertrag +, Aufwand −); das Ergebnis bleibt bei jeder Zuordnung "
                "gleich. Die Beträge je Jahr stehen unter jeder BWA im Block "
                "„Herleitung Sonderposten“.")
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:G2")
    ws.row_dimensions[2].height = 45

    ws["A4"] = "Verkauf ausweisen"
    ws["A4"].font = Font(bold=True)
    ws["B4"] = vorgabe.get("verkauf", NETTO)
    ws["B4"].fill = FILL_EINGABE
    ws["C4"] = ("netto: nur der Veräußerungsgewinn bzw. -verlust; brutto: Verkaufspreis als "
                "Ertrag, Verkaufskosten und Buchwertabgang als Aufwand")
    dv = DataValidation(type="list", formula1=f'"{NETTO},{BRUTTO}"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B4")
    _name(wb, "zuo_Verkauf", f"'{ZUORDNUNG}'!$B$4")

    kopf = ["Posten", "BWA-Nr.", "BWA-Zeile", "Ertrag/Aufwand", "Standard",
            "Summe Planjahre (Alle Objekte)", "Erläuterung"]
    for spalte, text in enumerate(kopf, start=1):
        c = ws.cell(row=6, column=spalte, value=text)
        c.font, c.fill = Font(bold=True), FILL_KOPF
        c.alignment = Alignment(wrap_text=True)
    # Liste der zulässigen Zielzeilen rechts daneben
    for spalte, text in enumerate(["zulässige Nr.", "Bezeichnung", "Art"], start=9):
        c = ws.cell(row=6, column=spalte, value=text)
        c.font, c.fill = Font(bold=True), FILL_KOPF
    bezeichnung = dict(BWA_ZEILEN)
    for zeile, nr in enumerate(ZIELZEILEN, start=7):
        ws.cell(row=zeile, column=9, value=nr)
        ws.cell(row=zeile, column=10, value=bezeichnung[nr])
        ws.cell(row=zeile, column=11, value="Ertrag" if nr in ERTRAGSZEILEN else "Aufwand")
    liste_ende = 6 + len(ZIELZEILEN)
    _name(wb, "zuo_Nummern", f"'{ZUORDNUNG}'!$I$7:$I${liste_ende}")
    _name(wb, "zuo_Liste", f"'{ZUORDNUNG}'!$I$7:$K${liste_ende}")

    erste = 7
    letzte = erste + len(POSTEN) - 1
    dv_nr = DataValidation(type="list", formula1="zuo_Nummern", allow_blank=False,
                           showErrorMessage=True, errorTitle="BWA-Nr.",
                           error="Nur die Nummern der Liste rechts (Spalte I).")
    ws.add_data_validation(dv_nr)
    for zeile, (key, text, standard, erlaeuterung, _) in zip(range(erste, letzte + 1), POSTEN):
        ws.cell(row=zeile, column=1, value=text)
        nr = ws.cell(row=zeile, column=2, value=vorgabe.get(key, standard))
        nr.fill = FILL_EINGABE
        dv_nr.add(nr.coordinate)
        _name(wb, f"zuo_{key}", f"'{ZUORDNUNG}'!$B${zeile}")
        for spalte, formel, fmt in (
                (3, f'=IFERROR(VLOOKUP($B{zeile},zuo_Liste,2,0),"ungültige Nr.")', "@"),
                (4, f'=IFERROR(VLOOKUP($B{zeile},zuo_Liste,3,0),"")', "@"),
                (5, standard, "0"),
                (6, f"=SUM(bwah_{key})", FMT_EURO)):
            c = ws.cell(row=zeile, column=spalte, value=formel)
            c.number_format = fmt
            if spalte != 5:
                c.fill = FILL_BERECHNET
        ws.cell(row=zeile, column=7, value=erlaeuterung)
    _name(wb, "zuo_Nr", f"'{ZUORDNUNG}'!$B${erste}:$B${letzte}")
    ws.conditional_formatting.add(
        f"B{erste}:C{letzte}",
        FormulaRule(formula=[f"COUNTIF(zuo_Nummern,$B{erste})=0"], fill=FILL_FEHLER))
    ws.conditional_formatting.add(
        f"B{erste}:B{letzte}",
        FormulaRule(formula=[f"$B{erste}<>$E{erste}"], fill=FILL_GEAENDERT))

    # Kontrolle gegen die Liquidität, je Planjahr eine Spalte ab B
    zeile = letzte + 3
    ws.cell(row=zeile, column=1, value="Kontrolle: BWA Alle Objekte gegen Blatt Liquidität "
                                       "(Szenario A)").font = FONT_ABSCHNITT
    zeile += 1
    jahr_z, bwa_z, liq_z, diff_z = zeile, zeile + 1, zeile + 2, zeile + 3
    for text, z_ in (("Jahr", jahr_z), ("Ergebnis lt. BWA Alle Objekte (Nr. 1353)", bwa_z),
                     ("Ergebnis vor Verlustvortrag lt. Liquidität", liq_z),
                     ("Differenz (muss 0 sein)", diff_z)):
        ws.cell(row=z_, column=1, value=text).font = Font(bold=z_ in (jahr_z, diff_z))
    for i in range(planjahre):
        sp = get_column_letter(2 + i)
        quelle = get_column_letter(SPALTE_PLAN + i)
        for z_, formel, fmt in (
                (jahr_z, f"=par_Startjahr+{i}", FMT_JAHR),
                (bwa_z, f"='{summe.title}'!{quelle}{z[1353]}", FMT_EURO),
                (liq_z, f"=INDEX(liq_ZvE,{i + 1})", FMT_EURO),
                (diff_z, f"=ROUND({sp}{bwa_z}-{sp}{liq_z},2)", FMT_EURO)):
            c = ws.cell(row=z_, column=2 + i, value=formel)
            c.number_format, c.fill = fmt, FILL_BERECHNET
    ende = get_column_letter(1 + planjahre)
    _name(wb, "zuo_Differenz", f"'{ZUORDNUNG}'!$B${diff_z}:${ende}${diff_z}")
    ws.conditional_formatting.add(
        f"B{diff_z}:{ende}{diff_z}", FormulaRule(formula=[f"B{diff_z}<>0"], fill=FILL_FEHLER))

    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["F"].width = 18
    ws.column_dimensions["G"].width = 50
    ws.column_dimensions["J"].width = 26
    ws.freeze_panes = "B7"
    return ws.title


def _name(wb, name: str, ref: str) -> None:
    from openpyxl.workbook.defined_name import DefinedName
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


# --- Sonderbereich Verkauf und Kauf ---


def _text(zelle, text: str):
    """Beschriftung als Text, auch wenn sie mit „=“ beginnt wie in der Kanzleivorlage."""
    zelle.value = text
    if isinstance(text, str) and text.startswith("="):
        zelle.data_type = "s"
    return zelle


class _Block:
    """Schreibt Zeilen eines Blocks; Formeln nennen andere Zeilen über {schlüssel}."""

    def __init__(self, ws, zeile: int):
        self.ws, self.zeile, self.adr = ws, zeile, {}

    def leer(self, n: int = 1) -> None:
        self.zeile += n

    def titel(self, text, font=FONT_ABSCHNITT) -> None:
        self.ws.cell(row=self.zeile, column=1, value=text).font = font
        self.zeile += 1

    def kopf(self, *texte) -> None:
        for spalte, text in enumerate(texte, start=1):
            c = self.ws.cell(row=self.zeile, column=spalte, value=self._ersetzen(text))
            c.font = Font(bold=True)
            c.fill = FILL_KOPF
            c.alignment = Alignment(wrap_text=True)
        self.zeile += 1

    def zeile_(self, key, text, *formeln, fmt=FMT_EURO, fett=False) -> None:
        """Zeile mit Beschriftung in A und Formeln ab B; {key} = Zelle B, {key.C} = Zelle C."""
        if key:
            self.adr[key] = self.zeile
        _text(self.ws.cell(row=self.zeile, column=1), text).font = Font(bold=fett)
        for spalte, formel in enumerate(formeln, start=2):
            if formel is None:
                continue
            c = self.ws.cell(row=self.zeile, column=spalte, value=self._ersetzen(formel))
            c.number_format = fmt
            c.font = Font(bold=fett)
            if isinstance(formel, str) and formel.startswith("="):
                c.fill = FILL_BERECHNET
        self.zeile += 1

    def _ersetzen(self, formel):
        if not isinstance(formel, str):
            return formel

        def adresse(m):
            key, _, spalte = m.group(1).partition(".")
            return f"${spalte or 'B'}${self.adr[key]}"
        return re.sub(r"\{(\w+(?:\.[A-Z])?)\}", adresse, formel)


def _vorgang(b: _Block, n: int) -> None:
    """Ergebnis- und Detailsicht eines Verkaufs (Zeile n im Blatt Verkäufe)."""
    def vk(name):
        return f"INDEX({name},{n})"

    def rl(name):
        return f"N(INDEX({name},{n}))"

    def ok(ausdruck):
        return f'=IF({{ok}}=1,{ausdruck},"")'

    def halten(feld, satz):
        return f"N(INDEX({feld},MATCH({{id}},obj_ID,0)))*(1+{satz})^({{x}}-par_Basisjahr)"

    def alt(name):
        return f'IF({{rl}}="",0,SUMIFS({name},prg_Quelle,{{rl}},prg_Jahr,{{x}}))'

    def bw_halten(jahr):
        return f"SUMIFS(prg_BuchwertHalten,prg_ID,{{id}},prg_Jahr,{jahr})"

    status = vk("vk_Status")
    b.titel(f'=IF({vk("vk_ID")}="","Vorgang {n}: kein Verkauf erfasst","Vorgang {n}: Verkauf "'
            f'&{vk("vk_ID")}&IFERROR(" "&INDEX(obj_Name,MATCH({vk("vk_ID")},obj_ID,0)),"")'
            f'&" zum Ende "&{vk("vk_Jahr")})')
    b.zeile_("id", "Kostenstelle (ObjektID)", f'=IF({vk("vk_ID")}="","",{vk("vk_ID")})',
             fmt="@")
    b.zeile_("status", "Status Verkauf", f'=IF({{id}}="","",{status})', fmt="@")
    b.zeile_("ok", "rechnet mit (1 = ja)",
             f'=IF(OR({{status}}="{GUELTIG[0]}",{{status}}="{GUELTIG[1]}"),1,0)', fmt="0")
    b.zeile_("jahr", "Verkaufsjahr (Verkauf zum Jahresende)", ok(vk("vk_Jahr")), fmt=FMT_JAHR)
    b.zeile_("rl", "RücklageID (§ 6b)", ok(f'IF({vk("rl_ID")}="","",{vk("rl_ID")})'), fmt="@")
    # Kaufjahr&"0" / 10: leere Annahmeformeln ("") zählen als 0 statt #WERT; geht in
    # jeder Excel-Version ohne Matrixformel (MAXIFS erst ab Excel 2019)
    letzter_kauf = (f'IF({{rl}}="",0,SUMPRODUCT(MAX((ne_Quelle={{rl}})'
                    f'*(ne_Status="{STATUS_OK}")*(ne_Kaufjahr&"0")/10)))')
    b.zeile_("x", "Vergleichsjahr: erstes volles Jahr nach Verkauf und Kauf",
             ok(f"MIN(MAX({{jahr}},{letzter_kauf})+1,par_Endjahr)"), fmt=FMT_JAHR)
    b.leer()

    # Planung im Vergleichsjahr (unten); die Ergebnissicht fasst sie zusammen
    b.titel("Ergebnissicht (für Berichterstattung)")
    b.zeile_("preis", "Bewertung/Erlös (Verkaufspreis)", ok(f"N({vk('vk_Preis')})"))
    reinvest = "+".join(f'SUMIFS({f},ne_Quelle,{{rl}},ne_Status,"{STATUS_OK}")'
                        for f in ("ne_Kaufpreis", "ne_Nebenkosten"))
    b.zeile_("reinvest", "Reinvestition (Kaufpreis und Nebenkosten der Neuobjekte)",
             ok(f'IF({{rl}}="",0,{reinvest})'))
    kredit = f'SUMIFS(ne_Kredit,ne_Quelle,{{rl}},ne_Status,"{STATUS_OK}")'
    b.zeile_("kredit", "Kredit der Neuobjekte", ok(f'IF({{rl}}="",0,{kredit})'))
    b.zeile_("uebertrag", "Übertrag § 6b EStG", ok(f"-({rl('rl_UebGeb')}+{rl('rl_UebGuB')})"))
    b.zeile_("steuer", "Steuer auf den Gewinn ca. (Verkaufsjahr, Auflösung im Fristjahr)",
             ok(f"-(N({vk('vk_Gewinn')})-{rl('rl_Betrag')}+{rl('rl_Aufloesung')}"
                f"+{rl('rl_Zuschlag')})*par_Steuersatz"))
    # Verkauf zum Jahresende: Restschuld nach der Rate des Verkaufsjahrs
    b.zeile_("abloesung", "Ablösung Darlehen des Objekts (Restschuld Ende Verkaufsjahr)",
             ok("-SUMIFS(prg_RestschuldHalten,prg_ID,{id},prg_Jahr,{jahr})"))
    b.zeile_("anlage", "Kapitalanlage (Nettoerlös − Reinvestition + Kredit − Steuer ca. − "
             "Ablösung)",
             ok(f"N({vk('vk_Nettoerloes')})-{{reinvest}}+{{kredit}}+{{steuer}}+{{abloesung}}"))
    b.zeile_(None, "Restschuld Kredite der Neuobjekte im Vergleichsjahr",
             ok(alt("prg_Restschuld")))
    b.leer()
    b.kopf(f'=IF({{ok}}=1,"Vergleich im Jahr "&{{x}},"Vergleich")', "Ausgangsfall (halten)",
           "Alternative (Verkauf, Kauf, Anlage)", "Differenz")
    zus = [("e_miete", "Mietertrag (netto)", "{p_miete}+{p_einn}"),
           ("e_kapital", "Kapitalertrag", "{p_zins}"),
           ("e_aufwand", "Aufwand (Erhaltung, weitere Ausgaben, AfA; ohne Zinsen)",
            "{p_erh}+{p_ausg}+{p_afa}"),
           ("e_zinsen", "Zinsen Darlehen", "{p_zinsen}"),
           ("e_ergebnis", "vorläufiges Ergebnis", "{p_ergebnis}"),
           ("e_vor", "liquider Überschuss vor Steuern", "{p_cashflow}"),
           ("e_steuer", "Steuern ca.", "-{p_ergebnis}*par_Steuersatz"),
           ("e_nach", "liquider Überschuss nach Steuern ca.", "{e_vor}+{e_steuer}")]
    ergebnis_zeile = b.zeile
    b.leer(len(zus) + 1)   # Platz; gefüllt, sobald die Planung steht

    b.titel("Detailsicht (Einzelauflistung)")
    b.kopf("Verkauf", "gesamt", "G+B", "Gebäude")
    q = f"N({vk('vk_AnteilGuB')})"
    b.zeile_("anteil", "Anteil", ok("1"), ok(q), ok(f"1-{q}"), fmt=FMT_PROZENT)
    b.zeile_("vp", "Veräußerungspreis", "={preis}", ok("{preis}*{anteil.C}"),
             ok("{preis}*{anteil.D}"))
    kosten = f"N({vk('vk_Kosten')})"
    b.zeile_("vkosten", "./. Veräußerungskosten", ok(f"-{kosten}"),
             ok(f"-{kosten}*{{anteil.C}}"), ok(f"-{kosten}*{{anteil.D}}"))
    b.zeile_("bw", "./. Buchwert (Gebäude zum Ende des Verkaufsjahrs, G+B = AK)",
             ok("{bw.C}+{bw.D}"), ok(f"-N({vk('vk_AKGuB')})"),
             ok(f"-N({vk('vk_BuchwertGeb')})"))
    b.zeile_("gewinn", "= Veräußerungsgewinn", ok(f"N({vk('vk_Gewinn')})"),
             ok(f"N({vk('vk_GewinnGuB')})"), ok(f"N({vk('vk_GewinnGeb')})"), fett=True)
    b.zeile_("ruecklage", "Rücklage § 6b (gebildet im Verkaufsjahr)", ok(rl("rl_Betrag")),
             ok(rl("rl_GuB")), ok(rl("rl_Geb")))
    b.zeile_("uebertragen", "übertragen auf Neuobjekte (G+B-Gewinn auf G+B und Gebäude)",
             ok("{uebertragen.C}+{uebertragen.D}"), ok(rl("rl_UebGuB")), ok(rl("rl_UebGeb")))
    b.zeile_("frist", "Fristjahr", ok(f'IF({{rl}}="","",{rl("rl_Fristjahr")})'), fmt=FMT_JAHR)
    b.zeile_("aufl", "aufgelöst im Fristjahr", ok(rl("rl_Aufloesung")))
    b.zeile_("zuschlag", "Gewinnzuschlag", ok(rl("rl_Zuschlag")))
    b.leer()

    b.kopf(f'=IF({{ok}}=1,"Planung im Jahr "&{{x}},"Planung")', "Ausgangsfall (halten)",
           "Alternative (Verkauf, Kauf, Anlage)", "Differenz")
    afa_h = f"{bw_halten('{x}-1')}-{bw_halten('{x}')}"
    planung = [
        ("p_miete", "Mieten", halten("obj_MieteBasis", "par_Mietsteig"), alt("prg_Miete")),
        ("p_einn", "weitere Einnahmen", halten("obj_EinnBasis", "par_Mietsteig"),
         alt("prg_Einnahmen")),
        ("p_zins", "Zinsertrag Kapitalanlage", "0", "{anlage}*par_Alternativrendite"),
        ("p_erh", "./. Erhaltung (mit Alterung und Großmaßnahme)",
         "-SUMIFS(prg_ErhaltungHalten,prg_ID,{id},prg_Jahr,{x})",
         "-" + alt("prg_Erhaltung")),
        ("p_ausg", "./. weitere Ausgaben", "-" + halten("obj_AusgBasis", "par_Kostensteig"),
         "-" + alt("prg_Ausgaben")),
        ("p_afa", "./. Abschreibungen (Steuerbilanz)", f"-({afa_h})", "-" + alt("prg_AfA")),
        ("p_zinsen", "./. Zinsen Darlehen",
         "-SUMIFS(prg_ZinsHalten,prg_ID,{id},prg_Jahr,{x})", "-" + alt("prg_KreditZins")),
        ("p_ergebnis", "= vorläufiges Ergebnis", "{p_miete}+{p_einn}+{p_zins}+{p_erh}+{p_ausg}"
         "+{p_afa}+{p_zinsen}", None),
        ("p_tilgung", "./. Tilgungen", "-SUMIFS(prg_TilgungHalten,prg_ID,{id},prg_Jahr,{x})",
         "-" + alt("prg_Tilgung")),
        ("p_afa_zurueck", "+ Abschreibungen", "-{p_afa}", None),
        ("p_cashflow", "= Cash Flow", "{p_ergebnis}+{p_tilgung}+{p_afa_zurueck}", None),
    ]
    for key, text, ausgang, alternative in planung:
        if alternative is None:   # Summenzeile: gleiche Rechnung in B und C
            alternative = re.sub(r"\{(\w+)\}", r"{\1.C}", ausgang)
        b.zeile_(key, text, ok(ausgang), ok(alternative), ok(f"{{{key}.C}}-{{{key}}}"),
                 fett=text.startswith("="))

    # Ergebnissicht: Zusammenfassung der Planung
    weiter = b.zeile
    b.zeile = ergebnis_zeile
    for key, text, formel in zus:
        alternative = re.sub(r"\{(\w+)\}", r"{\1.C}", formel)
        b.zeile_(key, text, ok(formel), ok(alternative), ok(f"{{{key}.C}}-{{{key}}}"),
                 fett=key in ("e_ergebnis", "e_nach"))
    b.zeile = weiter
    b.leer(2)


def _kauf(b: _Block, anzahl: int) -> None:
    """Detailsicht je Neuobjekt, eine Spalte je Zeile im Blatt Neuobjekte."""
    b.titel("Detailsicht Kauf (je Neuobjekt)")
    b.kopf("Neuobjekt", *[f"Neuobjekt {m}" for m in range(1, anzahl + 1)])

    def zeile(text, ausdruck, fmt=FMT_EURO, fett=False):
        formeln = []
        for m in range(1, anzahl + 1):
            ne_id = f"INDEX(ne_ID,{m})"
            formeln.append(f'=IF({ne_id}="","",{ausdruck.format(m=m, id=ne_id, sp="{sp}")})')
        _text(b.ws.cell(row=b.zeile, column=1), text).font = Font(bold=fett)
        for spalte, formel in enumerate(formeln, start=2):
            formel = formel.replace("{sp}", get_column_letter(spalte))
            c = b.ws.cell(row=b.zeile, column=spalte, value=formel)
            c.number_format = fmt
            c.fill = FILL_BERECHNET
            c.font = Font(bold=fett)
        b.zeile += 1
        return b.zeile - 1

    def ne(name):
        return f"INDEX({name},{{m}})"

    zeile("NeuID (Kostenstelle)", "{id}", "@", fett=True)
    zeile("Name", f'IF({ne("ne_Name")}="","",{ne("ne_Name")})', "@")
    zeile("Status", ne("ne_Status"), "@")
    for nr, (_, quelle, _) in enumerate(QUELLEN, start=1):
        zeile(f"Quelle {nr} RücklageID", f'IF({ne(quelle)}="","",{ne(quelle)})', "@")
    kj = zeile("Kaufjahr (Kauf zum Jahresende)", ne("ne_Kaufjahr"), FMT_JAHR)
    zeile("Kaufpreis", f"N({ne('ne_Kaufpreis')})")
    zeile("Kaufnebenkosten", f"N({ne('ne_Nebenkosten')})")
    zeile("Anteil G+B", f"N({ne('ne_AnteilGuB')})", FMT_PROZENT)
    zeile("AK G+B", f"N({ne('ne_AKGuBNeu')})")
    zeile("AK Gebäude", f"N({ne('ne_AKGebNeu')})")
    zeile("ü1 Gebäudegewinn auf Gebäude", f"-N({ne('ne_Ue1')})")
    zeile("ü2 G+B-Gewinn auf G+B", f"-N({ne('ne_Ue2')})")
    zeile("ü3 G+B-Gewinn auf Gebäude", f"-N({ne('ne_Ue3')})")
    zeile("Übertrag § 6b EStG gesamt", f"-N({ne('ne_UeGesamt')})", fett=True)
    zeile("AfA-Bemessungsgrundlage Gebäude", f"N({ne('ne_AfABasis')})", fett=True)
    zeile("steuerliche AK G+B", f"N({ne('ne_AKGuB')})")
    zeile("AfA-Methode", f'IF({ne("ne_AfAMethode")}="","linear",{ne("ne_AfAMethode")})', "@")
    zeile("AfA-Satz", f"N({ne('ne_AfASatz')})", FMT_PROZENT)
    zeile("Kaufpreis + Nebenkosten", f"N({ne('ne_AKGesamt')})")
    zeile("./. Einsatz Verkaufserlös (Quell-Verkäufe)", f"-N({ne('ne_Erloes')})")
    zeile("= Finanzierungsbedarf", f"N({ne('ne_Bedarf')})", fett=True)
    zeile("davon Kredit", f"N({ne('ne_Kredit')})")
    zeile("davon Eigenmittel", f"N({ne('ne_Eigen')})")
    jahr = zeile("erstes volles Jahr", f"{{sp}}{kj}+1", FMT_JAHR)

    def prg(name):
        return f"SUMIFS({name},prg_ID,{{id}},prg_Jahr,{{sp}}{jahr})"

    miete = zeile("Mieten", prg("prg_Miete"))
    erh = zeile("./. Erhaltung", f"-{prg('prg_Erhaltung')}")
    afa = zeile("./. Abschreibungen (Steuerbilanz)", f"-{prg('prg_AfA')}")
    zinsen = zeile("./. Zinsen Kredit", f"-{prg('prg_KreditZins')}")
    erg = zeile("= vorläufiges Ergebnis",
                f"{{sp}}{miete}+{{sp}}{erh}+{{sp}}{afa}+{{sp}}{zinsen}", fett=True)
    tilgung = zeile("./. Tilgung Kredit", f"-{prg('prg_Tilgung')}")
    zeile("= Cash Flow", f"{{sp}}{erg}-{{sp}}{afa}+{{sp}}{tilgung}", fett=True)


def blatt_sonderbereich(wb, modell: Modell) -> None:
    ws = wb.create_sheet(blattname(SONDERBEREICH, wb.sheetnames))
    ws["A1"] = "Sonderbereich Verkauf und Kauf"
    ws["A1"].font = FONT_TITEL
    ws["A2"] = ("Steuerbilanz; Ausgangsfall mit Zins und Tilgung der Darlehen des Objekts, "
                "Alternative mit den Krediten der Neuobjekte, das Darlehen des verkauften "
                "Objekts wird aus dem Erlös abgelöst. Neuobjekte "
                "zählen beim Verkauf ihrer Quelle 1. Je Verkauf "
                "eine Ergebnissicht und eine Detailsicht; Steuern ca. = Ergebnis × "
                "Grenzsteuersatz, ohne Verlustvortrag. Die genaue Steuer je Jahr steht im "
                "Blatt Liquidität.")
    b = _Block(ws, 4)
    b.kopf("Endvermögen nach latenter Steuer am Ende des Rasters", "A Plan (§ 6b)",
           "B sofort versteuern", "C versteuern und kaufen", "Baseline halten")
    b.zeile_(None, "Endvermögen", "=INDEX(vg_A,1)", "=INDEX(vg_B,1)", "=INDEX(vg_C,1)",
             "=INDEX(vg_Baseline,1)", fett=True)
    b.zeile_(None, "Differenz zu A", None, "=$B${0}-C{0}".format(b.zeile - 1),
             "=$B${0}-D{0}".format(b.zeile - 1), "=$B${0}-E{0}".format(b.zeile - 1))
    b.leer()

    vorgaenge = min(max(len(modell.verkaeufe), MIN_VORGAENGE), modell.kapazitaet.verkaeufe)
    for n in range(1, vorgaenge + 1):
        b.adr = {}
        _vorgang(b, n)
    neu = min(max(len(modell.neuobjekte), MIN_VORGAENGE), modell.kapazitaet.neuobjekte)
    _kauf(b, neu)
    b.zeile_(None, f"Weitere Vorgänge: das Blatt zeigt {vorgaenge} Verkäufe und {neu} "
             "Neuobjekte, jeweils die ersten Zeilen der Eingabeblätter. Mehr beim nächsten "
             "Generieren.", fmt="@")

    ws.column_dimensions["A"].width = 58
    for spalte in range(2, max(neu, 4) + 2):
        ws.column_dimensions[get_column_letter(spalte)].width = 22
    ws.freeze_panes = "B4"
