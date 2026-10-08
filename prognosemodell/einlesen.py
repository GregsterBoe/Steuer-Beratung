"""Einleseschicht für die Kanzlei-Excel mit einem Blatt je Kostenstelle.

Jedes Blatt folgt der DATEV-BWA Form 01: Spalte B trägt die BWA-Zeilennummer,
die Kopfzeile die Spaltenjahre. Gelesen wird die Jahresspalte des Basisjahrs,
gesucht wird über die BWA-Nummer, nie über die Zeilenposition.

Ändert sich das Quellformat, wird nur dieses Modul angepasst (Projektplan, Abschnitt 16).
"""

import dataclasses
import re
import unicodedata
from difflib import SequenceMatcher
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

from .modelle import (AFA_DEGRESSIV, AFA_LINEAR, ART_BGA, ART_FINANZ, ART_GEBAEUDE, ART_GUB,
                      ART_IM_BAU, ART_SONSTIGE, METHODE_KEINE, ZUORDNUNG_BEZEICHNUNG,
                      ZUORDNUNG_KOST1, ZUORDNUNG_MEHRDEUTIG, Anlage, Objekt)
from .vorlagen import NEUKAUF_MARKE

# Feste Stellen im Blatt
ZELLE_ID = "B2"        # Kostenstelle, z. B. "KSt 1"
ZELLE_NAME = "C2"      # Objektbezeichnung, z. B. "KC 24+26"
KOPFZEILE = 4          # B "Nr.", danach Jahres- und Monatsspalten
KENNUNG = "Nr."        # Inhalt von B4; nur Blätter mit dieser Kennung sind Kostenstellen
SPALTE_NR = 2          # B
# Jahreskopf: Zahl (2026) oder Text "Jahr 2026" / "Plan 2027"; Monatsspalten tragen ein Datum
JAHRESKOPF = re.compile(r"^(?:(?:jahr|plan)\s+)?(\d{4})$", re.IGNORECASE)
# Summenblatt über alle Kostenstellen: B2 nur "KSt" ohne Nummer oder C2 "Alle Objekte"
SUMMENBLATT_ID = "kst"
SUMMENBLATT_NAME = "alle objekte"

# BWA-Zeilen je Modellfeld; fehlende Zeilen zählen als 0
BWA_MIETE = (1020,)
BWA_EINNAHMEN = (1090,)
BWA_ERHALTUNG = (1250,)
BWA_AUSGABEN = (1100, 1120, 1140, 1150, 1160, 1180, 1200, 1220, 1260)
BWA_ABSCHREIBUNG = (1240,)  # Basisjahr: Abgleich; Planjahre mit Wert: AfA-Plan
# Neukauf-Kostenstellen: Position im Blatt Neukauf-KSt -> BWA-Zeilen
NEUKAUF_POSITIONEN = {"miete": BWA_MIETE, "einnahmen": BWA_EINNAHMEN,
                      "erhaltung": BWA_ERHALTUNG, "ausgaben": BWA_AUSGABEN}


# Aufschlüsselung der Abschreibungen unter der BWA (ohne BWA-Nr., Beschriftung in Spalte C):
# Block „Buchwert, JE“ mit einer Zeile je Anlagengruppe, Block „Abschreibungen JW“ mit der
# Jahres-AfA je Gruppe; „Abschreibungen MW“ (Monatswerte) beendet die Blöcke.
# Gelesen wird die Spalte des Jahres Stand Anlagenverzeichnis (Wj-Ende, im Muster F).
SPALTE_TEXT = 3        # C
BLOCK_BUCHWERT = "buchwert"
BLOCK_AFA_JAHR = "abschreibungen jw"
BLOCK_ENDE = "abschreibungen mw"


class EinleseFehler(ValueError):
    pass


@dataclass
class LaufendeWerte:
    """Werte eines Kostenstellenblatts im Basisjahr."""
    blatt: str
    objekt_id: str
    name: Optional[str]
    miete: float
    weitere_einnahmen: float
    erhaltung: float
    weitere_ausgaben: float
    abschreibung: float
    # Ist-Werte je BWA-Nr. und Spalte im Ausgabelayout (vorlagen: F, G Vorjahre,
    # H–S Monate, T Basisjahr); für das BWA-Blatt der Mappe
    ist: dict = field(default_factory=dict)
    # Aufschlüsselung der Abschreibungen je Anlagengruppe (Buchwert Stand, Jahres-AfA)
    anlagen: list = field(default_factory=list)
    # schon geplante AfA (BWA 1240) je Jahr nach dem Basisjahr, etwa bis 2046 fortgeschrieben
    afa_plan: dict = field(default_factory=dict)
    # Neukauf-Kostenstelle (Blatt hinter „KSt 9999“): kein Bestandsobjekt, sondern Planwerte
    # für ein Neuobjekt; jahre = {Position: {Jahr: Wert}} ab dem Basisjahr, siehe NEUKAUF_POSITIONEN
    neukauf: bool = False
    jahre: dict = field(default_factory=dict)
    # Neukauf: alle BWA-Zeilen der Planspalten {BWA-Nr.: {Jahr: Wert}} für das Blatt der Mappe
    plan: dict = field(default_factory=dict)


def _kopfjahr(wert) -> Optional[int]:
    """Jahr aus einem Spaltenkopf; None für Monatsspalten (Datum) und sonstige Köpfe."""
    if isinstance(wert, bool):
        return None
    if isinstance(wert, (int, float)):
        return int(wert)
    if isinstance(wert, str):
        treffer = JAHRESKOPF.match(wert.strip())
        return int(treffer.group(1)) if treffer else None
    return None


def _jahresspalte(ws, basisjahr: int) -> Optional[int]:
    for zelle in ws[KOPFZEILE]:
        if _kopfjahr(zelle.value) == basisjahr:
            return zelle.column
    return None


def _bwa_zeilen(ws) -> dict:
    zeilen = {}
    for zeile in range(KOPFZEILE + 1, ws.max_row + 1):
        nr = ws.cell(row=zeile, column=SPALTE_NR).value
        if isinstance(nr, (int, float)):
            zeilen[int(nr)] = zeile
    return zeilen


def _summe(ws, ws_formeln, zeilen: dict, nummern: tuple, spalte: int) -> float:
    summe = 0.0
    for nr in nummern:
        if nr not in zeilen:
            continue
        wert = ws.cell(row=zeilen[nr], column=spalte).value
        if wert is None:
            formel = ws_formeln.cell(row=zeilen[nr], column=spalte).value
            if isinstance(formel, str) and formel.startswith("="):
                raise EinleseFehler(
                    f"Blatt {ws.title!r}, BWA {nr}: Formel ohne berechneten Wert. "
                    "Datei in Excel öffnen, speichern und erneut einlesen.")
            continue
        if not isinstance(wert, (int, float)):
            raise EinleseFehler(f"Blatt {ws.title!r}, BWA {nr}: kein Zahlenwert ({wert!r})")
        summe += wert
    return summe


def _ist_spalten(ws, basisjahr: int) -> dict:
    """Spalte im Blatt -> Spalte im Ausgabelayout, für Vorjahre, Monate und Basisjahr."""
    from .vorlagen import SPALTE_JAHR, SPALTE_MONATE, SPALTE_VORJAHRE
    ziel = {}
    for zelle in ws[KOPFZEILE]:
        wert = zelle.value
        if isinstance(wert, date):  # Monatsspalte; datetime ist auch ein date
            if wert.year == basisjahr:
                ziel[zelle.column] = SPALTE_MONATE + wert.month - 1
            continue
        jahr = _kopfjahr(wert)
        if jahr in (basisjahr - 2, basisjahr - 1):
            ziel[zelle.column] = SPALTE_VORJAHRE + jahr - (basisjahr - 2)
        elif jahr == basisjahr:
            ziel[zelle.column] = SPALTE_JAHR
    return ziel


def _ist_werte(ws, zeilen: dict, basisjahr: int) -> dict:
    spalten = _ist_spalten(ws, basisjahr)
    ist = {}
    for nr, zeile in zeilen.items():
        werte = {ziel: ws.cell(row=zeile, column=quelle).value
                 for quelle, ziel in spalten.items()}
        werte = {k: v for k, v in werte.items()
                 if isinstance(v, (int, float)) and not isinstance(v, bool)}
        if werte:
            ist[nr] = werte
    return ist


def _plan_werte(ws, zeilen: dict, nummern: tuple, basisjahr: int) -> dict:
    """Jahr -> Summe der BWA-Zeilen in den Jahresspalten nach dem Basisjahr.

    Nur Zahlen zählen, auch 0; leere Zellen bleiben der Fortschreibung des Modells.
    """
    plan = {}
    for zelle in ws[KOPFZEILE]:
        jahr = _kopfjahr(zelle.value)
        if jahr is None or jahr <= basisjahr:
            continue
        werte = [_zahl(ws.cell(row=zeilen[nr], column=zelle.column).value)
                 for nr in nummern if nr in zeilen]
        werte = [w for w in werte if w is not None]
        if werte:
            plan[jahr] = sum(werte)
    return plan


def _ist_marke(ws) -> bool:
    """Blatt der Kostenstelle „KSt 9999“: dahinter folgen die Neukauf-Kostenstellen."""
    def norm(wert):
        return " ".join(str(wert or "").split()).lower()
    return NEUKAUF_MARKE.lower() in (norm(ws.title), norm(ws[ZELLE_ID].value))


def _formatfehler(ws, basisjahr: int, neukauf: bool = False) -> Optional[str]:
    """Grund, warum das Blatt kein lesbares Kostenstellenblatt ist; None = Format passt.
    Neukauf-Kostenstellen brauchen keine Spalte des Basisjahrs."""
    kennung = ws.cell(row=KOPFZEILE, column=SPALTE_NR).value
    if not (isinstance(kennung, str) and kennung.strip().lower() == KENNUNG.lower()):
        return f"kein {KENNUNG!r} in Zeile {KOPFZEILE}, Spalte B"
    objekt_id = ws[ZELLE_ID].value
    if objekt_id is None or str(objekt_id).strip() == "":
        return f"keine Kostenstelle in {ZELLE_ID}"
    name = ws[ZELLE_NAME].value
    if str(objekt_id).strip().lower() == SUMMENBLATT_ID or \
            (isinstance(name, str) and name.strip().lower() == SUMMENBLATT_NAME):
        return "Summenblatt aller Kostenstellen"
    if _jahresspalte(ws, basisjahr) is None and not neukauf:
        return f"keine Spalte {basisjahr} in Zeile {KOPFZEILE}"
    return None


def _zahl(wert) -> Optional[float]:
    if isinstance(wert, bool) or not isinstance(wert, (int, float)):
        return None
    return float(wert)


def _abschreibungsbloecke(ws, objekt_id: str, stand: int) -> list:
    """Anlagengruppen aus der Aufschlüsselung unter der BWA; leer, wenn sie fehlt.

    Buchwert und Jahres-AfA einer Gruppe werden über die Beschriftung verbunden: gleich,
    sonst eine Beschriftung Anfang der anderen („TG“ zu „TG 24“), sonst die Reihenfolge.
    Gruppen ohne Buchwert und ohne AfA (etwa „sonstige“ leer) entfallen.
    """
    spalte = _jahresspalte(ws, stand)
    if spalte is None:
        return []
    bloecke, block = {BLOCK_BUCHWERT: [], BLOCK_AFA_JAHR: []}, None
    for zeile in range(KOPFZEILE + 1, ws.max_row + 1):
        text = ws.cell(row=zeile, column=SPALTE_TEXT).value
        if not isinstance(text, str) or not text.strip():
            continue  # Zwischensummen ohne Beschriftung
        if ws.cell(row=zeile, column=SPALTE_NR).value not in (None, ""):
            block = None  # BWA-Zeile
            continue
        kennung = text.strip().lower()
        if kennung.startswith(BLOCK_ENDE):
            block = None
        elif kennung.startswith(BLOCK_AFA_JAHR):
            block = BLOCK_AFA_JAHR
        elif kennung.startswith(BLOCK_BUCHWERT):
            block = BLOCK_BUCHWERT
        elif block:
            bloecke[block].append((text.strip(), _zahl(ws.cell(row=zeile, column=spalte).value)))
    buchwerte, afas = bloecke[BLOCK_BUCHWERT], list(bloecke[BLOCK_AFA_JAHR])

    def partner(i, name):
        for kriterium in (lambda a: a.lower() == name.lower(),
                          lambda a: name.lower().startswith(a.lower())
                          or a.lower().startswith(name.lower())):
            treffer = [j for j, (a, _) in enumerate(afas) if a is not None and kriterium(a)]
            if len(treffer) == 1:
                return treffer[0]
        return i if i < len(afas) and afas[i][0] is not None else None

    gruppen = []
    for i, (name, buchwert) in enumerate(buchwerte):
        j = partner(i, name)
        afa = afas[j][1] if j is not None else None
        if j is not None:
            afas[j] = (None, None)   # jede AfA-Zeile nur einmal
        if not buchwert and not afa:
            continue
        gruppen.append(Anlage(
            nr=f"{objekt_id} {name}", bezeichnung=name, objekt_id=objekt_id, art=ART_GEBAEUDE,
            methode=AFA_LINEAR, bw_stand=buchwert or 0.0, afa=afa or 0.0))
    return gruppen


def lese_kostenstellen(pfad, basisjahr: int, stand: int = None) -> tuple:
    """Liest alle Kostenstellenblätter der Kanzlei-Excel.

    Blätter mit anderem Format (etwa Annahmen oder Übersichten) werden übersprungen.
    Blätter hinter der Kostenstelle „KSt 9999“ sind Neukauf-Kostenstellen (neukauf=True):
    Planwerte je Jahr für Neuobjekte, keine Bestandsobjekte.
    stand: Jahr der Aufschlüsselung der Abschreibungen (Wj-Ende), Standard Basisjahr − 1.
    Rückgabe: (je Kostenstellenblatt eine LaufendeWerte-Zeile,
    Liste (Blatttitel, Grund) der übersprungenen Blätter).
    """
    stand = basisjahr - 1 if stand is None else stand
    pfad = Path(pfad)
    werte_wb = load_workbook(pfad, data_only=True)   # berechnete Werte
    formel_wb = load_workbook(pfad, data_only=False)  # nur für Fehlermeldungen
    ergebnis, gesehen, uebersprungen = [], {}, []
    neukauf = False   # ab dem Blatt hinter „KSt 9999“
    for ws in werte_wb.worksheets:
        grund = _formatfehler(ws, basisjahr, neukauf)
        marke = _ist_marke(ws)
        if grund:
            uebersprungen.append((ws.title, grund))
            neukauf = neukauf or marke
            continue
        objekt_id = str(ws[ZELLE_ID].value).strip()
        if objekt_id in gesehen:
            raise EinleseFehler(
                f"Kostenstelle {objekt_id!r} doppelt (Blätter {gesehen[objekt_id]!r} "
                f"und {ws.title!r})")
        gesehen[objekt_id] = ws.title
        name = ws[ZELLE_NAME].value
        spalte = _jahresspalte(ws, basisjahr)
        zeilen = _bwa_zeilen(ws)
        wf = formel_wb[ws.title]

        def summe(nummern):
            return _summe(ws, wf, zeilen, nummern, spalte) if spalte else 0.0

        if neukauf:
            ergebnis.append(LaufendeWerte(
                blatt=ws.title, objekt_id=objekt_id,
                name=str(name).strip() if name is not None else None,
                miete=summe(BWA_MIETE), weitere_einnahmen=summe(BWA_EINNAHMEN),
                erhaltung=summe(BWA_ERHALTUNG), weitere_ausgaben=summe(BWA_AUSGABEN),
                abschreibung=summe(BWA_ABSCHREIBUNG), neukauf=True,
                ist=_ist_werte(ws, zeilen, basisjahr),
                jahre={pos: _plan_werte(ws, zeilen, nummern, basisjahr - 1)
                       for pos, nummern in NEUKAUF_POSITIONEN.items()},
                plan={nr: w for nr in zeilen
                      if (w := _plan_werte(ws, zeilen, (nr,), basisjahr))}))
            continue
        neukauf = marke
        ergebnis.append(LaufendeWerte(
            blatt=ws.title,
            objekt_id=objekt_id,
            name=str(name).strip() if name is not None else None,
            miete=summe(BWA_MIETE),
            weitere_einnahmen=summe(BWA_EINNAHMEN),
            erhaltung=summe(BWA_ERHALTUNG),
            weitere_ausgaben=summe(BWA_AUSGABEN),
            abschreibung=summe(BWA_ABSCHREIBUNG),
            ist=_ist_werte(ws, zeilen, basisjahr),
            anlagen=_abschreibungsbloecke(ws, objekt_id, stand),
            afa_plan=_plan_werte(ws, zeilen, BWA_ABSCHREIBUNG, basisjahr),
        ))
    if not ergebnis:
        gruende = "; ".join(f"{t!r}: {g}" for t, g in uebersprungen)
        raise EinleseFehler(f"{pfad.name}: kein lesbares Kostenstellenblatt ({gruende})")
    return ergebnis, uebersprungen


def zusammenfuehren(stammdaten: list, laufende: list) -> list:
    """Laufende Werte in die Objekte übernehmen (Schlüssel: ObjektID).

    Kostenstellen ohne Stammdaten werden als neue Objekte angelegt; deren
    steuerliche Pflichtfelder bleiben leer und die Statusspalte meldet sie.
    Objekte ohne Kostenstellenblatt bleiben unverändert.
    """
    nach_id = {o.objekt_id: o for o in stammdaten}
    ergebnis = list(stammdaten)
    for lw in laufende:
        if lw.neukauf:
            continue
        felder = dict(miete=lw.miete, weitere_einnahmen=lw.weitere_einnahmen,
                      erhaltung=lw.erhaltung, weitere_ausgaben=lw.weitere_ausgaben,
                      afa_bwa=lw.abschreibung)   # 0 = keine AfA mehr
        if lw.objekt_id in nach_id:
            alt = nach_id[lw.objekt_id]
            neu = dataclasses.replace(alt, name=alt.name or lw.name, **felder)
            ergebnis[ergebnis.index(alt)] = neu
        else:
            ergebnis.append(Objekt(objekt_id=lw.objekt_id, name=lw.name, **felder))
    return ergebnis


# --- Anlagenverzeichnis (DATEV-Export „Inventarübersicht“, Projektplan Abschnitt 21) ---
#
# Eine Zeile je Anlagegut, Spalten über die Kopfzeile gesucht. Die Kostenstelle steht in
# KOST1; sie verweist auf das Kostenstellenblatt (KOST1 1 = „KSt 1“).

INVENTAR_SPALTEN = {
    "konto": "Konto", "nr": "Inventar", "bezeichnung": "Inventarbezeichnung",
    "datum": "AHK-Datum", "ahk": "AHK Wj-Ende", "bw_stand": "Buchw. Wj-Ende",
    "afa_ende": "N-AfA Wj-Ende", "afa_beginn": "N-AfA Wj-Beginn",
    "sonder_ende": "S-Abschr. Wj-Ende", "sonder_beginn": "S-Abschr. Wj-Beginn",
    "afa_art": "AfA-Art", "satz": "AfA-%", "kost1": "KOST1", "abgang": "Abgang",
}
INVENTAR_PFLICHT = ("nr", "bw_stand", "kost1")
# Art je Kontenbereich (SKR04, 3- oder 4-stellig ohne führende Null); im Blatt änderbar
KONTEN_ART = [(200, 239, ART_GUB), (240, 399, ART_GEBAEUDE), (400, 699, ART_BGA),
              (700, 799, ART_IM_BAU), (800, 999, ART_FINANZ)]
# AfA-Art lt. DATEV geht vor dem Konto
AFA_ART_ART = [("lin.geb", ART_GEBAEUDE), ("anlag./bau", ART_IM_BAU), ("finanzanl", ART_FINANZ)]
AFA_ART_OHNE = ("keine afa", "anlag./bau", "finanzanl")
KST_PRAEFIX = "KSt "
STAND_IM_NAMEN = re.compile(r"(?<!\d)(20\d\d)(?!\d)")


def art_aus(konto, afa_art) -> str:
    text = (afa_art or "").strip().lower()
    for anfang, art in AFA_ART_ART:
        if text.startswith(anfang):
            return art
    if isinstance(konto, (int, float)):
        for von, bis, art in KONTEN_ART:
            if von <= konto <= bis:
                return art
    return ART_SONSTIGE


def methode_aus(afa_art, satz) -> str:
    text = (afa_art or "").strip().lower()
    if not text or text.startswith(AFA_ART_OHNE) or not satz:
        return METHODE_KEINE
    return AFA_DEGRESSIV if "degr" in text else AFA_LINEAR


def _datum(wert) -> Optional[date]:
    if isinstance(wert, datetime):
        return wert.date()
    if isinstance(wert, date):
        return wert
    if isinstance(wert, str) and wert.strip():
        try:
            return datetime.strptime(wert.strip(), "%d.%m.%Y").date()
        except ValueError:
            raise EinleseFehler(f"Datum nicht lesbar: {wert!r}") from None
    return None


def _text(wert) -> Optional[str]:
    if wert is None:
        return None
    if isinstance(wert, float) and wert.is_integer():
        wert = int(wert)
    text = str(wert).strip()
    return text or None


def stand_aus_dateiname(pfad) -> Optional[int]:
    """Jahr im Dateinamen, etwa Inventar_2025.xlsx; None, wenn keins oder mehrere."""
    treffer = set(STAND_IM_NAMEN.findall(Path(pfad).stem))
    return int(treffer.pop()) if len(treffer) == 1 else None


def lese_inventar(pfad) -> tuple:
    """Liest das Anlagenverzeichnis. Rückgabe: (Anlagen, Anzahl abgegangener Anlagen).

    Abgegangene Anlagen (Datum in Abgang) entfallen. KOST1 bleibt zunächst roh;
    die ObjektID setzt ordne_anlagen_zu.
    """
    pfad = Path(pfad)
    wb = load_workbook(pfad, data_only=True)
    for ws in wb.worksheets:
        for kopf in range(1, min(ws.max_row, 10) + 1):
            texte = {str(c.value).strip(): c.column for c in ws[kopf] if c.value is not None}
            spalten = {k: texte.get(v) for k, v in INVENTAR_SPALTEN.items()}
            if all(spalten[k] for k in INVENTAR_PFLICHT):
                break
        else:
            continue
        break
    else:
        pflicht = ", ".join(INVENTAR_SPALTEN[k] for k in INVENTAR_PFLICHT)
        raise EinleseFehler(f"{pfad.name}: keine Inventarübersicht (Kopf mit {pflicht})")

    def wert(zeile, key):
        spalte = spalten[key]
        return ws.cell(row=zeile, column=spalte).value if spalte else None

    def zahl(zeile, key):
        w = wert(zeile, key)
        if w in (None, ""):
            return None
        if not isinstance(w, (int, float)) or isinstance(w, bool):
            raise EinleseFehler(f"{pfad.name}, Zeile {zeile}, {INVENTAR_SPALTEN[key]}: "
                                f"kein Zahlenwert ({w!r})")
        return float(w)

    anlagen, abgang, gesehen = [], 0, {}
    for zeile in range(kopf + 1, ws.max_row + 1):
        nr = _text(wert(zeile, "nr"))
        if nr is None:
            continue
        if nr in gesehen:
            raise EinleseFehler(f"{pfad.name}: Inventar {nr} doppelt (Zeilen {gesehen[nr]} "
                                f"und {zeile})")
        gesehen[nr] = zeile
        if _text(wert(zeile, "abgang")):
            abgang += 1
            continue
        konto = zahl(zeile, "konto")
        afa_art = _text(wert(zeile, "afa_art"))
        satz = zahl(zeile, "satz")
        satz = satz / 100 if satz is not None else None
        afa_stand = sum(v or 0.0 for v in (zahl(zeile, "afa_ende"), zahl(zeile, "sonder_ende"))) \
            - sum(v or 0.0 for v in (zahl(zeile, "afa_beginn"), zahl(zeile, "sonder_beginn")))
        anlagen.append(Anlage(
            nr=nr, bezeichnung=_text(wert(zeile, "bezeichnung")),
            konto=int(konto) if konto is not None else None,
            kost1=_text(wert(zeile, "kost1")), datum=_datum(wert(zeile, "datum")),
            ahk=zahl(zeile, "ahk"), bw_stand=zahl(zeile, "bw_stand") or 0.0,
            afa_art=afa_art, satz=satz, art=art_aus(konto, afa_art),
            methode=methode_aus(afa_art, satz), afa_stand=round(afa_stand, 2)))
    return anlagen, abgang


# Zuordnung über die Inventarbezeichnung, wenn KOST1 fehlt: die Wörter der Bezeichnung
# werden einzeln mit ObjektID und Name der Kostenstelle verglichen (unscharf, damit
# „Musterstr.“ zu „Musterstraße“ und kleine Tippfehler passen). Wörter für die Art der
# Anlage tragen nichts zum Objekt bei.
BEZ_FUELLWOERTER = {
    "und", "u", "der", "die", "das", "den", "des", "dem", "in", "im", "am", "an", "zu", "zum",
    "zur", "mit", "fuer", "von", "vom", "bei", "auf", "aus", "nach", "ohne", "nr", "kst",
    "grund", "boden", "gub", "gb", "grundstueck", "grundstuecke", "gebaeude", "wohngebaeude",
    "geschaeftsgebaeude", "buerogebaeude", "haus", "wohnhaus", "wohnbau", "bau", "anbau",
    "neubau", "umbau", "aussenanlage", "aussenanlagen", "anlage", "anlagen", "anteil",
    "objekt", "kostenstelle", "kostenst", "wohnung", "wohnungen", "whg", "einbau", "einbauten",
    "sanierung", "modernisierung", "erweiterung", "teil", "gesamt", "alle",
}
BEZ_AEHNLICH = 0.8        # Mindestähnlichkeit zweier Wörter (difflib-Quote)
BEZ_ABSTAND = 0.5         # Vorsprung des besten Objekts vor dem zweitbesten
# Kostenstellennummer in der Bezeichnung: „KSt 5“, „Kostenstelle 5“, „Kostenst. 5“
BEZ_KST = re.compile(r"\b(?:kst|kostenst(?:elle)?)\.?\s*(\d+)\b", re.IGNORECASE)


def _woerter(text) -> tuple:
    """(Wörter, Zahlen) eines Texts: klein, Umlaute ausgeschrieben, „straße“ zu „str“."""
    text = str(text or "").lower()
    for alt, neu in (("ä", "ae"), ("ö", "oe"), ("ü", "ue"), ("ß", "ss")):
        text = text.replace(alt, neu)
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    woerter, zahlen = set(), set()
    for t in re.findall(r"[a-z]+|\d+", text):
        if t.isdigit():
            zahlen.add(t.lstrip("0") or "0")
            continue
        t = re.sub(r"(str)asse$", r"\1", t)
        if t not in BEZ_FUELLWOERTER and len(t) >= 2:
            woerter.add(t)
    return woerter, zahlen


def _aehnlich(a: str, b: str) -> float:
    if a == b:
        return 1.0
    if min(len(a), len(b)) < 4:       # kurze Kürzel nur exakt („KC“, „OG“)
        return 0.0
    if a.startswith(b) or b.startswith(a):
        return 0.9
    quote = SequenceMatcher(None, a, b).ratio()
    return quote if quote >= BEZ_AEHNLICH else 0.0


def treffer_bezeichnung(bezeichnung, objekte: dict) -> list:
    """Objekte, deren ObjektID oder Name zur Bezeichnung passen, als Liste
    (Punkte, ObjektID, passende Wörter) absteigend; ohne passendes Wort leer.
    objekte: ObjektID -> Liste von Texten (ID, Name)."""
    woerter, zahlen = _woerter(bezeichnung)
    ergebnis = []
    for oid, texte in objekte.items():
        o_woerter, o_zahlen = set(), set()
        for t in texte:
            w, z = _woerter(t)
            o_woerter |= w
            o_zahlen |= z
        passend, punkte = [], 0.0
        for w in sorted(woerter):
            bestes = max((_aehnlich(w, o) for o in o_woerter), default=0.0)
            if bestes:
                punkte += bestes
                passend.append(w)
        if not passend:
            continue
        gleiche_zahlen = sorted(zahlen & o_zahlen)
        punkte += 0.5 * len(gleiche_zahlen)
        ergebnis.append((round(punkte, 3), oid, passend + gleiche_zahlen))
    return sorted(ergebnis, key=lambda t: (-t[0], str(t[1])))


def _objekttexte(objekte) -> dict:
    """ObjektID -> Texte für den Abgleich; objekte als ObjektIDs oder Objekte mit Name."""
    texte = {}
    for o in objekte:
        oid = getattr(o, "objekt_id", o)
        if oid:
            texte.setdefault(oid, [str(oid)])
            name = getattr(o, "name", None)
            if name:
                texte[oid].append(str(name))
    return texte


def ordne_anlagen_zu(anlagen: list, objekte) -> list:
    """ObjektID aus KOST1: die ObjektID mit derselben Endnummer (1 zu „KSt 1“),
    sonst „KSt <KOST1>“. Ohne KOST1 über die Inventarbezeichnung (treffer_bezeichnung),
    gekennzeichnet in Zuordnung; passt nichts eindeutig, bleibt die Anlage ohne Objekt.
    objekte: ObjektIDs oder Objekte (dann zählt auch der Name)."""
    texte = _objekttexte(objekte)
    nach_nummer = {}
    for oid in texte:
        treffer = re.search(r"(\d+)\s*$", str(oid))
        if treffer:
            nach_nummer.setdefault(int(treffer.group(1)), []).append(oid)

    def aus_kost(kost1):
        kost = int(kost1) if kost1.isdigit() else None
        kandidaten = nach_nummer.get(kost, []) if kost is not None else []
        return kandidaten[0] if len(kandidaten) == 1 else (
            f"{KST_PRAEFIX}{kost}" if kost is not None else kost1)

    ergebnis = []
    for a in anlagen:
        oid, zuordnung = a.objekt_id, a.zuordnung
        if oid is None and a.kost1:
            oid, zuordnung = aus_kost(a.kost1), ZUORDNUNG_KOST1
        elif oid is None and a.bezeichnung:
            kst = BEZ_KST.search(a.bezeichnung)
            treffer = treffer_bezeichnung(a.bezeichnung, texte)
            if kst:
                oid = aus_kost(kst.group(1))
                zuordnung = f"{ZUORDNUNG_BEZEICHNUNG}: KSt {kst.group(1)}"
            elif len(treffer) == 1 or (treffer and treffer[0][0] - treffer[1][0]
                                       >= BEZ_ABSTAND - 1e-9):
                _, oid, woerter = treffer[0]
                zuordnung = f"{ZUORDNUNG_BEZEICHNUNG}: {', '.join(woerter)}"
            elif treffer:
                beste = [t[1] for t in treffer if treffer[0][0] - t[0] < BEZ_ABSTAND - 1e-9]
                zuordnung = f"{ZUORDNUNG_MEHRDEUTIG}: {', '.join(beste)}"
        ergebnis.append(dataclasses.replace(a, objekt_id=oid, zuordnung=zuordnung))
    return ergebnis


def anlagen_zusammenfuehren(inventar: list, laufende: list) -> tuple:
    """Anlagen für das Blatt Anlagen: das Anlagenverzeichnis, dazu für Kostenstellen ohne
    Anlage darin die Gruppen aus der Aufschlüsselung im Kostenstellenblatt.

    Rückgabe: (Anlagen, Abgleich je Kostenstelle mit beiden Quellen als Liste
    (ObjektID, Buchwert lt. Verzeichnis, lt. Kostenstelle, AfA p. a. lt. Kostenstelle)).
    """
    mit_inventar = {a.objekt_id for a in inventar if a.objekt_id}
    ergebnis, abgleich = list(inventar), []
    for lw in laufende:
        if not lw.anlagen:
            continue
        if lw.objekt_id in mit_inventar:
            bw = sum(a.bw_stand or 0.0 for a in inventar if a.objekt_id == lw.objekt_id
                     and a.art in (ART_GEBAEUDE, ART_BGA, ART_IM_BAU))
            abgleich.append((lw.objekt_id, bw, sum(g.bw_stand for g in lw.anlagen),
                             sum(g.afa for g in lw.anlagen)))
        else:
            ergebnis.extend(lw.anlagen)
    return ergebnis, abgleich
