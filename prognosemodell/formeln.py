"""Formel-Bausteine.

Formeln werden in englischer Excel-Syntax geschrieben (so legt openpyxl sie ab);
Excel zeigt sie in der deutschen Oberfläche automatisch als WENN, ZÄHLENWENN usw.
"""

from openpyxl.utils import get_column_letter

from .modelle import (AFA_DEGRESSIV, FEHLER, NEU_FELDER, NEU_SPALTEN, OBJEKT_FELDER, PROGNOSE_SPALTEN,
                      RUECKLAGE_JAHR_SPALTEN, RUECKLAGE_SPALTEN, STATUS_6B_UNZULAESSIG,
                      STATUS_OK, SZ_A, SZ_B, SZ_BASELINE, SZENARIEN, VERKAUF_FELDER,
                      VERKAUF_SPALTEN, WARNUNG, Szenario, aus_spalten, liq_spalten)


def spalte(key: str) -> str:
    """Spaltenbuchstabe eines Objektfelds."""
    for i, f in enumerate(OBJEKT_FELDER, start=1):
        if f.key == key:
            return get_column_letter(i)
    raise KeyError(key)


def status_objekt(zeile: int) -> str:
    """Plausibilitätsstatus je Objektzeile; leer, solange keine ObjektID steht."""
    id_ = f"${spalte('objekt_id')}{zeile}"
    pflicht = [f"${spalte(f.key)}{zeile}" for f in OBJEKT_FELDER if f.pflicht]
    kaufjahr = f"${spalte('kaufjahr')}{zeile}"
    restbw = f"${spalte('restbuchwert')}{zeile}"
    ak_geb = f"${spalte('ak_gebaeude')}{zeile}"
    return (
        f'=IF({id_}="","",'
        f'IF(COUNTA({",".join(pflicht)})<{len(pflicht)},"Pflichtfeld fehlt",'
        f'IF(COUNTIF(obj_ID,{id_})>1,"ObjektID doppelt",'
        f'IF({kaufjahr}>par_Basisjahr,"Kaufjahr nach Basisjahr",'
        f'IF({restbw}>{ak_geb},"Restbuchwert > AK Gebäude",'
        f'"OK")))))'
    )


# --- Prognosematrix (Etappe 2, Projektplan Abschnitt 8) ---
#
# Jede Zeile des Objektblatts hat einen festen Block aufeinanderfolgender
# Jahreszeilen; die Stammdaten kommen per INDEX aus genau dieser Objektzeile.
# Der Buchwert des Vorjahres steht deshalb direkt in der Zeile darüber; im
# ersten Prognosejahr ist es der Restbuchwert aus dem Objektblatt.


def pspalte(key: str) -> str:
    """Spaltenbuchstabe einer Prognosespalte."""
    for i, s in enumerate(PROGNOSE_SPALTEN, start=1):
        if s.key == key:
            return get_column_letter(i)
    raise KeyError(key)


def _p(key: str, zeile: int) -> str:
    return f"${pspalte(key)}{zeile}"


def _stamm(name: str, objekt_nr: int) -> str:
    """Stammdatum aus Zeile objekt_nr des Objektblatts (1 = erstes Objekt)."""
    return f"INDEX({name},{objekt_nr})"


def _leer_oder(zeile: int, ausdruck: str) -> str:
    return f'=IF({_p("id", zeile)}="","",{ausdruck})'


def _indexiert(name: str, satz: str, zeile: int, objekt_nr: int, aktiv: bool = True) -> str:
    ausdruck = f"{_stamm(name, objekt_nr)}*(1+{satz})^({_p('jahr', zeile)}-par_Basisjahr)"
    return f"{ausdruck}*{_p('aktiv', zeile)}" if aktiv else ausdruck


def _verkaufsjahr(zeile: int) -> str:
    """Verkaufsjahr laut Blatt Verkäufe; ohne Verkauf 9999, also nie."""
    treffer = f"INDEX(vk_Jahr,MATCH({_p('id', zeile)},vk_ID,0))"
    return f'IFERROR(IF({treffer}="",9999,{treffer}),9999)'


def prognose_zeile(zeile: int, objekt_nr: int, erstes_jahr: bool) -> dict:
    """Formeln einer Prognosezeile je Spaltenschlüssel.

    objekt_nr ist die Zeile im Objektblatt (1 = erstes Objekt). Leere Objektzeilen
    ergeben leere Prognosezeilen.
    """
    id_obj = _stamm("obj_ID", objekt_nr)
    bw_vor = (_stamm("obj_Restbuchwert", objekt_nr) if erstes_jahr
              else _p("buchwert", zeile - 1))
    afa_voll = f"{_stamm('obj_AKGebaeude', objekt_nr)}*{_stamm('obj_AfASatz', objekt_nr)}"

    def indexiert(name, satz, aktiv=True):
        return _leer_oder(zeile, _indexiert(name, satz, zeile, objekt_nr, aktiv))

    return {
        "id": f'=IF({id_obj}="","",{id_obj})',
        "jahr": "=par_Startjahr" if erstes_jahr else f"={_p('jahr', zeile - 1)}+1",
        # 0 ab dem Jahr nach dem Verkauf
        "aktiv": _leer_oder(zeile, f"IF({_p('jahr', zeile)}<={_verkaufsjahr(zeile)},1,0)"),
        "miete": indexiert("obj_MieteBasis", "par_Mietsteig"),
        "einnahmen": indexiert("obj_EinnBasis", "par_Mietsteig"),
        "erhaltung": indexiert("obj_ErhBasis", "par_Erhaltsteig"),
        "ausgaben": indexiert("obj_AusgBasis", "par_Kostensteig"),
        # keine AfA über den Restbuchwert hinaus
        "afa": _leer_oder(zeile, f"MIN({afa_voll},{bw_vor})*{_p('aktiv', zeile)}"),
        "buchwert": _leer_oder(zeile, f"MAX({bw_vor}-{_p('afa', zeile)},0)"),
        "ergebnis": _leer_oder(
            zeile,
            f"({_p('miete', zeile)}+{_p('einnahmen', zeile)}-{_p('erhaltung', zeile)}"
            f"-{_p('ausgaben', zeile)}-{_p('afa', zeile)})*{_p('aktiv', zeile)}"),
        "verkehrswert": indexiert("obj_Verkehrswert", "par_Wertsteig", aktiv=False),
        # Verkauf zum Jahresende: im Verkaufsjahr schon nicht mehr im Bestand
        "bestand": _leer_oder(zeile, f"IF({_p('jahr', zeile)}<{_verkaufsjahr(zeile)},1,0)"),
        "neu": _leer_oder(zeile, "0"),
        # G+B wird nicht abgeschrieben; ob das Objekt noch im Bestand ist, filtert die Auswertung
        "buchwert_gub": _leer_oder(zeile, _stamm("obj_AKGuB", objekt_nr)),
        # Buchwert ohne Verkauf: Restbuchwert minus volle AfA je Jahr seit dem Basisjahr
        "buchwert_halten": _leer_oder(
            zeile, f"MAX({_stamm('obj_Restbuchwert', objekt_nr)}"
                   f"-({_p('jahr', zeile)}-par_Basisjahr)*{afa_voll},0)"),
        # Szenarien B und C: Bestandsobjekte rechnen wie im Plan
        "mit_quelle": _leer_oder(zeile, "0"),
        "afa_ohne6b": _leer_oder(zeile, _p("afa", zeile)),
        "buchwert_ohne6b": _leer_oder(zeile, _p("buchwert", zeile)),
        "buchwert_gub_ohne6b": _leer_oder(zeile, _p("buchwert_gub", zeile)),
    }


def prognose_zeile_neu(zeile: int, neu_nr: int, erstes_jahr: bool) -> dict:
    """Prognosezeile eines Neuobjekts (Etappe 6); neu_nr = Zeile im Blatt Neuobjekte.

    Kauf zum Ende des Kaufjahrs: Im Kaufjahr steht der Buchwert auf der AfA-Basis
    und das Objekt ist im Bestand, Miete, Erhaltung und AfA laufen ab dem Folgejahr.
    Alles bleibt 0, solange die Zeile nicht im Modell ist (ne_Gueltig).
    """
    def ne(name):
        return f"INDEX({name},{neu_nr})"

    t, kj, gueltig = _p("jahr", zeile), ne("ne_Kaufjahr"), f'{ne("ne_Gueltig")}=1'
    aktiv = f"{_p('aktiv', zeile)}=1"

    def afa(basis, bw_key, afa_key):
        """AfA und Buchwert ab der AfA-Basis basis, Vorjahresbuchwert aus Spalte bw_key."""
        bw_vor = "0" if erstes_jahr else _p(bw_key, zeile - 1)
        # Degressiv vom Restbuchwert; linear über die Restnutzungsdauer, sobald das mehr ist
        # (§ 7 Abs. 5a Satz 4 EStG). Restnutzungsdauer zu Jahresbeginn, mindestens 1 Jahr.
        rnd = f'MAX(1/{ne("ne_AfASatz")}-({t}-{kj}-1),1)'
        degressiv = f"MIN(MAX({bw_vor}*par_AfADegressiv,{bw_vor}/{rnd}),{bw_vor})"
        linear = f'MIN({basis}*{ne("ne_AfASatz")},{bw_vor})'
        return (
            _leer_oder(zeile, f'IF({aktiv},IF({ne("ne_AfAMethode")}="{AFA_DEGRESSIV}",'
                              f'{degressiv},{linear}),0)'),
            _leer_oder(zeile, f'IF({gueltig},IF({t}<{kj},0,IF({t}={kj},{basis},'
                              f'MAX({bw_vor}-{_p(afa_key, zeile)},0))),0)'),
        )

    afa_plan, buchwert_plan = afa(ne("ne_AfABasis"), "buchwert", "afa")
    afa_ohne, buchwert_ohne = afa(ne("ne_AKGebNeu"), "buchwert_ohne6b", "afa_ohne6b")

    def ab_kauf(basis, satz):
        return _leer_oder(zeile, f"IF({aktiv},{basis}*(1+{satz})^({t}-{kj}),0)")

    return {
        "id": f'=IF({ne("ne_ID")}="","",{ne("ne_ID")})',
        "jahr": "=par_Startjahr" if erstes_jahr else f"={_p('jahr', zeile - 1)}+1",
        "aktiv": _leer_oder(zeile, f"IF({gueltig},IF({t}>{kj},1,0),0)"),
        "miete": ab_kauf(f'{ne("ne_Kaufpreis")}*{ne("ne_Mietrendite")}', "par_Mietsteig"),
        "einnahmen": _leer_oder(zeile, "0"),
        "erhaltung": ab_kauf(f'{ne("ne_Kaufpreis")}*{ne("ne_ErhQuote")}', "par_Erhaltsteig"),
        "ausgaben": _leer_oder(zeile, "0"),
        "afa": afa_plan,
        "buchwert": buchwert_plan,
        "ergebnis": _leer_oder(
            zeile,
            f"({_p('miete', zeile)}+{_p('einnahmen', zeile)}-{_p('erhaltung', zeile)}"
            f"-{_p('ausgaben', zeile)}-{_p('afa', zeile)})*{_p('aktiv', zeile)}"),
        # Wert = Kaufpreis, ab dem Kaufjahr mit der Wertsteigerung fortgeschrieben
        "verkehrswert": _leer_oder(zeile, f'IF({gueltig},IF({t}>={kj},'
                                          f'{ne("ne_Kaufpreis")}*(1+par_Wertsteig)^({t}-{kj}),0),0)'),
        "bestand": _leer_oder(zeile, f"IF({gueltig},IF({t}>={kj},1,0),0)"),
        "neu": _leer_oder(zeile, "1"),
        "buchwert_gub": _leer_oder(zeile, f'IF({gueltig},IF({t}>={kj},{ne("ne_AKGuB")},0),0)'),
        # gehört nicht zur Baseline
        "buchwert_halten": _leer_oder(zeile, "0"),
        # Szenarien B und C: ohne Übertragung der Rücklage, volle AK als AfA-Basis
        "mit_quelle": _leer_oder(zeile, ne("ne_MitQuelle")),
        "afa_ohne6b": afa_ohne,
        "buchwert_ohne6b": buchwert_ohne,
        "buchwert_gub_ohne6b": _leer_oder(zeile, f'IF({gueltig},IF({t}>={kj},'
                                                 f'{ne("ne_AKGuBNeu")},0),0)'),
    }


# --- Verkäufe (Etappe 4, Projektplan Abschnitt 11) ---
#
# Verkauf zum Jahresende: Der Gebäudebuchwert ist der Prognosewert am Ende des
# Verkaufsjahrs, also nach dessen AfA. Den Erlös teilt der Anteil G+B laut
# Kaufvertrag auf, ersatzweise der Verkehrswertanteil aus dem Objektblatt.


def vspalte(key: str) -> str:
    """Spaltenbuchstabe im Blatt Verkäufe (Eingaben, dann berechnete Spalten)."""
    for i, s in enumerate(VERKAUF_FELDER + VERKAUF_SPALTEN, start=1):
        if s.key == key:
            return get_column_letter(i)
    raise KeyError(key)


def _v(key: str, zeile: int) -> str:
    return f"${vspalte(key)}{zeile}"


def verkauf_zeile(zeile: int) -> dict:
    """Formeln der berechneten Verkaufsspalten; leer ohne ObjektID oder bei unbekannter ID."""
    id_ = _v("objekt_id", zeile)

    def obj(name):
        return f"INDEX({name},MATCH({id_},obj_ID,0))"

    def wenn(ausdruck):
        return f'=IF(OR({id_}="",COUNTIF(obj_ID,{id_})=0),"",{ausdruck})'

    quote_vertrag = _v("anteil_gub", zeile)
    return {
        "vorbesitz": wenn(f"{_v('jahr', zeile)}-{obj('obj_Kaufjahr')}"),
        "buchwert_geb": wenn(f"SUMIFS(prg_Buchwert,prg_ID,{id_},prg_Jahr,{_v('jahr', zeile)})"),
        "ak_gub": wenn(obj("obj_AKGuB")),
        "nettoerloes": wenn(f"{_v('preis', zeile)}-{_v('kosten', zeile)}"),
        # Kaufvertrag vor Verkehrswert; ohne beides leer, der Status meldet es
        "quote_gub": wenn(f'IF({quote_vertrag}<>"",{quote_vertrag},'
                          f'IF({obj("obj_VKQuoteGeb")}="","",1-{obj("obj_VKQuoteGeb")}))'),
        "erloes_gub": wenn(f'IF({_v("quote_gub", zeile)}="","",'
                           f'{_v("nettoerloes", zeile)}*{_v("quote_gub", zeile)})'),
        "erloes_geb": wenn(f'IF({_v("quote_gub", zeile)}="","",'
                           f'{_v("nettoerloes", zeile)}-{_v("erloes_gub", zeile)})'),
        "gewinn_geb": wenn(f'IF({_v("quote_gub", zeile)}="","",'
                           f'{_v("erloes_geb", zeile)}-{_v("buchwert_geb", zeile)})'),
        # G+B wird nicht abgeschrieben: Buchwert = AK
        "gewinn_gub": wenn(f'IF({_v("quote_gub", zeile)}="","",'
                           f'{_v("erloes_gub", zeile)}-{_v("ak_gub", zeile)})'),
        "gewinn": wenn(f'IF({_v("quote_gub", zeile)}="","",'
                       f'{_v("gewinn_geb", zeile)}+{_v("gewinn_gub", zeile)})'),
    }


def status_verkauf(zeile: int) -> str:
    """Plausibilitätsstatus je Verkaufszeile; leer, solange keine ObjektID steht."""
    id_, jahr = _v("objekt_id", zeile), _v("jahr", zeile)
    return (
        f'=IF({id_}="","",'
        f'IF(COUNTIF(obj_ID,{id_})=0,"ObjektID unbekannt",'
        f'IF(COUNTIF(vk_ID,{id_})>1,"Objekt mehrfach verkauft",'
        f'IF({jahr}="","Verkaufsjahr fehlt",'
        f'IF(OR({jahr}<par_Startjahr,{jahr}>par_Endjahr),"Verkaufsjahr außerhalb Raster",'
        f'IF({_v("preis", zeile)}="","Verkaufspreis fehlt",'
        f'IF({_v("quote_gub", zeile)}="","Aufteilung fehlt: Anteil G+B oder Verkehrswertanteil",'
        f'IF(AND({_v("nutzung_6b", zeile)}="ja",{_v("vorbesitz", zeile)}<par_6bVorbesitz),'
        f'"{STATUS_6B_UNZULAESSIG}",'
        f'"{STATUS_OK}"))))))))'
    )


# --- Rücklagenspiegel § 6b (Etappe 5, Projektplan Abschnitt 12) ---
#
# Zeile n des Rücklagenblatts gehört zu Zeile n des Verkaufsblatts. Eine Rücklage
# entsteht nur bei § 6b ja und Status OK, und nur aus positiven Teilgewinnen:
# Gebäude und G+B sind getrennte Wirtschaftsgüter, ein Verlust des einen mindert
# die Rücklage aus dem anderen nicht, er wirkt sofort.


def _spalte_aus(spalten, key: str, versatz: int = 0) -> str:
    for i, s in enumerate(spalten, start=1 + versatz):
        if s.key == key:
            return get_column_letter(i)
    raise KeyError(key)


def rspalte(key: str) -> str:
    """Spaltenbuchstabe im Teil je Rücklage."""
    return _spalte_aus(RUECKLAGE_SPALTEN, key)


def rsspalte(key: str) -> str:
    """Spaltenbuchstabe im Spiegel je Jahr, rechts neben dem Teil je Rücklage (eine Spalte Abstand)."""
    return _spalte_aus(RUECKLAGE_JAHR_SPALTEN, key, versatz=len(RUECKLAGE_SPALTEN) + 1)


def _r(key: str, zeile: int) -> str:
    return f"${rspalte(key)}{zeile}"


def _rs(key: str, zeile: int) -> str:
    return f"${rsspalte(key)}{zeile}"


def ruecklage_zeile(zeile: int, verkauf_nr: int) -> dict:
    """Formeln einer Rücklagenzeile; verkauf_nr = Zeile im Verkaufsblatt (1 = erster Verkauf)."""
    def vk(name):
        return f"INDEX({name},{verkauf_nr})"

    def wenn(ausdruck):
        return (f'=IF(AND({vk("vk_Status")}="{STATUS_OK}",{vk("vk_6b")}="ja"),'
                f'{ausdruck},"")')

    frist = f'IF({vk("vk_6bNeubau")}="ja",par_6bFristNeubau,par_6bFrist)'
    return {
        "id": wenn(f'"RL-"&{vk("vk_ID")}'),
        "objekt_id": wenn(vk("vk_ID")),
        "jahr": wenn(vk("vk_Jahr")),
        "geb": wenn(f'MAX({vk("vk_GewinnGeb")},0)'),
        "gub": wenn(f'MAX({vk("vk_GewinnGuB")},0)'),
        "betrag": wenn(f"{_r('geb', zeile)}+{_r('gub', zeile)}"),
        "fristjahr": wenn(f"{_r('jahr', zeile)}+{frist}"),
        # Übertragung auf Neuobjekte (Etappe 6), Gebäudegewinn nur als ü1
        "ueb_geb": wenn(f"SUMIFS(ne_Ue1,ne_Quelle,{_r('id', zeile)})"),
        "ueb_gub": wenn(f"SUMIFS(ne_Ue2,ne_Quelle,{_r('id', zeile)})"
                        f"+SUMIFS(ne_Ue3,ne_Quelle,{_r('id', zeile)})"),
        # was bis zum Fristjahr nicht übertragen ist, wird dort aufgelöst
        "aufloesung": wenn(f"{_r('betrag', zeile)}-{_r('ueb_geb', zeile)}-{_r('ueb_gub', zeile)}"),
        # 6 % je vollem Jahr zwischen Bildung und Auflösung (§ 6b Abs. 7)
        "zuschlag": wenn(f"{_r('aufloesung', zeile)}*par_6bZuschlag"
                         f"*({_r('fristjahr', zeile)}-{_r('jahr', zeile)})"),
        "hinweis": wenn(f'IF(AND({_r("fristjahr", zeile)}>par_Endjahr,{_r("aufloesung", zeile)}>0),'
                        f'"Frist endet nach Prognoseende","")'),
    }


def ruecklage_jahr_zeile(zeile: int, erstes_jahr: bool) -> dict:
    """Formeln einer Zeile des Spiegels je Jahr."""
    jahr = _rs("jahr", zeile)
    # nur sauber berechnete Verkäufe; bei zu kurzer Vorbesitzzeit gilt der Gewinn, nur ohne Rücklage
    gewinne = "+".join(f'SUMIFS(vk_Gewinn,vk_Jahr,{jahr},vk_Status,"{status}")'
                       for status in (STATUS_OK, STATUS_6B_UNZULAESSIG))
    bestand_vor = "0" if erstes_jahr else _rs("bestand", zeile - 1)
    return {
        "jahr": "=par_Startjahr" if erstes_jahr else f"={_rs('jahr', zeile - 1)}+1",
        "gewinne": f"={gewinne}",
        "bildung": f"=SUMIFS(rl_Betrag,rl_Jahr,{jahr})",
        # Übertragung zum Ende des Kaufjahrs des Neuobjekts
        "uebertragung": f"=SUMIFS(ne_UeGesamt,ne_Kaufjahr,{jahr})",
        "aufloesung": f"=SUMIFS(rl_Aufloesung,rl_Fristjahr,{jahr})",
        "zuschlag": f"=SUMIFS(rl_Zuschlag,rl_Fristjahr,{jahr})",
        "steuerwirksam": (f"={_rs('gewinne', zeile)}-{_rs('bildung', zeile)}"
                          f"+{_rs('aufloesung', zeile)}+{_rs('zuschlag', zeile)}"),
        "bestand": (f"={bestand_vor}+{_rs('bildung', zeile)}-{_rs('uebertragung', zeile)}"
                    f"-{_rs('aufloesung', zeile)}"),
    }


# --- Neuobjekte und Übertragung (Etappe 6, Projektplan Abschnitt 13) ---
#
# Je Neuobjekt höchstens eine Quelle-Rücklage. Mehrere Neuobjekte können dieselbe
# Rücklage nutzen; sie bedienen sich in Zeilenreihenfolge, jede Zeile sieht nur,
# was die Zeilen darüber übrig gelassen haben. Reihenfolge der Übertragung:
# ü1 Gebäudegewinn auf Gebäude, ü2 G+B-Gewinn auf G+B, ü3 Rest G+B-Gewinn auf Gebäude.


def nspalte(key: str) -> str:
    """Spaltenbuchstabe im Blatt Neuobjekte (Eingaben, dann berechnete Spalten)."""
    return _spalte_aus(NEU_FELDER + NEU_SPALTEN, key)


def _n(key: str, zeile: int) -> str:
    return f"${nspalte(key)}{zeile}"


def _uebertragung_moeglich(zeile: int) -> str:
    """Bedingung: Zeile im Modell, Quelle bekannt, Kauf zwischen Bildungs- und Fristjahr."""
    q, kj = _n("quelle", zeile), _n("kaufjahr", zeile)
    rl = f"MATCH({q},rl_ID,0)"
    # verschachtelt, weil AND alle Argumente auswertet und MATCH sonst #N/A liefert
    return (f'IF(OR({_n("gueltig", zeile)}<>1,{q}=""),FALSE,IF(COUNTIF(rl_ID,{q})=0,FALSE,'
            f'AND({kj}>=INDEX(rl_Jahr,{rl}),{kj}<=INDEX(rl_Fristjahr,{rl}))))')


def neu_zeile(zeile: int) -> dict:
    """Formeln der berechneten Spalten je Neuobjekt; leer ohne NeuID."""
    id_, q, kj = _n("neu_id", zeile), _n("quelle", zeile), _n("kaufjahr", zeile)
    pflicht = [_n(f.key, zeile) for f in NEU_FELDER if f.pflicht]

    def wenn(ausdruck):
        return f'=IF({id_}="","",{ausdruck})'

    def darueber(key):
        """Spalte key in den Zeilen über dieser (Kopfzeile zählt als Text nicht mit)."""
        bst = nspalte(key)
        return f"${bst}$1:${bst}{zeile - 1}"

    def verbraucht(key):
        return f"SUMIFS({darueber(key)},{darueber('quelle')},{q})"

    moeglich = _uebertragung_moeglich(zeile)
    ak = f"({_n('kaufpreis', zeile)}+{_n('nebenkosten', zeile)})"
    ue1, ue2 = _n("ue1", zeile), _n("ue2", zeile)
    return {
        "gueltig": wenn(f"IF(AND(COUNTA({','.join(pflicht)})={len(pflicht)},"
                        f"COUNTIF(ne_ID,{id_})=1,COUNTIF(obj_ID,{id_})=0,"
                        f"{kj}>=par_Startjahr,{kj}<=par_Endjahr),1,0)"),
        # Nebenkosten im Verhältnis des Kaufpreises aufgeteilt und aktiviert
        "ak_gub_neu": wenn(f"{ak}*{_n('anteil_gub', zeile)}"),
        "ak_geb_neu": wenn(f"{ak}*(1-{_n('anteil_gub', zeile)})"),
        "rl_geb": wenn(f"IF({moeglich},INDEX(rl_Geb,MATCH({q},rl_ID,0))-{verbraucht('ue1')},0)"),
        "rl_gub": wenn(f"IF({moeglich},INDEX(rl_GuB,MATCH({q},rl_ID,0))"
                       f"-{verbraucht('ue2')}-{verbraucht('ue3')},0)"),
        "ue1": wenn(f"MIN({_n('rl_geb', zeile)},{_n('ak_geb_neu', zeile)})"),
        "ue2": wenn(f"MIN({_n('rl_gub', zeile)},{_n('ak_gub_neu', zeile)})"),
        "ue3": wenn(f"MIN({_n('rl_gub', zeile)}-{ue2},{_n('ak_geb_neu', zeile)}-{ue1})"),
        "ue_gesamt": wenn(f"{ue1}+{ue2}+{_n('ue3', zeile)}"),
        "afa_basis": wenn(f"{_n('ak_geb_neu', zeile)}-{ue1}-{_n('ue3', zeile)}"),
        "ak_gub": wenn(f"{_n('ak_gub_neu', zeile)}-{ue2}"),
        "mit_quelle": wenn(f'IF({q}="",0,1)'),
    }


def status_neu(zeile: int) -> str:
    """Plausibilitätsstatus je Neuobjekt; Fehler der Quelle lassen das Objekt im Modell."""
    id_, q, kj = _n("neu_id", zeile), _n("quelle", zeile), _n("kaufjahr", zeile)
    pflicht = [_n(f.key, zeile) for f in NEU_FELDER if f.pflicht]
    rl = f"MATCH({q},rl_ID,0)"
    return (
        f'=IF({id_}="","",'
        f'IF(COUNTA({",".join(pflicht)})<{len(pflicht)},"Pflichtfeld fehlt",'
        f'IF(COUNTIF(ne_ID,{id_})>1,"NeuID doppelt",'
        f'IF(COUNTIF(obj_ID,{id_})>0,"NeuID wie Bestandsobjekt",'
        f'IF(OR({kj}<par_Startjahr,{kj}>par_Endjahr),"Kaufjahr außerhalb Raster",'
        f'IF({q}="","{STATUS_OK}",'
        f'IF(COUNTIF(rl_ID,{q})=0,"Rücklage unbekannt, keine Übertragung",'
        f'IF({kj}<INDEX(rl_Jahr,{rl}),"Kauf vor Bildung der Rücklage, keine Übertragung",'
        f'IF({kj}>INDEX(rl_Fristjahr,{rl}),"Kauf nach Fristjahr, keine Übertragung",'
        f'"{STATUS_OK}")))))))))'
    )


# --- Liquidität und Auswertung (Etappe 7 und 8, Projektplan Abschnitte 14 und 15) ---
#
# Je Jahr eine Zeile, je Szenario eine Tabelle mit denselben Spalten nebeneinander:
# A Plan wie erfasst, B sofort versteuern und Kapital anlegen, C sofort versteuern und
# trotzdem kaufen, Baseline alles halten. Die Steuer rechnet mit Verlustvortrag ohne
# Mindestbesteuerung. Die Liquidität wird in allen Szenarien mit par_Alternativrendite
# verzinst. Alle Werte sind vor Finanzierung.


def _zelle(spalten, versatz: int = 0):
    """Zellbezug key, zeile in einer Tabelle, die nach versatz Spalten beginnt."""
    return lambda key, zeile: f"${_spalte_aus(spalten, key, versatz)}{zeile}"


def _tabelle(spalten_je_sz, sz: Szenario):
    """Zellbezug in der Tabelle des Szenarios; Tabellen in SZENARIEN-Reihenfolge, eine Spalte Abstand."""
    versatz = 0
    for anderes in SZENARIEN:
        spalten = spalten_je_sz(anderes)
        if anderes == sz:
            return _zelle(spalten, versatz)
        versatz += len(spalten) + 1
    raise KeyError(sz.key)


def _jahr(c, zeile: int, erstes_jahr: bool) -> str:
    return "=par_Startjahr" if erstes_jahr else f"={c('jahr', zeile - 1)}+1"


def _vor(c, key: str, zeile: int, erstes_jahr: bool) -> str:
    """Wert des Vorjahres; vor dem ersten Prognosejahr 0."""
    return "0" if erstes_jahr else c(key, zeile - 1)


def _steuer(c, zeile: int, erstes_jahr: bool, zve: str) -> dict:
    """Steuer mit Verlustvortrag: Verluste mindern spätere Gewinne, ein Rücktrag entfällt."""
    vor = _vor(c, "vortrag", zeile, erstes_jahr)
    genutzt = c("vortrag_genutzt", zeile)
    return {
        "vortrag_genutzt": f"=MIN({vor},MAX({zve},0))",
        "bemessung": f"=MAX({zve},0)-{genutzt}",
        "vortrag": f"={vor}-{genutzt}+MAX(-{zve},0)",
        "steuer": f"={c('bemessung', zeile)}*par_Steuersatz",
    }


def _summe_jahr(name: str, jahr: str, *bedingungen: str) -> str:
    return f"SUMIFS({name},prg_Jahr,{jahr}{''.join(',' + b for b in bedingungen)})"


def _summe_zeilen(name: str, jahr: str, sz: Szenario, *bedingungen: str) -> str:
    """Prognosesumme je Jahr über die Zeilen, die im Szenario vorkommen."""
    if sz == SZ_B:
        bedingungen += ("prg_MitQuelle,0",)
    return _summe_jahr(name, jahr, *bedingungen)


def liquiditaet_zeile(sz: Szenario):
    """Formelfunktion (zeile, erstes_jahr) der Liquiditätstabelle des Szenarios."""
    c = _tabelle(liq_spalten, sz)

    def zeile_formeln(zeile: int, erstes_jahr: bool) -> dict:
        t = c("jahr", zeile)
        if sz == SZ_BASELINE:
            laufend = _liq_baseline(c, zeile, t, erstes_jahr)
        else:
            laufend = _liq_plan(c, t, sz)
        kum_vor = _vor(c, "kum", zeile, erstes_jahr)
        return {
            "jahr": _jahr(c, zeile, erstes_jahr),
            **laufend,
            "ergebnis": f"={c('einnahmen', zeile)}-{c('ausgaben', zeile)}-{c('afa', zeile)}",
            # Zins auf den Stand am Vorjahresende, Zufluss zum Jahresende
            "zins": f"={kum_vor}*par_Alternativrendite",
            "zve": f"={c('ergebnis', zeile)}+{c('verkauf', zeile)}+{c('zins', zeile)}",
            **_steuer(c, zeile, erstes_jahr, c("zve", zeile)),
            "zufluss": (f"={c('einnahmen', zeile)}-{c('ausgaben', zeile)}+{c('zins', zeile)}"
                        f"+{c('verkaufserloes', zeile)}-{c('steuer', zeile)}"
                        f"-{c('kauf', zeile)}"),
            "kum": f"={kum_vor}+{c('zufluss', zeile)}",
        }
    return zeile_formeln


def _liq_plan(c, t: str, sz: Szenario) -> dict:
    """Szenarien A, B, C: Summen aus der Prognose, Verkäufe in allen gleich."""
    afa = "prg_AfA" if sz == SZ_A else "prg_AfAOhne6b"
    # A: sofort versteuerte Gewinne, Auflösung und Zuschlag; B und C: jeder Gewinn sofort
    verkauf = "rls_Steuerwirksam" if sz == SZ_A else "rls_Gewinne"
    neu = ",ne_MitQuelle,0" if sz == SZ_B else ""
    erloese = "+".join(f'SUMIFS(vk_Nettoerloes,vk_Jahr,{t},vk_Status,"{status}")'
                       for status in (STATUS_OK, STATUS_6B_UNZULAESSIG))
    return {
        "einnahmen": f"={_summe_zeilen('prg_Miete', t, sz)}+{_summe_zeilen('prg_Einnahmen', t, sz)}",
        "ausgaben": f"={_summe_zeilen('prg_Erhaltung', t, sz)}+{_summe_zeilen('prg_Ausgaben', t, sz)}",
        "afa": f"={_summe_zeilen(afa, t, sz)}",
        "verkauf": f"=SUMIFS({verkauf},rls_Jahr,{t})",
        # nur Verkäufe, deren Gewinn auch versteuert wird (gleicher Filter wie der Rücklagenspiegel)
        "verkaufserloes": f"={erloese}",
        # steuerneutraler Teil: Erlös minus Veräußerungsgewinn = Buchwert Gebäude + AK G+B
        "rueckfluss": f"={erloese}-SUMIFS(rls_Gewinne,rls_Jahr,{t})",
        "kauf": (f"=SUMIFS(ne_Kaufpreis,ne_Kaufjahr,{t},ne_Gueltig,1{neu})"
                 f"+SUMIFS(ne_Nebenkosten,ne_Kaufjahr,{t},ne_Gueltig,1{neu})"),
    }


def _liq_baseline(c, zeile: int, t: str, erstes_jahr: bool) -> dict:
    """Baseline: alle Bestandsobjekte werden gehalten, ohne Verkäufe und Neuobjekte."""
    def stamm(name, satz):
        return f'SUMIFS({name},obj_ID,"<>")*(1+{satz})^({t}-par_Basisjahr)'

    bw_vor = ('SUMIFS(obj_Restbuchwert,obj_ID,"<>")' if erstes_jahr
              else _summe_jahr("prg_BuchwertHalten", f"{t}-1"))
    return {
        "einnahmen": (f"={stamm('obj_MieteBasis', 'par_Mietsteig')}"
                      f"+{stamm('obj_EinnBasis', 'par_Mietsteig')}"),
        "ausgaben": (f"={stamm('obj_ErhBasis', 'par_Erhaltsteig')}"
                     f"+{stamm('obj_AusgBasis', 'par_Kostensteig')}"),
        # AfA = Rückgang des Buchwerts bei Halten
        "afa": f"={bw_vor}-{_summe_jahr('prg_BuchwertHalten', t)}",
        "verkauf": "=0",
        "verkaufserloes": "=0",
        "rueckfluss": "=0",
        "kauf": "=0",
    }


def auswertung_zeile(sz: Szenario):
    """Formelfunktion (zeile, erstes_jahr) der Auswertungstabelle des Szenarios."""
    c = _tabelle(aus_spalten, sz)

    def zeile_formeln(zeile: int, erstes_jahr: bool) -> dict:
        t = c("jahr", zeile)

        def liq(name):
            return f"=SUMIFS({sz.liq}_{name},{sz.liq}_Jahr,{t})"

        if sz == SZ_BASELINE:
            bestand = "prg_Neu,0"
            gebaeude, gub = "prg_BuchwertHalten", "prg_BuchwertGuB"
        else:
            bestand = "prg_Bestand,1"
            gebaeude, gub = (("prg_Buchwert", "prg_BuchwertGuB") if sz == SZ_A
                             else ("prg_BuchwertOhne6b", "prg_BuchwertGuBOhne6b"))
        ruecklage = f"=SUMIFS(rls_Bestand,rls_Jahr,{t})" if sz == SZ_A else "=0"
        return {
            "jahr": _jahr(c, zeile, erstes_jahr),
            "ergebnis": liq("Ergebnis"),
            "verkauf": liq("Verkauf"),
            "zins": liq("Zins"),
            "guv": f"={c('ergebnis', zeile)}+{c('verkauf', zeile)}+{c('zins', zeile)}",
            "steuer": liq("Steuer"),
            "nach_steuer": f"={c('guv', zeile)}-{c('steuer', zeile)}",
            "steuer_kum": f"={_vor(c, 'steuer_kum', zeile, erstes_jahr)}+{c('steuer', zeile)}",
            "verkehrswert": f"={_summe_zeilen('prg_Verkehrswert', t, sz, bestand)}",
            "buchwert": (f"={_summe_zeilen(gebaeude, t, sz, bestand)}"
                         f"+{_summe_zeilen(gub, t, sz, bestand)}"),
            "stille_reserven": f"={c('verkehrswert', zeile)}-{c('buchwert', zeile)}",
            "ruecklage": ruecklage,
            "vortrag": liq("Vortrag"),
            "liquiditaet": liq("Kum"),
            "vermoegen": f"={c('verkehrswert', zeile)}+{c('liquiditaet', zeile)}",
            # Steuer, wenn alle Objekte zum Verkehrswert verkauft und die Rücklage aufgelöst würde
            "latente_steuer": (f"=MAX({c('stille_reserven', zeile)}+{c('ruecklage', zeile)}"
                               f"-{c('vortrag', zeile)},0)*par_Steuersatz"),
            "vermoegen_netto": f"={c('vermoegen', zeile)}-{c('latente_steuer', zeile)}",
        }
    return zeile_formeln


def vergleich_zeile(praefix_art: str, name: str, art: str, sz: Szenario) -> str:
    """Kennzahl eines Szenarios im Blatt Vergleich: Endwert oder Summe über alle Jahre."""
    bereich = f"{sz.liq if praefix_art == 'liq' else sz.aus}_{name}"
    return f"=INDEX({bereich},par_Prognosejahre)" if art == "ende" else f"=SUM({bereich})"


# --- Plausibilitätsprüfungen (Etappe 9, Projektplan Abschnitt 19) ---
#
# Je Prüfung die Anzahl betroffener Zeilen. SUMPRODUCT statt COUNTIF, weil Excel und
# LibreOffice Formelzellen mit "" bei COUNTIF(…,"<>") unterschiedlich zählen.
# Verkäufe und Rücklagen haben dieselben Zeilen, rl_* und vk_* lassen sich kombinieren.


def _nicht_ok(status: str, ausnahme: str = None) -> str:
    weitere = f'*({status}<>"{ausnahme}")' if ausnahme else ""
    return f'=SUMPRODUCT(({status}<>"")*({status}<>"{STATUS_OK}"){weitere})'


def pruefung_anzahl() -> dict:
    genannt = "COUNTIF(ne_Quelle,rl_ID)"
    # Restbetrag über einen halben Cent, damit Rundung nicht anschlägt
    rest = '(rl_ID<>"")*(rl_Aufloesung>0.005)'
    gueltig = f'((vk_Status="{STATUS_OK}")+(vk_Status="{STATUS_6B_UNZULAESSIG}"))'
    im_fenster = "+".join(
        f'COUNTIFS(vk_Jahr,">"&(vk_Jahr-par_DOJahre),vk_Jahr,"<="&vk_Jahr,vk_Status,"{st}")'
        for st in (STATUS_OK, STATUS_6B_UNZULAESSIG))
    return {
        "objekte": _nicht_ok("obj_Status"),
        "verkaeufe": _nicht_ok("vk_Status", STATUS_6B_UNZULAESSIG),
        "neuobjekte": _nicht_ok("ne_Status"),
        "steuerwelt": '=IF(par_Steuerwelt="GmbH",0,1)',
        "vorbesitz": f'=SUMPRODUCT(--(vk_Status="{STATUS_6B_UNZULAESSIG}"))',
        "teiluebertrag": f"=SUMPRODUCT({rest}*({genannt}>0))",
        "ohne_reinvest": f"=SUMPRODUCT({rest}*({genannt}=0))",
        # je gültigem Verkauf: Verkäufe im Zeitraum, der mit seinem Verkaufsjahr endet
        "grundstueckshandel": f"=SUMPRODUCT({gueltig}*(({im_fenster})>par_DOGrenze))",
        "frist_ende": '=SUMPRODUCT(--(rl_Hinweis<>""))',
        "liquiditaet": "=SUMPRODUCT(--(liq_Kum<-0.005))",
        "verkehrswert": f'=SUMPRODUCT((obj_Status="{STATUS_OK}")*(obj_Verkehrswert=""))',
    }


def pruefung_ergebnis(zeile: int) -> str:
    """Ergebnis einer Prüfzeile: OK oder die Art (Fehler, Warnung, Hinweis)."""
    return f'=IF($D{zeile}=0,"{STATUS_OK}",$C{zeile})'


def pruefung_summen() -> dict:
    return {
        "pr_Fehler": f'=COUNTIF(pr_Ergebnis,"{FEHLER}")',
        "pr_Warnungen": f'=COUNTIF(pr_Ergebnis,"{WARNUNG}")',
        "pr_Gesamt": (f'=IF(pr_Fehler>0,pr_Fehler&" {FEHLER}",'
                      f'IF(pr_Warnungen>0,pr_Warnungen&" Warnung(en)","{STATUS_OK}"))'),
    }
