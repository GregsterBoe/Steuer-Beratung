"""Eingaben aus einer älteren Mappe in eine neu erzeugte übernehmen (--uebernehmen).

Übernommen werden nur eingetippte Werte (keine Formeln): Formeln der alten Mappe sind
Annahmen, Verknüpfungen oder Rechnung, die neue Mappe bringt ihre eigenen mit. Spalten
werden über die Überschrift zugeordnet, nicht über die Position; Felder, die es in der
alten Version noch nicht gab, behalten ihre Vorbelegung.

Werte, die beim Einlesen aus der Kostenstellen-Datei entstanden sind (gleich der
ausgeblendeten Kopie „eingelesen“), übernimmt die Mappe nicht, sobald die neue Mappe dafür
selbst einen eingelesenen Wert hat: dann gilt der neue Stand der Kostenstellen-Datei. Die
Kostenstellenblätter selbst werden nicht übernommen, sie kommen aus der Kostenstellen-Datei.
"""

from datetime import date, datetime

from openpyxl import load_workbook

from . import formeln
from .bwa import POSTEN, ZUORDNUNG
from .modelle import (ANLAGE_FELDER, ANLAGE_SPALTEN, NEU_FELDER, NEU_SPALTEN,
                      NEUKAUF_POSITIONEN, OBJEKT_EINGELESEN, OBJEKT_FELDER, PARAMETER,
                      STATUS_UEBERSCHRIFT, VARIANTEN_KOPF, VERKAUF_FELDER, VERKAUF_SPALTEN,
                      ZUORDNUNG_KOST1, prognosejahre)

# Anlagen: aus dem neuen Anlagenverzeichnis kommen Beträge und Daten; von Hand gepflegt
# sind die Zuordnung zum Objekt und eine eingetippte AfA
ANLAGEN_VON_HAND = ("objekt_id", "zuordnung", "afa")
VARIANTEN_ERSTE = 5


def _ist_eingabe(wert) -> bool:
    """Eingetippter Wert: Zahl, Text (keine Formel), Datum; leer zählt nicht."""
    if isinstance(wert, str):
        return wert.strip() != "" and not wert.startswith("=")
    return isinstance(wert, (bool, int, float, date, datetime))


def _kopftext(wert) -> str:
    """Überschrift ohne Pflicht-Stern und Leerraum, für den Vergleich alt gegen neu."""
    if not isinstance(wert, str):
        return ""
    return " ".join(wert.replace("*", " ").split())


class Bericht:
    def __init__(self):
        self.zeilen = []
        self.anzahl = {}

    def zaehle(self, blatt: str, n: int = 1) -> None:
        if n:
            self.anzahl[blatt] = self.anzahl.get(blatt, 0) + n

    def hinweis(self, text: str) -> None:
        self.zeilen.append(text)

    def ausgabe(self) -> list:
        kopf = [f"  {blatt}: {n} Wert(e)" for blatt, n in self.anzahl.items()]
        return (["übernommen aus der alten Mappe:"] + (kopf or ["  keine Eingaben gefunden"])
                + [f"  Hinweis: {z}" for z in self.zeilen])


def _zeilen(wb, name: str, erste: int, letzte: int) -> range:
    """Datenzeilen der Tabelle laut benanntem Bereich der Mappe (z. B. obj_ID); ohne den
    Namen erste bis letzte. Unter den Tabellen stehen Hinweistexte, die nicht dazugehören."""
    try:
        _, ref = next(iter(wb.defined_names[name].destinations))
        von, bis = ref.replace("$", "").split(":")
        return range(int("".join(c for c in von if c.isdigit())),
                     int("".join(c for c in bis if c.isdigit())) + 1)
    except (KeyError, StopIteration, ValueError):
        return range(erste, letzte + 1)


def _parameter_alt(ws) -> dict:
    """Parameterwerte der alten Mappe je Name (Spalte C), Wert in Spalte B."""
    werte = {}
    for zeile in range(1, ws.max_row + 1):
        name = ws.cell(row=zeile, column=3).value
        if isinstance(name, str) and name.startswith("par_"):
            werte[name] = (zeile, ws.cell(row=zeile, column=2).value)
    return werte


def _parameter(alt, neu, bericht: Bericht, geschuetzt: set) -> None:
    if "Parameter" not in alt.sheetnames:
        return
    ws_alt, ws_neu = alt["Parameter"], neu["Parameter"]
    alte = _parameter_alt(ws_alt)
    neue = _parameter_alt(ws_neu)
    for name, (_, wert) in alte.items():
        if name not in neue:
            if _ist_eingabe(wert):
                bericht.hinweis(f"Parameter {name} gibt es nicht mehr (alter Wert {wert}).")
            continue
        zeile, neu_wert = neue[name]
        if not _ist_eingabe(wert) or wert == neu_wert:
            continue
        if name in geschuetzt or not _ist_eingabe(neu_wert) and neu_wert is not None:
            # Basisjahr, beim Erzeugen gesetzte Werte und berechnete Parameter bleiben
            bericht.hinweis(f"Parameter {name}: alter Wert {wert} nicht übernommen, "
                            f"die neue Mappe rechnet mit {neu_wert}.")
            continue
        ws_neu.cell(row=zeile, column=2, value=wert)
        bericht.zaehle("Parameter")


def _spalten(ws_alt, kopfzeile: int, felder, berechnet: list, blatt: str,
             bericht: Bericht) -> dict:
    """Spalte alt → Spalte neu je Eingabefeld, zugeordnet über die Überschrift.

    Die Eingabefelder der alten Mappe reichen bis zur ersten berechneten Spalte; dort
    gefundene Überschriften ohne Gegenstück in der neuen Version meldet der Bericht.
    """
    neu = {_kopftext(f.ueberschrift): i for i, f in enumerate(felder, start=1)}
    ende = {_kopftext(s.ueberschrift) for s in berechnet} | {STATUS_UEBERSCHRIFT}
    zuordnung = {}
    for spalte in range(1, ws_alt.max_column + 1):
        text = _kopftext(ws_alt.cell(row=kopfzeile, column=spalte).value)
        if text in ende:
            break
        if text in neu:
            zuordnung[spalte] = neu[text]
        elif text:
            bericht.hinweis(f"{blatt}: Spalte „{text}“ gibt es nicht mehr, nicht übernommen.")
    return zuordnung


def _zeilenweise(alt, neu, blatt: str, felder, berechnet: list, kapazitaet: int,
                 bericht: Bericht, bereich: str) -> None:
    """Verkäufe und Neuobjekte: Zeile n bleibt Zeile n (Rücklage und Reinvestition hängen
    an der Zeilennummer)."""
    if blatt not in alt.sheetnames:
        return
    ws_alt, ws_neu = alt[blatt], neu[blatt]
    spalten = _spalten(ws_alt, 1, felder, berechnet, blatt, bericht)
    for zeile in _zeilen(alt, bereich, 2, ws_alt.max_row):
        if zeile > kapazitaet + 1:
            bericht.hinweis(f"{blatt}: mehr Zeilen als in der neuen Mappe, ab Zeile {zeile} "
                            "nicht übernommen.")
            break
        # von Hand erfasste Zeile: leere Felder waren gewollt leer, die neue Mappe soll dort
        # keine Annahme rechnen (Felder, die es alt nicht gab, behalten ihre Vorbelegung)
        erfasst = _ist_eingabe(ws_alt.cell(row=zeile, column=1).value)
        for alt_sp, neu_sp in spalten.items():
            wert = ws_alt.cell(row=zeile, column=alt_sp).value
            if _ist_eingabe(wert):
                if ws_neu.cell(row=zeile, column=neu_sp).value != wert:
                    ws_neu.cell(row=zeile, column=neu_sp, value=wert)
                    bericht.zaehle(blatt)
            elif erfasst and wert is None:
                ws_neu.cell(row=zeile, column=neu_sp).value = None


def _id_zeilen(ws, spalte: int, erste: int, letzte: int) -> tuple:
    """(Zeile je eingetippter ID, freie Zeilen) einer Eingabetabelle."""
    belegt, frei = {}, []
    for zeile in range(erste, letzte + 1):
        wert = ws.cell(row=zeile, column=spalte).value
        if _ist_eingabe(wert):
            belegt[str(wert).strip()] = zeile
        else:
            frei.append(zeile)
    return belegt, frei


def _zielzeile(schluessel: str, belegt: dict, frei: list, ws, spalte: int):
    """Zeile der neuen Mappe mit derselben ID; sonst die nächste freie, mit ID."""
    if schluessel in belegt:
        return belegt[schluessel]
    if not frei:
        return None
    zeile = frei.pop(0)
    ws.cell(row=zeile, column=spalte, value=schluessel)
    belegt[schluessel] = zeile
    return zeile


def _objekte(alt, neu, bericht: Bericht, kapazitaet: int) -> None:
    if "Objekte" not in alt.sheetnames:
        return
    ws_alt, ws_neu = alt["Objekte"], neu["Objekte"]
    spalten = _spalten(ws_alt, 1, OBJEKT_FELDER, [], "Objekte", bericht)
    # ausgeblendete Kopien der eingelesenen Werte, alt über die Überschrift, neu über formeln
    titel = {f"eingelesen: {f.ueberschrift}": f.key for f in OBJEKT_FELDER
             if f.key in OBJEKT_EINGELESEN}
    import_alt = {}
    for spalte in range(1, ws_alt.max_column + 1):
        key = titel.get(ws_alt.cell(row=1, column=spalte).value)
        if key:
            import_alt[key] = spalte
    import_neu = {key: formeln.ispalte(key) for key in titel.values()}
    key_je_spalte = {i: f.key for i, f in enumerate(OBJEKT_FELDER, start=1)}
    id_alt = next((a for a, n in spalten.items() if key_je_spalte[n] == "objekt_id"), None)
    if id_alt is None:
        bericht.hinweis("Objekte: Spalte ObjektID nicht gefunden, Blatt nicht übernommen.")
        return
    belegt, frei = _id_zeilen(ws_neu, 1, 2, kapazitaet + 1)
    for zeile in _zeilen(alt, "obj_ID", 2, ws_alt.max_row):
        oid = ws_alt.cell(row=zeile, column=id_alt).value
        if not _ist_eingabe(oid):
            continue
        ziel = _zielzeile(str(oid).strip(), belegt, frei, ws_neu, 1)
        if ziel is None:
            bericht.hinweis(f"Objekte: kein Platz mehr für {oid}.")
            continue
        for alt_sp, neu_sp in spalten.items():
            key = key_je_spalte[neu_sp]
            if key == "objekt_id":
                continue
            wert = ws_alt.cell(row=zeile, column=alt_sp).value
            if not _ist_eingabe(wert):
                continue
            if key in import_alt:
                eingelesen = ws_alt.cell(row=zeile, column=import_alt[key]).value
                frisch = ws_neu[f"{import_neu[key]}{ziel}"].value
                if wert == eingelesen and frisch is not None:
                    continue   # alter Stand der Kostenstellen-Datei, neuer gilt
            if ws_neu.cell(row=ziel, column=neu_sp).value != wert:
                ws_neu.cell(row=ziel, column=neu_sp, value=wert)
                bericht.zaehle("Objekte")


def _anlagen(alt, neu, bericht: Bericht) -> None:
    if "Anlagen" not in alt.sheetnames:
        return
    ws_alt, ws_neu = alt["Anlagen"], neu["Anlagen"]
    spalten = _spalten(ws_alt, 2, ANLAGE_FELDER, ANLAGE_SPALTEN, "Anlagen", bericht)
    key_je_spalte = {i: f.key for i, f in enumerate(ANLAGE_FELDER, start=1)}
    alt_je_key = {key_je_spalte[n]: a for a, n in spalten.items()}
    if "nr" not in alt_je_key:
        return
    letzte = ws_neu.max_row
    belegt, frei = _id_zeilen(ws_neu, 1, 3, letzte)
    # ohne neues Anlagenverzeichnis: alle Anlagen der alten Mappe übernehmen
    alles = not belegt
    fehlen = 0
    for zeile in _zeilen(alt, "anl_Nr", 3, ws_alt.max_row):
        nr = ws_alt.cell(row=zeile, column=alt_je_key["nr"]).value
        if not _ist_eingabe(nr):
            continue
        nr = str(nr).strip()
        if not alles and nr not in belegt:
            fehlen += 1
            continue
        ziel = _zielzeile(nr, belegt, frei, ws_neu, 1)
        if ziel is None:
            continue
        zuordnung = ws_alt.cell(row=zeile, column=alt_je_key["zuordnung"]).value \
            if "zuordnung" in alt_je_key else None
        for key, alt_sp in alt_je_key.items():
            if key == "nr" or not alles and key not in ANLAGEN_VON_HAND:
                continue
            # ObjektID aus KOST1 kommt aus dem neuen Verzeichnis, nur Zuordnung von Hand zählt
            if not alles and key in ("objekt_id", "zuordnung") and zuordnung == ZUORDNUNG_KOST1:
                continue
            wert = ws_alt.cell(row=zeile, column=alt_sp).value
            neu_sp = spalten[alt_sp]
            if _ist_eingabe(wert) and ws_neu.cell(row=ziel, column=neu_sp).value != wert:
                ws_neu.cell(row=ziel, column=neu_sp, value=wert)
                bericht.zaehle("Anlagen")
    if fehlen:
        bericht.hinweis(f"Anlagen: {fehlen} Anlage(n) der alten Mappe fehlen im neuen "
                        "Anlagenverzeichnis und wurden nicht übernommen.")


def _jahreswerte(ws_alt, ws_neu, alt_zeile: int, neu_zeile: int, j0_alt: int, imp_alt: int,
                 j0_neu: int, imp_neu: int, n_alt: int, n_neu: int, versatz: int) -> int:
    """Eingetippte Jahreswerte einer Zeile; versatz = Basisjahr alt − Basisjahr neu."""
    anzahl = 0
    for i in range(n_alt):
        k = i + versatz
        if not 0 <= k < n_neu:
            continue
        wert = ws_alt.cell(row=alt_zeile, column=j0_alt + i).value
        if not _ist_eingabe(wert):
            continue
        eingelesen = ws_alt.cell(row=alt_zeile, column=imp_alt + i).value
        frisch = ws_neu.cell(row=neu_zeile, column=imp_neu + k).value
        if wert == eingelesen and frisch is not None:
            continue
        ws_neu.cell(row=neu_zeile, column=j0_neu + k, value=wert)
        anzahl += 1
    return anzahl


def _afa_plan(alt, neu, bericht: Bericht, versatz: int, jahre_alt: int,
              kapazitaet: int) -> None:
    if "AfA-Plan" not in alt.sheetnames:
        return
    ws_alt, ws_neu = alt["AfA-Plan"], neu["AfA-Plan"]
    jahre, j0 = prognosejahre(), 4
    belegt, frei = _id_zeilen(ws_neu, 1, 3, kapazitaet + 2)
    for zeile in _zeilen(alt, "afp_ID", 3, ws_alt.max_row):
        oid = ws_alt.cell(row=zeile, column=1).value
        if not _ist_eingabe(oid):
            continue
        ziel = _zielzeile(str(oid).strip(), belegt, frei, ws_neu, 1)
        if ziel is None:
            continue
        bericht.zaehle("AfA-Plan", _jahreswerte(
            ws_alt, ws_neu, zeile, ziel, j0, j0 + jahre_alt + 1, j0, j0 + jahre + 1,
            jahre_alt, jahre, versatz))


def _neukauf(alt, neu, bericht: Bericht, versatz: int, jahre_alt: int,
             kapazitaet: int) -> None:
    if "Neukauf-KSt" not in alt.sheetnames:
        return
    ws_alt, ws_neu = alt["Neukauf-KSt"], neu["Neukauf-KSt"]
    je, erste = len(NEUKAUF_POSITIONEN), formeln.NK_ERSTE
    j0, _, imp, n = formeln.nk_spalten()
    n_alt = jahre_alt + 1
    imp_alt = j0 + 2 * n_alt + 2
    belegt, frei = {}, []
    for block in range(kapazitaet):
        kopf = erste + block * je
        wert = ws_neu.cell(row=kopf, column=1).value
        if _ist_eingabe(wert):
            belegt[str(wert).strip()] = kopf
        else:
            frei.append(kopf)
    kopf_alt = erste
    ende = _zeilen(alt, "nk_ID", erste, ws_alt.max_row).stop
    while kopf_alt < ende:
        kst = ws_alt.cell(row=kopf_alt, column=1).value
        if _ist_eingabe(kst):
            ziel = _zielzeile(str(kst).strip(), belegt, frei, ws_neu, 1)
            if ziel is not None:
                name = ws_alt.cell(row=kopf_alt, column=2).value
                if _ist_eingabe(name) and ws_neu.cell(row=ziel, column=2).value is None:
                    ws_neu.cell(row=ziel, column=2, value=name)
                for k in range(je):
                    bericht.zaehle("Neukauf-KSt", _jahreswerte(
                        ws_alt, ws_neu, kopf_alt + k, ziel + k, j0, imp_alt, j0, imp,
                        n_alt, n, versatz))
        kopf_alt += je


def _zuordnung(alt, neu, bericht: Bericht) -> None:
    if ZUORDNUNG not in alt.sheetnames or ZUORDNUNG not in neu.sheetnames:
        return
    ws_alt, ws_neu = alt[ZUORDNUNG], neu[ZUORDNUNG]
    if _ist_eingabe(ws_alt["B4"].value) and ws_alt["B4"].value != ws_neu["B4"].value:
        ws_neu["B4"] = ws_alt["B4"].value
        bericht.zaehle(ZUORDNUNG)
    zeile_neu = {ws_neu.cell(row=7 + i, column=1).value: 7 + i for i in range(len(POSTEN))}
    for zeile in range(7, ws_alt.max_row + 1):
        text = ws_alt.cell(row=zeile, column=1).value
        wert = ws_alt.cell(row=zeile, column=2).value
        if text in zeile_neu and _ist_eingabe(wert) \
                and ws_neu.cell(row=zeile_neu[text], column=2).value != wert:
            ws_neu.cell(row=zeile_neu[text], column=2, value=wert)
            bericht.zaehle(ZUORDNUNG)


def _varianten(alt, neu, bericht: Bericht) -> None:
    if "Varianten" not in alt.sheetnames:
        return
    ws_alt, ws_neu = alt["Varianten"], neu["Varianten"]
    breite = len(VARIANTEN_KOPF)
    for zeile in _zeilen(alt, "var_Bezeichnung", VARIANTEN_ERSTE, ws_alt.max_row):
        werte = [ws_alt.cell(row=zeile, column=s).value for s in range(1, breite + 1)]
        if any(_ist_eingabe(w) for w in werte):
            for s, w in enumerate(werte, start=1):
                ws_neu.cell(row=zeile, column=s, value=w)
            bericht.zaehle("Varianten")


def uebernehmen(neu, pfad, geschuetzt=(), kapazitaet=None) -> list:
    """Eingaben der Mappe unter pfad in die Mappe neu schreiben; liefert den Bericht.

    geschuetzt: Parameter, die beim Erzeugen gesetzt wurden und nicht überschrieben werden.
    """
    from .modelle import Kapazitaet
    kapazitaet = kapazitaet or Kapazitaet()
    alt = load_workbook(pfad)
    bericht = Bericht()
    basis_neu = next(p.wert for p in PARAMETER if p.name == "par_Basisjahr")
    alte_par = _parameter_alt(alt["Parameter"]) if "Parameter" in alt.sheetnames else {}
    basis_alt = alte_par.get("par_Basisjahr", (None, basis_neu))[1]
    jahre_alt = alte_par.get("par_Prognosejahre", (None, prognosejahre()))[1]
    if not isinstance(basis_alt, int):
        basis_alt = basis_neu
    if not isinstance(jahre_alt, int):
        jahre_alt = prognosejahre()
    if basis_alt != basis_neu:
        bericht.hinweis(f"Basisjahr alt {basis_alt}, neu {basis_neu}: Jahreswerte in AfA-Plan "
                        "und Neukauf-KSt sind auf ihr Kalenderjahr verschoben; Werte der "
                        "Objekte (Basisjahr) bitte prüfen.")

    _parameter(alt, neu, bericht, set(geschuetzt) | {"par_Basisjahr"})
    _objekte(alt, neu, bericht, kapazitaet.objekte)
    _anlagen(alt, neu, bericht)
    _afa_plan(alt, neu, bericht, basis_alt - basis_neu, jahre_alt, kapazitaet.objekte)
    _neukauf(alt, neu, bericht, basis_alt - basis_neu, jahre_alt, kapazitaet.neukauf)
    _zeilenweise(alt, neu, "Verkäufe", VERKAUF_FELDER, VERKAUF_SPALTEN,
                 kapazitaet.verkaeufe, bericht, "vk_ID")
    _zeilenweise(alt, neu, "Neuobjekte", NEU_FELDER, NEU_SPALTEN,
                 kapazitaet.neuobjekte, bericht, "ne_ID")
    _zuordnung(alt, neu, bericht)
    _varianten(alt, neu, bericht)
    return bericht.ausgabe()
